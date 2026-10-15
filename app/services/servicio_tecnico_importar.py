"""Servicio Técnico: traer equipos desde Excel, CSV y PDF («importador»).

Flujo (un lote = una fila de st_importaciones):
  1. crear_lote      guarda cada archivo en st_documentos (aún sin equipo) y lo
                     lee con reglas fijas: tablas de Excel/CSV por sus títulos;
                     hojas tipo ficha, PDF con texto y .txt con el lector fijo.
  2. leer_siguiente  con IA, una parte por llamada (cada llamada es corta): la
                     IA ubica las columnas que las reglas no entendieron y lee
                     las fichas y los PDF. Lo que propone y no está en el
                     documento se descarta (servicio_tecnico_ia).
  3. revisar         estado de cada fila (lista / revisar / no se puede), equipos
                     que ya existen (serial o IMEI), cliente del CRM (correo,
                     teléfono o nombre). Nada se guarda todavía.
  4. importar        una sola transacción: clientes nuevos en el CRM, equipos (o
                     completa los que ya existían, sin pisar nada), hoja de
                     vida, plan y recordatorios, documentos adjuntos.
  5. deshacer        retira lo que creó el lote (activo = FALSE) y devuelve los
                     campos que había completado: nada se borra.

Carpeta de cada equipo: la empresa a la que pertenece (sus equipos van juntos)
o «Particulares» si es de una persona. Se decide con reglas
(servicio_tecnico_empresas) y, si el nombre es dudoso o el documento no dice la
empresa con su etiqueta, con la IA en un paso más de la lectura (uno por archivo).

Privacidad: los datos del cliente (nombre, teléfono, correo, documento) se
sacan SOLO con reglas fijas. La única excepción es la carpeta: el encabezado del
documento y los nombres dudosos van a la IA LOCAL del negocio (nunca a la nube),
sin correos, teléfonos ni documentos, y con los nombres de persona tapados. Las
columnas de claves y contraseñas no se importan.
"""

import csv
import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime, timedelta, timezone

from database import get_db_cursor
from services import servicio_tecnico_empresas as emp
from services import servicio_tecnico_ficha as fichas
from services import servicio_tecnico_lector as lector
from services import servicio_tecnico_service as st
from services import servicio_tecnico_tipos as tipos
from services import servicio_tecnico_versiones as ver
from services.servicio_tecnico_service import ErrorServicio

# ── Límites ─────────────────────────────────────────────────────
EXTENSIONES = {'xlsx': 'excel', 'xlsm': 'excel', 'csv': 'csv', 'txt': 'texto', 'pdf': 'pdf'}
ACEPTA = '.xlsx,.xlsm,.csv,.txt,.pdf'
ARCHIVOS_MAX = 30
ARCHIVO_MAX_BYTES = 20 * 1024 * 1024
EQUIPOS_MAX = 500                  # equipos por lote
FILAS_LEER_MAX = 5000              # filas de datos que se leen de una tabla
HOJAS_MAX = 40
COLUMNAS_MAX = 80
PAGINAS_MAX = 60
HISTORIAL_MAX = 60                 # mantenimientos por equipo
LEYENDO_SEGUNDOS = 150             # una lectura con IA «en curso» más vieja se repite
INTENTOS_MAX = 8                   # reintentos mientras el motor de IA se prepara
ESTADOS_LOTE = {
    'leyendo': 'Leyendo los archivos',
    'revision': 'Por revisar',
    'importada': 'Importada',
    'deshecha': 'Deshecha',
    'descartada': 'Descartada',
}
VENCIDOS = (('siguiente', 'Pasarlos al siguiente según la frecuencia'),
            ('hoy', 'Ponerlos para hoy, para llamar al cliente'))


# ── Texto ───────────────────────────────────────────────────────
def _norm(texto):
    """minúsculas, sin tildes ni signos: «Nº de Serie:» → «n de serie»."""
    t = unicodedata.normalize('NFD', str(texto or '').lower())
    t = ''.join(ch for ch in t if unicodedata.category(ch) != 'Mn')
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', t).split())


def _celda(v):
    """Valor de una celda como texto: fechas dd/mm/aaaa y enteros sin «.0»."""
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y') if (v.hour, v.minute) == (0, 0) else v.strftime('%d/%m/%Y %H:%M')
    if isinstance(v, date):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, bool):
        return 'Sí' if v else 'No'
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return ' '.join(str(v).split())


def _letra(j):
    letras = ''
    j += 1
    while j:
        j, resto = divmod(j - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


# ── Fechas, números, teléfonos ──────────────────────────────────
_MES = {'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6, 'jul': 7, 'ago': 8, 'sep': 9, 'set': 9,
        'oct': 10, 'nov': 11, 'dic': 12}
_RE_FECHAS = (
    (re.compile(r'(?<!\d)(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})(?!\d)'), 'amd'),
    (re.compile(r'(?<!\d)(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4}|\d{2})(?![\d/\-.]*\d)'), 'dma'),
    (re.compile(r'(?<!\d)(\d{1,2})[-/.](ene|feb|mar|abr|may|jun|jul|ago|sep|set|oct|nov|dic)[a-z]*\.?[-/.]'
                r'(\d{4}|\d{2})(?!\d)'), 'dMa'),
    (re.compile(r'(?<!\d)(\d{1,2})\s*(?:de\s+)?(ene|feb|mar|abr|may|jun|jul|ago|sep|set|oct|nov|dic)[a-z]*\.?\s*'
                r'(?:de\s+|del\s+)?(\d{4})(?!\d)'), 'dMa'),
    (re.compile(r'(?<![a-z])(ene|feb|mar|abr|may|jun|jul|ago|sep|set|oct|nov|dic)[a-z]*\.?\s*(?:de\s+|del\s+)?'
                r'(\d{4})(?!\d)'), 'Ma'),
)


def _armar_fecha(a, m, d):
    try:
        a, m, d = int(a), int(m), int(d)
        if a < 100:
            a += 2000 if a <= date.today().year % 100 + 1 else 1900
        f = date(a, m, d)
    except (TypeError, ValueError):
        return None
    return f if 1990 <= f.year <= 2100 else None


def fechas_en(texto):
    """[(fecha, inicio, fin)] de las fechas de un texto, en orden. Día
    primero, como en Colombia (si el «mes» pasa de 12, se lee mes/día)."""
    t = str(texto or '').lower()
    ocupado, salida = [], []
    for patron, orden in _RE_FECHAS:
        for m in patron.finditer(t):
            if any(m.start() < fin and m.end() > ini for ini, fin in ocupado):
                continue
            if orden == 'amd':
                f = _armar_fecha(m.group(1), m.group(2), m.group(3))
            elif orden == 'dma':
                f = _armar_fecha(m.group(3), m.group(2), m.group(1)) or _armar_fecha(m.group(3), m.group(1), m.group(2))
            elif orden == 'dMa':
                f = _armar_fecha(m.group(3), _MES[m.group(2)], m.group(1))
            else:
                f = _armar_fecha(m.group(2), _MES[m.group(1)], 1)
            if f:
                ocupado.append((m.start(), m.end()))
                salida.append((f, m.start(), m.end()))
    salida.sort(key=lambda x: x[1])
    return salida


def _fecha_valor(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 30000 <= v <= 60000:                      # número de serie de fecha de Excel
            return date(1899, 12, 30) + timedelta(days=int(v))
        return None
    encontradas = fechas_en(_celda(v))
    return encontradas[0][0] if encontradas else None


def _iso(valor):
    if not valor:
        return None
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor)[:10])
    except ValueError:
        return None


def _numero(v):
    if v in (None, ''):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    t = re.sub(r'[^\d,.]', '', str(v))
    if not t:
        return None
    if re.fullmatch(r'\d{1,3}(\.\d{3})+(,\d+)?', t):
        t = t.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'\d{1,3}(,\d{3})+(\.\d+)?', t):
        t = t.replace(',', '')
    else:
        t = t.replace(',', '.')
    try:
        return float(t)
    except ValueError:
        return None


_CORREO = re.compile(r'[\w.+-]+@[\w-]+(?:\.[\w-]+)+')
_TELEFONO_TXT = re.compile(r'(?<!\d)(?:\+?57[\s-]?)?(?:3\d{2}|60\d)[\s.-]?\d{3}[\s.-]?\d{4}(?!\d)')


def _telefono(v):
    """Celular o fijo de Colombia: 10 dígitos, o None."""
    digitos = re.sub(r'\D', '', _celda(v))
    if len(digitos) == 12 and digitos.startswith('57'):
        digitos = digitos[2:]
    return digitos if len(digitos) == 10 and digitos[0] in '36' else None


def _formato_tel(digitos):
    return f'{digitos[:3]} {digitos[3:6]} {digitos[6:]}' if digitos and len(digitos) == 10 else digitos


def _correo(v):
    m = _CORREO.search(_celda(v))
    return m.group(0).lower()[:150] if m else None


def meses_de(texto):
    """Frecuencia en meses: «6», «cada 4 meses», «semestral», «2 veces al año»…"""
    t = _norm(texto)
    if not t:
        return None
    m = re.search(r'cada\s+(\d{1,2})\s+mes', t) or re.search(r'(?<!\d)(\d{1,2})\s+mes', t)
    if m and 1 <= int(m.group(1)) <= 36:
        return int(m.group(1))
    for n, palabras in ((1, ('mensual',)), (2, ('bimestral',)), (3, ('trimestral',)), (4, ('cuatrimestral',)),
                        (6, ('semestral', 'semestre')), (12, ('anual', 'cada ano', 'una vez al ano', '1 vez al ano'))):
        if any(re.search(r'(?<![a-z])' + p + r'(?![a-z])', t) for p in palabras):
            return n
    m = re.search(r'(\d{1,2})\s+veces\s+al\s+ano', t)
    if m and int(m.group(1)) and 12 % int(m.group(1)) == 0:
        return 12 // int(m.group(1))
    if re.fullmatch(r'\d{1,2}', t) and 1 <= int(t) <= 36:
        return int(t)
    return None


_CORRECTIVO = re.compile(r'repar|arregl|reempl|falla|dano|no enciende|no prende|no carga|garantia|soldad|reballing|'
                         r'sustitu|cambio de (?:pantalla|teclado|bateria|disco|fuente|board|tarjeta|puerto|pin|display|'
                         r'flex|camara|parlante|microfono|bisagra|carcasa|cargador|tactil|vidrio|memoria|ram)')
_PREVENTIVO = re.compile(r'prevent|limpieza|mantenimiento|mtto|revisi|chequeo|ajuste|optimiz|actualiz|soplad|'
                         r'aspirad|desinfec|pasta termica|lubric')


def tipo_mantenimiento(texto):
    """'preventivo', 'correctivo' o None según lo que dice el texto."""
    n = _norm(texto)
    if not n:
        return None
    dice_prev, dice_corr = 'prevent' in n, 'correct' in n
    if dice_prev != dice_corr:
        return 'preventivo' if dice_prev else 'correctivo'
    if _CORRECTIVO.search(n):
        return 'correctivo'
    if _PREVENTIVO.search(n):
        return 'preventivo'
    return 'correctivo' if re.search(r'(?<![a-z])cambi', n) else None


def _sin_rotulo(desc):
    """«Preventivo - Limpieza…» → «Limpieza…» (el tipo ya queda aparte)."""
    resto = re.sub(r'(?i)^\s*(?:mantenimiento\s+)?(?:preventivo|correctivo)\s*[-–:,.|]?\s*', '', desc)
    return resto if len(resto) >= 4 else desc


# ── Títulos de columna → campo ──────────────────────────────────
# Destino de cada columna. «notas» guarda «Título: valor» en las notas del
# equipo (así ningún dato se pierde); «ignorar» descarta (números de fila);
# «secreto» no se importa nunca (claves, contraseñas, patrones).
_SINONIMOS = (
    ('tipo', 'tipo|tipo de equipo|tipo equipo|clase de equipo|clase|categoria|dispositivo|tipo de dispositivo|'
             'tipo dispositivo|elemento|tipo de activo|tipo de elemento'),
    ('descripcion', 'equipo|descripcion del equipo|descripcion equipo|articulo|producto|marca y modelo|marca modelo|'
                    'referencia del equipo|equipo marca modelo|nombre del producto'),
    ('marca', 'marca|fabricante|brand|manufacturer|marca equipo|marca del equipo|marca cpu|marca portatil|'
              'marca celular|marca torre'),
    ('modelo', 'modelo|model|referencia|ref|modelo equipo|modelo del equipo|modelo cpu|modelo portatil|'
               'modelo celular|referencia equipo|linea'),
    ('serial', 'serial|serie|no serie|no de serie|numero de serie|nro serie|nro de serie|n serie|num serie|sn|s n|'
               'serial number|service tag|serial equipo|serial del equipo|serial cpu|serial portatil|serial torre|'
               'serie equipo|serial celular|n de serie'),
    ('imei', 'imei|imei 1|imei1|imei principal'),
    ('extra.imei2', 'imei 2|imei2|imei secundario'),
    ('color', 'color'),
    ('sistema_operativo', 'sistema operativo|so|s o|os|sistema|version windows|windows|version de windows|'
                          'sistema op|sist operativo'),
    ('procesador', 'procesador|cpu|proc|processor|micro|microprocesador|chip'),
    ('ram', 'ram|memoria|memoria ram|mem|ram gb|memoria gb|capacidad ram|cantidad ram|mem ram'),
    ('almacenamiento', 'disco|disco duro|almacenamiento|dd|hdd|ssd|capacidad disco|capacidad de disco|storage|'
                       'capacidad|rom|almacenamiento interno|memoria interna|unidad de disco|disco gb|'
                       'capacidad almacenamiento|disco solido'),
    ('pantalla', 'pantalla|tamano pantalla|tamano de pantalla|pulgadas|display|tamano'),
    ('ficha_codigo', 'nro ficha tecnica|no ficha tecnica|numero ficha tecnica|numero de ficha tecnica|'
                     'nro de ficha tecnica|n ficha tecnica|codigo ficha tecnica|codigo de ficha tecnica|nro ficha|'
                     'no ficha|numero de ficha|nro de ficha|codigo ficha|codigo de ficha|ficha nro|ficha no'),
    ('especificaciones', 'caracteristicas|especificaciones|configuracion|caracteristicas tecnicas|'
                         'especificaciones tecnicas|descripcion tecnica|hardware|ficha tecnica|componentes'),
    ('extra.tipo_disco', 'tipo de disco|tipo disco|tecnologia disco|tecnologia del disco'),
    ('extra.ram_ranuras_total', 'ranuras ram|ranuras de ram|slots ram|ranuras|slots|ranuras de memoria'),
    ('extra.ram_ranuras_libres', 'ranuras libres|slots libres|ranuras ram libres'),
    ('extra.ram_maxima', 'ram maxima|memoria maxima|ram max'),
    ('extra.tarjeta_video', 'tarjeta de video|tarjeta video|grafica|tarjeta grafica|gpu|video'),
    ('extra.board', 'board|placa base|tarjeta madre|motherboard|mainboard'),
    ('extra.fuente_poder', 'fuente|fuente de poder|fuente poder'),
    ('extra.licencia_windows', 'licencia|licencia windows|licencia de windows|windows licenciado|licenciado|'
                               'licencia so'),
    ('extra.cargador', 'cargador|adaptador'),
    ('extra.bateria_salud', 'salud bateria|salud de bateria|salud de la bateria|estado bateria|estado de bateria|'
                            'bateria'),
    ('extra.bateria_desgaste', 'desgaste bateria|desgaste de bateria|desgaste de la bateria'),
    ('extra.operador', 'operador|operador movil|compania celular'),
    ('extra.version_sistema', 'version|version del sistema|version android|version ios|version software|'
                              'version de android'),
    ('extra.capacidad_va', 'va|capacidad va|potencia va|capacidad en va'),
    ('extra.capacidad_w', 'w|watts|vatios|potencia w|potencia watts|potencia|capacidad w'),
    ('extra.topologia', 'topologia|tipo de ups|tipo ups'),
    ('extra.autonomia_min', 'autonomia|autonomia min|autonomia minutos|tiempo de respaldo|respaldo'),
    ('extra.baterias_cantidad', 'baterias|numero de baterias|cantidad de baterias|no baterias|cant baterias|'
                                'no de baterias'),
    ('extra.baterias_tipo', 'tipo de bateria|tipo bateria|tipo de baterias'),
    ('extra.baterias_voltaje', 'voltaje bateria|voltaje baterias|voltaje de baterias|baterias voltaje|'
                               'referencia bateria|ref bateria|referencia baterias'),
    ('extra.baterias_ultimo_cambio', 'ultimo cambio de baterias|cambio de baterias|fecha cambio baterias|'
                                     'fecha cambio de baterias|ultimo cambio baterias|fecha de cambio de baterias'),
    ('extra.tipo_panel', 'tipo de panel|panel|tipo panel'),
    ('extra.resolucion', 'resolucion'),
    ('extra.smart', 'smart|smart tv'),
    ('extra.tecnologia', 'tecnologia|tipo de impresion|tipo impresora|tecnologia de impresion'),
    ('extra.multifuncional', 'multifuncional'),
    ('extra.conexion', 'conexion|conectividad|interfaz'),
    ('extra.contador_paginas', 'contador|contador de paginas|paginas impresas|contador paginas|'
                               'contador de impresiones'),
    ('extra.direccion_ip', 'ip|direccion ip|dir ip'),
    ('extra.firmware', 'firmware'),
    ('extra.entradas', 'entradas|puertos'),
    ('cliente_nombre', 'cliente|nombre|nombre cliente|nombre del cliente|nombres|nombre completo|propietario|dueno|'
                       'razon social|nombre o razon social|cliente nombre|nombres y apellidos|tercero|'
                       'nombre propietario'),
    ('cliente_empresa', 'empresa|entidad|compania|institucion|organizacion|empresa cliente'),
    ('cliente_telefono', 'telefono|tel|telefonos|celular|cel|movil|telefono cliente|celular cliente|'
                         'numero de contacto|telefono de contacto|telefono contacto|contacto telefono|'
                         'numero celular|tel celular|telefono celular|celular contacto'),
    ('cliente_whatsapp', 'whatsapp|wsp|wpp|whatsapp cliente|numero whatsapp'),
    ('cliente_email', 'correo|email|e mail|mail|correo electronico|correo cliente|email cliente|correo del cliente'),
    ('cliente_ciudad', 'ciudad|municipio|ciudad cliente'),
    ('cliente_direccion', 'direccion|dir|direccion cliente|direccion del cliente'),
    ('cliente_documento', 'cedula|cc|nit|documento|identificacion|nit cc|cc nit|no documento|numero de documento|'
                          'cedula nit|nit cedula|documento cliente'),
    ('contacto', 'contacto|datos de contacto'),
    ('numero', 'numero|no|n'),
    ('mant_fecha', 'fecha mantenimiento|fecha de mantenimiento|fecha del mantenimiento|ultimo mantenimiento|'
                   'fecha ultimo mantenimiento|fecha del ultimo mantenimiento|fecha servicio|fecha de servicio|'
                   'fecha del servicio|ultimo servicio|fecha revision|fecha de revision|fecha visita|'
                   'fecha de visita|fecha ingreso|fecha de ingreso|fecha recibido|fecha de recibido|'
                   'fecha reparacion|fecha de reparacion|fecha entrega|fecha de entrega|fecha mtto|'
                   'fecha ultimo mtto|ultimo mtto|fecha atencion|fecha de atencion'),
    ('fecha', 'fecha|dia'),
    ('mantenimiento', 'mantenimiento|mtto|mantenimientos'),
    ('mant_tipo', 'tipo de mantenimiento|tipo mantenimiento|tipo de servicio|tipo servicio|clase de mantenimiento|'
                  'preventivo correctivo|tipo mtto|tipo de mtto|tipo de trabajo|clase de servicio'),
    ('mant_descripcion', 'trabajo realizado|trabajos realizados|descripcion|descripcion del trabajo|'
                         'descripcion del servicio|actividad|actividades|actividades realizadas|servicio realizado|'
                         'mantenimiento realizado|solucion|reparacion|reparacion realizada|trabajo|detalle|'
                         'detalle del servicio|servicio|que se hizo|labor realizada|labores|procedimiento|'
                         'accion realizada|acciones'),
    ('mant_falla', 'falla|dano|problema|falla reportada|diagnostico|sintoma|motivo|motivo de ingreso|'
                   'falla presentada|danos|fallas'),
    ('mant_costo', 'valor|costo|precio|total|valor servicio|valor del servicio|costo servicio|costo del servicio|'
                   'valor total|cobro|valor cobrado|precio servicio|valor reparacion|costo reparacion'),
    ('mant_tecnico', 'tecnico|tecnico responsable|realizado por|atendido por|responsable del mantenimiento|'
                     'ingeniero|tecnico asignado'),
    ('mant_proximo', 'proximo mantenimiento|proximo|siguiente mantenimiento|fecha proximo mantenimiento|'
                     'proxima revision|proximo servicio|fecha proxima|proxima fecha|prox mantenimiento|prox mtto|'
                     'proximo mtto|fecha proximo mtto|proxima visita|fecha de proximo mantenimiento|'
                     'siguiente revision'),
    ('mant_frecuencia', 'frecuencia|periodicidad|frecuencia de mantenimiento|frecuencia mantenimiento|cada cuanto|'
                        'cada cuantos meses|frecuencia mtto|periodicidad mantenimiento'),
    ('notas', 'observaciones|observacion|obs|notas|nota|comentarios|comentario|estado|estado fisico|'
              'estado del equipo|accesorios|ubicacion|area|dependencia|oficina|sede|piso|usuario|usuario asignado|'
              'responsable|asignado a|asignado|funcionario|empleado|cargo|codigo|codigo inventario|'
              'codigo de inventario|placa|placa inventario|activo|activo fijo|no inventario|numero de inventario|'
              'inventario|nombre del equipo|nombre equipo|hostname|nombre de host|nombre de red|garantia|'
              'fecha de compra|fecha compra|fecha de adquisicion|proveedor|factura|antivirus|office|software|'
              'software instalado|programas|valor compra|valor del equipo|costo equipo|precio compra|monitor|'
              'teclado|mouse|perifericos|mac|direccion mac|usuario windows|orden|no orden|numero de orden|'
              'orden de servicio|os|area correspondiente|area asignada|area responsable|numeracion ultimo mantenimiento|'
              'numeracion|numero de mantenimientos|cantidad de mantenimientos'),
    ('ignorar', 'item|consecutivo|id|fila|registro|nro|num'),
)
# Para la IA y la pantalla: qué es cada destino.
DESTINOS = {
    'tipo': 'Tipo de equipo', 'descripcion': 'Equipo (texto con tipo, marca y modelo)', 'marca': 'Marca',
    'modelo': 'Modelo', 'serial': 'Serial', 'imei': 'IMEI', 'color': 'Color',
    'sistema_operativo': 'Sistema operativo', 'procesador': 'Procesador', 'ram': 'Memoria RAM',
    'almacenamiento': 'Almacenamiento / disco', 'pantalla': 'Pantalla (pulgadas)',
    'especificaciones': 'Características en texto libre (procesador, RAM, disco…)',
    'extra.tarjeta_video': 'Tarjeta de video', 'extra.tipo_disco': 'Tipo de disco (HDD, SSD…)',
    'extra.bateria_salud': 'Salud de la batería (%)', 'extra.capacidad_va': 'Capacidad de la UPS (VA)',
    'extra.version_sistema': 'Versión del sistema',
    'cliente_nombre': 'Nombre del cliente', 'cliente_empresa': 'Empresa del cliente',
    'cliente_telefono': 'Teléfono o celular del cliente', 'cliente_email': 'Correo del cliente',
    'cliente_ciudad': 'Ciudad', 'cliente_direccion': 'Dirección', 'cliente_documento': 'Cédula o NIT',
    'mant_fecha': 'Fecha del mantenimiento o servicio', 'mant_tipo': 'Tipo de mantenimiento (preventivo/correctivo)',
    'mant_descripcion': 'Qué se le hizo al equipo', 'mant_falla': 'Falla o problema', 'mant_costo': 'Valor cobrado',
    'mant_tecnico': 'Técnico', 'mant_proximo': 'Próximo mantenimiento',
    'mant_frecuencia': 'Cada cuánto se hace el mantenimiento', 'notas': 'Otro dato (va a las notas del equipo)',
    'ficha_codigo': 'Número de la ficha técnica del taller',
    'ignorar': 'No sirve (números de fila, vacía)',
}
_ESPECIALES = {'contacto', 'numero', 'fecha', 'mantenimiento'}
_PERIFERICOS = {'monitor', 'teclado', 'mouse', 'raton', 'parlante', 'parlantes', 'diadema', 'webcam', 'camara',
                'cargador', 'adaptador', 'regulador'}
_IDENTIFICA = {'serial', 'serie', 'marca', 'modelo', 'referencia', 'ref', 'placa', 'codigo', 'sn'}
_SECRETOS = {'clave', 'claves', 'contrasena', 'contrasenas', 'password', 'pass', 'pin', 'puk', 'patron',
             'desbloqueo'}


# Palabras sueltas que dicen lo mismo aunque vayan dentro de un título más
# largo («Serial del portátil», «Capacidad del disco»). Las genéricas («equipo»,
# «nombre», «tipo», «valor»…) solo cuentan si son todo el título: «Resp. del
# equipo» no es la descripción del equipo.
_PALABRAS_ESPECIFICAS = {
    'serial', 'serie', 'imei', 'marca', 'modelo', 'procesador', 'ram', 'cpu', 'ssd', 'hdd', 'disco',
    'almacenamiento', 'pantalla', 'pulgadas', 'correo', 'email', 'telefono', 'celular', 'whatsapp', 'ciudad',
    'direccion', 'cedula', 'nit', 'observaciones', 'observacion', 'notas', 'comentarios', 'garantia', 'proveedor',
    'factura', 'ubicacion', 'dependencia', 'oficina', 'antivirus', 'office', 'hostname', 'frecuencia',
    'periodicidad', 'propietario', 'tecnico', 'color', 'resolucion', 'topologia', 'firmware', 'caracteristicas',
    'especificaciones', 'configuracion'}
# Lo que se ignora para comparar un título: formatos y unidades («Fecha
# (dd/mm/aaaa)», «RAM (GB)», «Valor ($)»).
_RUIDO = {'dd', 'mm', 'aa', 'aaaa', 'yyyy', 'gb', 'mb', 'tb', 'ghz', 'mhz', 'cop', 'pesos', 'opcional',
          'obligatorio', 'x'}


def _indices():
    exactos, contenidos = {}, []
    for destino, lista in _SINONIMOS:
        for s in lista.split('|'):
            s = _norm(s)
            if not s:
                continue
            exactos.setdefault(s, destino)
            if ' ' in s or s in _PALABRAS_ESPECIFICAS:
                contenidos.append((s, destino))
    contenidos.sort(key=lambda x: -len(x[0]))
    return exactos, contenidos


_EXACTOS, _CONTENIDOS = _indices()


def destino_de(titulo):
    """Campo al que va una columna según su título, o None si no se reconoce."""
    h = _norm(titulo)
    if not h:
        return None
    palabras = set(h.split())
    if palabras & _SECRETOS:
        return 'secreto'
    if h in _EXACTOS:
        return _EXACTOS[h]
    if palabras & _PERIFERICOS and palabras & _IDENTIFICA:
        return 'notas'                       # «Serial monitor»: es de otro aparato
    sin_ruido = ' '.join(w for w in h.split() if w not in _RUIDO)
    if sin_ruido in _EXACTOS:
        return _EXACTOS[sin_ruido]
    for s, destino in _CONTENIDOS:
        if re.search(r'(?<![a-z0-9])' + re.escape(s) + r'(?![a-z0-9])', h):
            return destino
    return None


def _consecutivo(valores):
    nums = [_numero(v) for v in valores]
    nums = [n for n in nums if n is not None]
    if len(nums) < 3 or len(nums) < 0.8 * len(valores):
        return False
    return sum(1 for a, b in zip(nums, nums[1:]) if b - a == 1) >= 0.8 * (len(nums) - 1)


def _resolver(destino, valores):
    """Ajusta el destino por lo que trae la columna («Contacto» con celulares
    es teléfono; «Mantenimiento» con fechas es la fecha del servicio)."""
    vals = [v for v in valores if _celda(v)][:60]
    if not vals or destino not in _ESPECIALES | {'cliente_nombre', 'cliente_telefono', 'mant_fecha',
                                                 'mant_proximo', 'tipo'}:
        return destino

    def parte(pred):
        return sum(1 for v in vals if pred(v)) / len(vals)
    if destino == 'tipo':
        servicio = parte(lambda v: not tipos.tipo_desde_texto(_celda(v)) and len(_celda(v)) <= 30
                         and tipo_mantenimiento(_celda(v)) is not None)
        return 'mant_tipo' if servicio >= .6 else destino
    if destino in ('contacto', 'numero', 'cliente_nombre', 'cliente_telefono'):
        if parte(lambda v: _correo(v) is not None) >= .6:
            return 'cliente_email'
        if parte(lambda v: _telefono(v) is not None) >= .6:
            return 'cliente_telefono'
        if destino == 'numero':
            return 'ignorar' if _consecutivo(vals) else 'notas'
        return 'cliente_telefono' if destino == 'cliente_telefono' else 'cliente_nombre'
    if parte(lambda v: _fecha_valor(v) is not None) >= .6:
        return 'mant_fecha' if destino == 'mantenimiento' else destino
    if destino == 'mantenimiento':
        cortos = parte(lambda v: tipo_mantenimiento(_celda(v)) is not None and len(_celda(v)) <= 30)
        return 'mant_tipo' if cortos >= .6 else 'mant_descripcion'
    return 'notas'


# ── Marcas, modelos y características en texto libre ────────────
MARCAS = ('Hewlett Packard', 'HP', 'Dell', 'Lenovo', 'Asus', 'Acer', 'Apple', 'Samsung', 'Xiaomi', 'Redmi',
          'Motorola', 'Huawei', 'Honor', 'Oppo', 'Realme', 'Vivo', 'Nokia', 'LG', 'Sony', 'Toshiba', 'MSI',
          'Gigabyte', 'Compaq', 'Janus', 'Kalley', 'Challenger', 'Hyundai', 'Epson', 'Canon', 'Brother', 'Ricoh',
          'Kyocera', 'Xerox', 'Lexmark', 'APC', 'CDP', 'Forza', 'Tripp Lite', 'Eaton', 'Powest', 'Unitec',
          'Microsoft', 'Nintendo', 'TP-Link', 'Mikrotik', 'Ubiquiti', 'Hikvision', 'Dahua', 'Tecno', 'Infinix',
          'ZTE', 'Alcatel', 'Google', 'OnePlus', 'Panasonic', 'Philips', 'TCL', 'Hisense', 'Caixun', 'Olimpo',
          'Zebra', 'Bixolon', 'Gateway', 'Lanix', 'Haier', 'Compumax', 'Sharp', 'Fujitsu', 'Razer', 'Alienware',
          'Packard Bell', 'Benq', 'ViewSonic', 'AOC', 'Kodak', 'Lenco', 'Imaco', 'Polaroid', 'Matica', 'Evolis',
          'Datacard', 'Fargo', 'Zkteco')
_MARCA_CANONICA = {'hewlett packard': 'HP', 'asustek': 'Asus'}


def _marca_canonica(texto):
    """La marca como la conoce el sistema, aunque el fabricante la escriba con su
    razón social: «HP Inc.» → «HP», «Dell Inc» → «Dell», «Matica Technologies AG»
    → «Matica». Una marca que no se conoce queda como viene (con mayúscula inicial)."""
    texto = ' '.join(str(texto or '').split())
    n = _norm(texto)
    if not n:
        return None
    for clave, marca in _MARCA_CANONICA.items():
        if n == clave or n.startswith(clave + ' '):
            return marca
    for marca in sorted(MARCAS, key=len, reverse=True):
        m = _norm(marca)
        if n == m or n.startswith(m + ' '):
            return _MARCA_CANONICA.get(m, marca)
    return texto.title()


_MODELOS_TIPO = (
    ('tablet', re.compile(r'\b(ipad|galaxy\s+tab|tab\s+[as]\d|mediapad|matepad|lenovo\s+tab|fire\s+hd)', re.I)),
    ('celular', re.compile(r'\b(iphone|galaxy\s+[asmz]\d|galaxy\s+note|redmi|poco\s|moto\s?[eg]\d|pixel\s\d|'
                           r'huawei\s+[py]\d|honor\s+\d|oppo\s+a\d|realme|tecno\s|infinix|zte\s+blade|alcatel)', re.I)),
    ('computador', re.compile(r'\b(optiplex|prodesk|elitedesk|thinkcentre|ideacentre|ideacenter|veriton|imac|'
                              r'all[\s-]?in[\s-]?one|aio\b|'
                              r'precision\s+t)', re.I)),
    ('portatil', re.compile(r'\b(macbook|thinkpad|ideapad|yoga|legion|pavilion|probook|elitebook|envy|spectre|omen|'
                            r'victus|inspiron|latitude|xps|vivobook|zenbook|tuf\s|rog\s|aspire|swift|nitro|predator|'
                            r'travelmate|satellite|tecra|galaxy\s+book|surface\s+laptop|chromebook|\d{3}\s?g\d\b)', re.I)),
    ('impresora', re.compile(r'\b(ecotank|laserjet|deskjet|officejet|smart\s+tank|ink\s+tank|pixma|maxify|'
                             r'imageclass|ecosys|epson\s+l\d|dcp-|hl-|mfc-)', re.I)),
    ('ups', re.compile(r'\b(back-?ups|smart-?ups|\d{3,4}\s?va\b)', re.I)),
    ('consola', re.compile(r'\b(playstation|ps[345]\b|xbox|nintendo)', re.I)),
)
_SO = re.compile(r'\b(windows\s*(?:xp|vista|7|8\.1|8|10|11)(?:\s*(?:pro(?:fessional)?|home|hogar|enterprise|'
                 r'education|single language))?|win\s?(?:7|8|10|11)(?:\s*pro)?|ubuntu(?:\s*[\d.]+)?|linux(?:\s+\w+)?|'
                 r'mac\s?os(?:\s+[\w.]+)?|macos(?:\s+[\w.]+)?|android\s*\d+(?:\.\d+)?|ios\s*\d+(?:\.\d+)?|'
                 r'chrome\s?os)\b', re.I)
_CPU = re.compile(r'\b(core\s+ultra\s*\d\s*\d{3}\w*|(?:intel\s+)?(?:core\s*)?i[3579](?:[\s-]?\d{3,5}[a-z]{0,3})?|'
                  r'ryzen\s*[3579](?:\s*(?:pro\s*)?\d{3,5}[a-z]{0,3})?|celeron(?:\s*[a-z]?\d{3,5}\w*)?|'
                  r'pentium(?:\s*(?:gold|silver)?\s*[a-z]?\d{3,5}\w*)?|athlon(?:\s*\w+)?|apple\s+m\d(?:\s+(?:pro|max|ultra))?|'
                  r'snapdragon\s*\w+|mediatek\s*\w+|helio\s*\w+|dimensity\s*\d+|exynos\s*\d+|xeon(?:\s*\w+)?)\b', re.I)


def _marca_modelo(texto):
    """(marca, modelo) de un texto como «Portátil HP 240 G8 Core i5»."""
    t = ' '.join(str(texto or '').split())
    for marca in sorted(MARCAS, key=len, reverse=True):
        m = re.search(r'(?<![A-Za-z0-9])' + re.escape(marca) + r'(?![A-Za-z0-9])', t, re.I)
        if m:
            resto = re.split(r'[,;(|]| - | / ', t[m.end():])[0]
            resto = re.sub(r'(?i)\b(core|intel|amd|ryzen|celeron|pentium|ram|ssd|hdd|windows|win|\d+\s*gb|\d+\s*tb)\b.*$',
                           '', resto).strip(' -:')
            return _MARCA_CANONICA.get(marca.lower(), marca), (resto[:120] or None)
    return None, None


def _gb_tb(numero, unidad):
    n = _numero(numero)
    if n is None:
        return None
    unidad = (unidad or '').upper().rstrip('B')
    if unidad not in ('G', 'T'):
        unidad = 'T' if n < 16 else 'G'
    return f'{n:g} {unidad}B'


def specs_de_texto(texto):
    """Procesador, RAM, disco y sistema de un texto libre («Core i5 8GB 256 SSD
    Win 10»). Devuelve (columnas, extras); solo lo que se reconoce."""
    t = ' ' + ' '.join(str(texto or '').split()) + ' '
    col, ext = {}, {}
    cpu = _CPU.search(t)
    if cpu:
        valor = cpu.group(1).strip()
        m = re.fullmatch(r'(?i)(?:intel\s+)?(?:core\s*)?(i[3579])(?:[\s-]?(\d{3,5}[a-z]{0,3}))?', valor)
        if m:
            valor = f'Core {m.group(1).lower()}' + (f'-{m.group(2).upper()}' if m.group(2) else '')
        col['procesador'] = valor[:160]
    candidatos = []
    # (un número pegado a letras, como el «1» de «PM9B1 NVMe», es parte de una referencia, no la capacidad)
    for patron in (r'(?i)(?<![a-z0-9.,])(\d{1,4}(?:[.,]\d)?)\s*(gb|tb|g|t)?\s*(?:de\s+)?(ssd|hdd|nvme|m\.2|emmc|disco(?:\s+duro|\s+s[oó]lido)?)\b',
                   r'(?i)\b(ssd|hdd|nvme|m\.2|emmc|disco(?:\s+duro|\s+s[oó]lido)?)\s*(?:de\s+)?[:\-]?\s*(\d{1,4}(?:[.,]\d)?)\s*(gb|tb|g|t)?\b'):
        for disco in re.finditer(patron, t):
            if disco.group(1)[0].isdigit():
                cap, unidad, clase = disco.group(1), disco.group(2), disco.group(3)
            else:
                clase, cap, unidad = disco.group(1), disco.group(2), disco.group(3)
            capacidad = _gb_tb(cap, unidad)
            clase_n = _norm(clase)
            tipo_disco = ('SSD NVMe' if clase_n in ('nvme', 'm 2') else 'SSD SATA' if clase_n in ('ssd', 'disco solido')
                          else 'eMMC' if clase_n == 'emmc' else 'HDD')
            n = _numero(cap) or 0
            en_gb = n * 1024 if (capacidad or '').endswith('TB') else n
            candidatos.append((en_gb, capacidad, tipo_disco))
    if candidatos:
        grandes = [c for c in candidatos if c[0] >= 16]
        _en_gb, capacidad, tipo_disco = max(grandes) if grandes else candidatos[0]
        if capacidad:
            col['almacenamiento'] = f"{'SSD' if tipo_disco.startswith('SSD') else tipo_disco} {capacidad}"
        ext['tipo_disco'] = tipo_disco
    ram = (re.search(r'(?i)(\d{1,3})\s*(?:gb|g)\b\s*(?:de\s+)?(?:ram|memoria|ddr\d\w?)', t)
           or re.search(r'(?i)\b(?:ram|memoria)\s*(?:de\s*)?[:\-]?\s*(\d{1,3})\s*(?:gb|g)\b', t))
    if ram and 1 <= int(ram.group(1)) <= 256:
        ddr = re.search(r'(?i)\b(lp)?ddr\d\w?\b', t)
        col['ram'] = f'{int(ram.group(1))} GB' + (f' {ddr.group(0).upper()}' if ddr else '')
    if 'ram' not in col or 'almacenamiento' not in col:
        tamanos = [(_numero(m.group(1)), m.group(2).upper()) for m in re.finditer(r'(?i)(\d{1,4})\s*(gb|tb)\b', t)]
        tamanos = [(n, u) for n, u in tamanos if n]
        if len(tamanos) >= 2:
            en_gb = sorted(tamanos, key=lambda x: x[0] * (1024 if x[1] == 'TB' else 1))
            if 'ram' not in col and en_gb[0][1] == 'GB' and en_gb[0][0] <= 64:
                col['ram'] = f'{en_gb[0][0]:g} GB'
            if 'almacenamiento' not in col and (en_gb[-1][1] == 'TB' or en_gb[-1][0] >= 64):
                col['almacenamiento'] = f'{en_gb[-1][0]:g} {en_gb[-1][1]}'
    so = _SO.search(t)
    if so:
        valor = re.sub(r'(?i)^win(?!dows)\s?', 'Windows ', so.group(1)).strip()
        col['sistema_operativo'] = valor[:1].upper() + valor[1:120]
    return col, ext


_CAPACIDAD = re.compile(r'(?i)(?<![a-z0-9.,])(\d{1,4}(?:[.,]\d{1,2})?)\s*(gb|tb)\b')
_CLASE_DISCO = re.compile(r'(?i)(?<![a-z0-9])(nvme|m\.2|ssd|hdd|emmc|s[oó]lido|mec[aá]nico)(?![a-z0-9])')


def discos_de_texto(texto):
    """[(tipo, capacidad)] de cada disco de la fila «Unidad de Disco» de una ficha,
    p. ej. «SSD 239GB M.2 / HDD 932GB» → [('SSD SATA', '239 GB'), ('HDD', '932 GB')].
    El tipo es None si no lo dice; un número pegado a letras («PM9B1») es parte de
    la referencia del disco, no su capacidad."""
    discos = []
    for trozo in re.split(r'\s*/\s*|\s*\+\s*|\s*;\s*|\s+y\s+|\s*\|\s*', str(texto or '')):
        capacidad = _CAPACIDAD.search(trozo)
        if not capacidad:
            continue
        clases = {_norm(m.group(1)) for m in _CLASE_DISCO.finditer(trozo)}
        tipo = ('SSD NVMe' if 'nvme' in clases else 'SSD SATA' if clases & {'ssd', 'solido'}
                else 'SSD NVMe' if 'm 2' in clases else 'HDD' if clases & {'hdd', 'mecanico'}
                else 'eMMC' if 'emmc' in clases else None)
        discos.append((tipo, _gb_tb(capacidad.group(1), capacidad.group(2))))
    return discos[:4]


def _cpu_o_gpu(texto):
    """'cpu', 'gpu' o None según lo que describe el texto."""
    cpu = re.search(r'(?i)\b(core|ryzen|razen|celeron|pentium|athlon|xeon|i[3579][\s-]?\d{3,5}|i\d-\d{4,5}|'
                    r'\d{1,2}(?:th|st|nd|rd)\s+gen|apple\s+m\d|snapdragon|mediatek|exynos)\b', texto or '')
    gpu = re.search(r'(?i)\b(graphi\w*|gr[aá]fic\w*|radeon|geforce|gtx|rtx|uhd|iris|vega|quadro)\b', texto or '')
    if cpu and not gpu:
        return 'cpu'
    if gpu and not cpu:
        return 'gpu'
    return None


def _descruzar_cpu_gpu(datos):
    """Algunas fichas traen el procesador y la tarjeta gráfica cruzados
    («Procesador: Intel UHD Graphics 770» y «Tarjeta gráfica: Core i7-13700»):
    los pone en su sitio. Devuelve True si los cambió."""
    grafica = next((k for k in datos if _norm(k) in ('tarjeta grafica', 'tarjeta de video', 'grafica', 'gpu')), None)
    if not grafica or not datos.get('Procesador') or not datos.get(grafica):
        return False
    if _cpu_o_gpu(datos['Procesador']) == 'gpu' and _cpu_o_gpu(datos[grafica]) == 'cpu':
        datos['Procesador'], datos[grafica] = datos[grafica], datos['Procesador']
        return True
    return False


def _tipo_por_datos(equipo):
    """Tipo deducido de los datos del equipo, o None."""
    extras = equipo.get('extras') or {}
    texto = ' '.join(str(equipo.get(c) or '') for c in ('marca', 'modelo'))
    for codigo, patron in _MODELOS_TIPO:
        if patron.search(texto):
            return codigo
    if equipo.get('imei') or extras.get('imei2'):
        return 'celular'
    if any(extras.get(k) for k in ('capacidad_va', 'capacidad_w', 'baterias_cantidad', 'autonomia_min')):
        return 'ups'
    if equipo.get('procesador') or equipo.get('ram'):
        return 'portatil' if (equipo.get('pantalla') or extras.get('bateria_desgaste') or extras.get('cargador')) \
            else 'computador'
    return None


# ── Una fila del lote ───────────────────────────────────────────
def _fila_vacia():
    return {'clase': 'equipo', 'incluir': True, 'editada': False, 'cliente': {}, 'usar_contacto': None,
            'equipo': {'tipo': None, 'extras': {}}, 'tipo_adivinado': False, 'historial': [],
            'plan': {'cada_meses': None, 'proximo': None}, 'notas': [], 'avisos': [], 'fuente': 'reglas',
            'resumen': None, 'documento_id': None, 'unir_con': None, 'adjuntar_a': None}


def _agregar_nota(fila, texto):
    texto = ' '.join(str(texto or '').split())[:300]
    if texto and texto not in fila['notas']:
        fila['notas'].append(texto)


def _poner_imei(fila, valor, etiqueta='IMEI'):
    crudo = _celda(valor)
    numeros = re.findall(r'\d[\d\s-]{13,20}\d', crudo) or [crudo]
    for numero in numeros:
        digitos = re.sub(r'\D', '', numero)
        if lector.imei_luhn(digitos):
            if not fila['equipo'].get('imei'):
                fila['equipo']['imei'] = digitos
            elif digitos != fila['equipo']['imei'] and not fila['equipo']['extras'].get('imei2'):
                fila['equipo']['extras']['imei2'] = digitos
        elif digitos:
            _agregar_nota(fila, f'{etiqueta} (revisar, no es válido): {numero.strip()}')
            fila['avisos'].append(f'El IMEI «{numero.strip()}» no es válido: quedó en las notas.')


_SIN_SERIAL = {'na', 'n a', 'no', 'no tiene', 'sin', 'sin serial', 'sin serie', 'ninguno', 'no aplica', 'x', 'xx',
               'xxx', 's n', 'sn', '0', '00', '000', 'pendiente', 'no registra', 'nr', 'ilegible', 'no legible',
               'desconocido', 'no visible', 'n d', 'nd'}


def _serial(v):
    s = _celda(v).strip(' .:-')
    if _norm(s) in _SIN_SERIAL or len(re.sub(r'[\W_]', '', s)) < 3:
        return None
    return s[:120]


def _cerrar(fila):
    """Tipo, largos, extras válidos para el tipo y notas: lo que no cabe en la
    ficha queda en las notas (ningún dato se pierde)."""
    e = fila['equipo']
    e['extras'] = {k: v for k, v in (e.get('extras') or {}).items() if v not in (None, '')}
    if not tipos.es_valido(e.get('tipo') or ''):
        deducido = _tipo_por_datos(e)
        e['tipo'] = deducido or 'otro'
        fila['tipo_adivinado'] = True
    tipo = e['tipo']
    if tipo in ('televisor', 'monitor') and e.get('pantalla') and not e['extras'].get('pulgadas'):
        numero = re.search(r'\d{2,3}(?:[.,]\d)?', str(e['pantalla']))
        if numero:
            e['extras']['pulgadas'] = numero.group(0).replace(',', '.')
    for c in ('ram', 'almacenamiento'):
        valor = str(e.get(c) or '').strip()
        if re.fullmatch(r'\d{1,4}(?:[.,]\d+)?', valor):
            n = _numero(valor)
            e[c] = (f'{n:g} GB' if c == 'ram' or n >= 16 else f'{n:g} TB') if n else None
    if re.fullmatch(r'\d{1,2}(?:[.,]\d)?', str(e.get('pantalla') or '').strip()):
        e['pantalla'] = str(e['pantalla']).strip().replace('.', ',') + '"'
    # Extras del tipo: los que no aplican o no tienen la forma esperada, a las notas.
    definidos = {c[0]: c for c in tipos.campos_extra(tipo)}
    for clave, valor in list(e['extras'].items()):
        campo = definidos.get(clave)
        valor_txt = _celda(valor)
        if campo and campo[2] == 'numero':
            n = re.sub(r'[%\s]', '', valor_txt).replace(',', '.')
            if re.fullmatch(r'\d+(?:\.\d+)?', n):
                e['extras'][clave] = n
                continue
        elif campo and campo[2] == 'fecha':
            f = _fecha_valor(valor)
            if f:
                e['extras'][clave] = f.isoformat()
                continue
        elif campo and campo[2] == 'opcion':
            opcion = next((o for o in campo[3] if _norm(o) == _norm(valor_txt)), None)
            if opcion:
                e['extras'][clave] = opcion
                continue
        elif campo:
            e['extras'][clave] = valor_txt[:200]
            continue
        etiqueta = campo[1] if campo else next((c[1] for codigo in tipos.CODIGOS for c in tipos.campos_extra(codigo)
                                               if c[0] == clave), clave)
        _agregar_nota(fila, f'{etiqueta}: {valor_txt}')
        e['extras'].pop(clave)
    e['extras'] = tipos.limpiar_extras(tipo, e['extras'])
    for c in st._COLUMNAS_EQUIPO:
        if e.get(c):
            e[c] = ' '.join(str(e[c]).split())[:st._LARGOS[c]]
        else:
            e.pop(c, None)
    if e.get('imei') and not lector.imei_luhn(e['imei']):
        _agregar_nota(fila, f"IMEI (revisar, no es válido): {e.pop('imei')}")
    fila['historial'] = _sin_repetidos(fila['historial'])[:HISTORIAL_MAX]
    return fila


def _sin_repetidos(entradas):
    vistos, salida = set(), []
    for h in sorted(entradas, key=lambda x: x['fecha']):
        clave = (h['fecha'], h['tipo'], _norm(h['descripcion'])[:40])
        if clave not in vistos:
            vistos.add(clave)
            salida.append(h)
    return salida


def tiene_datos(equipo):
    return any(equipo.get(c) for c in ('marca', 'modelo', 'serial', 'imei', 'procesador', 'sistema_operativo',
                                       'almacenamiento', 'ram')) or bool(equipo.get('extras'))


def tiene_datos_fila(fila):
    """¿La fila describe un equipo? (datos, o al menos el tipo dicho en el archivo)."""
    e = fila['equipo']
    return tiene_datos(e) or (e.get('tipo') not in (None, 'otro') and not fila.get('tipo_adivinado'))


def claves_de(equipo):
    """Lo que identifica al equipo: IMEI y serial normalizados (sin espacios,
    puntos ni guiones). Conjunto vacío si no tiene ninguno."""
    claves = set()
    if equipo.get('imei'):
        claves.add('imei:' + re.sub(r'\D', '', str(equipo['imei'])))
    serial = re.sub(r'[\s\-./]', '', str(equipo.get('serial') or '')).lower()
    if len(serial) >= 4:
        claves.add('serial:' + serial)
    return claves


def _unir(destino, origen):
    """Une en `destino` una fila del mismo equipo (mismo serial o IMEI)."""
    for c, v in origen['equipo'].items():
        if c == 'extras':
            for k, x in (v or {}).items():
                destino['equipo']['extras'].setdefault(k, x)
        elif v and not destino['equipo'].get(c):
            destino['equipo'][c] = v
    for k, v in origen['cliente'].items():
        if v and not destino['cliente'].get(k):
            destino['cliente'][k] = v
    destino['historial'] = _sin_repetidos(destino['historial'] + origen['historial'])[:HISTORIAL_MAX]
    for k in ('cada_meses', 'proximo'):
        if origen['plan'].get(k):
            if k == 'proximo' and destino['plan'].get(k):
                destino['plan'][k] = max(destino['plan'][k], origen['plan'][k])
            else:
                destino['plan'][k] = destino['plan'].get(k) or origen['plan'][k]
    for n in origen['notas']:
        _agregar_nota(destino, n)
    for a in origen['avisos']:
        if a not in destino['avisos']:
            destino['avisos'].append(a)


# ── Tablas (Excel y CSV) ────────────────────────────────────────
_UTILES_EQUIPO = {'tipo', 'descripcion', 'marca', 'modelo', 'serial', 'imei', 'procesador', 'especificaciones'}
_SE_JUNTAN = {'notas', 'mant_descripcion', 'mant_falla', 'especificaciones', 'descripcion'}
_HOJAS_DE_AYUDA = re.compile(r'^(instrucciones?|ejemplos?|ayuda|listas|parametros|configuracion|'
                             r'(formato|plantilla|modelo)s?( de ficha| en blanco| vacio| vacia)?|en blanco|blanco)$')


def _es_etiqueta(v):
    """¿Es el título de un dato («Marca», «Serial:», «Memoria RAM»)? Un valor
    con números o una frase («Cambio de disco por SSD») no lo es."""
    texto = _celda(v) if isinstance(v, str) else ''
    palabras = len(texto.split())
    if not texto or len(texto) > 40 or palabras > 4:
        return False
    if texto.endswith(':') or _norm(texto) in _EXACTOS:
        return destino_de(texto) not in (None, 'ignorar') or texto.endswith(':')
    return palabras <= 3 and not re.search(r'\d', texto) and destino_de(texto) not in (None, 'ignorar')


def _pares(fila):
    """Cuántas «Etiqueta | valor» trae la fila (formato de ficha, no de tabla)."""
    celdas = [v for v in fila if _celda(v)]
    return sum(1 for k, v in enumerate(celdas[:-1]) if _es_etiqueta(v) and not _es_etiqueta(celdas[k + 1]))


_DATOS_DEL_EQUIPO = _UTILES_EQUIPO | {'sistema_operativo', 'ram', 'almacenamiento', 'color', 'pantalla'}


def _pares_del_equipo(fila):
    """Como `_pares`, pero solo los datos del equipo (marca, serial, RAM…). El
    membrete de un documento («CODIGO | …», «EMPRESA | …», «VERSION | 1») no cuenta."""
    celdas = [v for v in fila if _celda(v)]
    return sum(1 for k, v in enumerate(celdas[:-1]) if _es_etiqueta(v) and not _es_etiqueta(celdas[k + 1])
               and destino_de(v) in _DATOS_DEL_EQUIPO)


def _buscar_encabezado(filas):
    """(índice, {col: destino}) de la fila de títulos de una tabla de equipos.
    None si la hoja no es una tabla (p. ej. una hoja de vida «Etiqueta | valor»)."""
    mejor = None
    for i, fila in enumerate(filas[:25]):
        textos = {j: v for j, v in enumerate(fila[:COLUMNAS_MAX]) if isinstance(v, str) and _celda(v)
                  and len(_celda(v)) <= 60}
        if len(textos) < 2:
            continue
        destinos = {j: destino_de(v) for j, v in textos.items()}
        utiles = {d for d in destinos.values() if d and d not in ('notas', 'ignorar', 'secreto')}
        if len(utiles) < 2 or not utiles & (_UTILES_EQUIPO | {'mant_descripcion', 'mant_fecha'}):
            continue
        if not utiles & _UTILES_EQUIPO and 'descripcion' not in utiles:
            continue
        reconocidas = sum(1 for d in destinos.values() if d)
        if _pares(fila) >= max(1, reconocidas // 2) and any(_pares(f2) for f2 in filas[i + 1:i + 4]):
            continue                                   # «Marca | HP | Modelo | 240 G8»: es una ficha
        if sum(1 for f2 in filas[:i] if _pares_del_equipo(f2)) >= 3:
            continue                                   # ficha con su tabla de historial debajo
        if not any(sum(1 for j in textos if j < len(f2) and _celda(f2[j])) >= 2 for f2 in filas[i + 1:i + 8]):
            continue                                   # títulos sin datos debajo
        puntaje = (len(utiles), -i)
        if mejor is None or puntaje > mejor[0]:
            mejor = (puntaje, i, {j: v for j, v in textos.items()})
    if not mejor:
        return None
    return mejor[1], mejor[2]


def _solo_titulos(filas):
    """¿La hoja es una fila de títulos y nada más (la plantilla sin llenar)?"""
    con_datos = [f for f in filas if any(_celda(v) for v in f)]
    if len(con_datos) != 1:
        return False
    utiles = {destino_de(v) for v in con_datos[0] if isinstance(v, str)} - {None, 'notas', 'ignorar', 'secreto'}
    return len(utiles) >= 2


def _parte_de_hoja(filas, titulo):
    """Una hoja → parte 'tabla' (títulos + filas) o 'texto' (ficha), o None si está vacía."""
    filas = [list(f) for f in filas]
    while filas and not any(_celda(v) for v in filas[-1]):
        filas.pop()
    if not filas:
        return None
    if _solo_titulos(filas):
        return None
    encabezado = _buscar_encabezado(filas)
    if encabezado:
        i, titulos = encabezado
        datos = []
        for k, fila in enumerate(filas[i + 1:], start=i + 2):
            if any(_celda(v) for v in fila):
                if [_norm(fila[j]) if j < len(fila) else '' for j in titulos] == [_norm(t) for t in titulos.values()]:
                    continue                           # títulos repetidos a mitad de la hoja
                datos.append((k, fila[:COLUMNAS_MAX]))
            if len(datos) >= FILAS_LEER_MAX:
                break
        # Columnas con datos pero sin título: también se miran.
        for j in sorted({j for _, f in datos for j, v in enumerate(f) if _celda(v)} - set(titulos)):
            titulos[j] = ''
        # Lo que hay encima de los títulos («INVENTARIO DE EQUIPOS — EMPRESA X»): dice de quién son.
        encabezado = [' · '.join(_celda(v) for v in f if _celda(v)) for f in filas[:i]][-8:]
        return {'tipo': 'tabla', 'titulo': titulo, 'titulos': {j: _celda(t) for j, t in sorted(titulos.items())},
                'filas': datos, 'fila_titulos': i + 1, 'encabezado': [x for x in encabezado if x]}
    texto = _hoja_a_texto(filas)
    return {'tipo': 'texto', 'titulo': titulo, 'texto': texto, 'paginas': [texto]} if texto.strip() else None


def _hoja_a_texto(filas):
    """Hoja tipo ficha → líneas «Etiqueta: valor» que entiende el lector."""
    lineas = []
    for fila in filas:
        celdas = [v for v in fila if _celda(v)]
        if not celdas:
            continue
        if isinstance(celdas[0], (datetime, date)) or (isinstance(celdas[0], str) and fechas_en(celdas[0])
                                                        and len(_celda(celdas[0])) <= 12):
            lineas.append(' | '.join(_celda(v) for v in celdas))   # fila de un historial: tal cual
            continue
        partes, k = [], 0
        while k < len(celdas):
            v = celdas[k]
            if isinstance(v, str) and destino_de(v) == 'secreto':
                k += 2                                 # «Contraseña | 1234»: no se copia
                continue
            if k + 1 < len(celdas) and (_es_etiqueta(v) or (isinstance(v, str) and v.strip().endswith(':'))) \
                    and not _es_etiqueta(celdas[k + 1]):
                if partes:
                    lineas.append(' | '.join(partes))
                    partes = []
                lineas.append(f'{_celda(v).rstrip(":").strip()}: {_celda(celdas[k + 1])}')
                k += 2
                continue
            partes.append(_celda(v))
            k += 1
        if partes:
            lineas.append(' | '.join(partes))
    return '\n'.join(lineas)


def _columnas_de_tabla(parte, mapeo_ia=None):
    """{col: destino} de una parte tabla (reglas + lo que dijo la IA), y las
    columnas que nadie reconoció."""
    columnas, sin_mapear, usados = {}, [], set()
    for j, titulo in parte['titulos'].items():
        j = int(j)
        destino = (mapeo_ia or {}).get(j) or (mapeo_ia or {}).get(str(j)) or destino_de(titulo)
        valores = [f[j] for _, f in parte['filas'] if j < len(f)]
        if destino is None:
            if any(_celda(v) for v in valores):
                sin_mapear.append(j)
                destino = 'notas'
            else:
                destino = 'ignorar'
        destino = _resolver(destino, valores)
        if destino in usados and destino not in _SE_JUNTAN and destino not in ('ignorar', 'secreto'):
            destino = 'cliente_whatsapp' if destino == 'cliente_telefono' and 'cliente_whatsapp' not in usados \
                else 'notas'
        usados.add(destino)
        columnas[j] = destino
    return columnas, sin_mapear


def _fila_de_tabla(fila_excel, columnas, titulos):
    fila = _fila_vacia()
    e, c = fila['equipo'], fila['cliente']
    reg = {}
    for j, destino in columnas.items():
        if j < len(fila_excel) and _celda(fila_excel[j]) and destino not in ('ignorar', 'secreto'):
            reg.setdefault(destino, []).append((titulos.get(j) or f'Columna {_letra(j)}', fila_excel[j]))

    def txt(d):
        return _celda(reg[d][0][1]) if d in reg else ''

    def todos(d):
        return [_celda(v) for _, v in reg.get(d, [])]
    if 'tipo' in reg:
        e['tipo'] = tipos.tipo_desde_texto(txt('tipo'))
        if not e['tipo']:
            _agregar_nota(fila, f"{reg['tipo'][0][0]}: {txt('tipo')}")
    descripcion = ' '.join(todos('descripcion'))
    if descripcion:
        e['tipo'] = e.get('tipo') or tipos.tipo_desde_texto(descripcion)
        marca, modelo = _marca_modelo(descripcion)
        if not txt('marca') and marca:
            e['marca'] = marca
        if not txt('modelo') and modelo:
            e['modelo'] = modelo
        if not marca and not txt('modelo'):
            # «Portátil gris de la recepción»: sin la palabra del tipo, lo demás es el modelo.
            resto = _sin_palabras_de_tipo(descripcion)
            if resto:
                e['modelo'] = resto[:120]
    for col in ('marca', 'modelo', 'color', 'sistema_operativo', 'procesador', 'ram', 'almacenamiento', 'pantalla'):
        if txt(col):
            e[col] = txt(col)
    if txt('modelo') and not e.get('tipo'):
        e['tipo'] = tipos.tipo_desde_texto(txt('modelo'))     # «Matica MC310 (Impresora de carnés)»
        if not txt('marca'):
            marca, modelo = _marca_modelo(txt('modelo'))
            if marca:
                e['marca'], e['modelo'] = marca, modelo or e['modelo']
    if 'serial' in reg:
        e['serial'] = _serial(reg['serial'][0][1])
    if txt('ficha_codigo') and re.search(r'[A-Za-z]', txt('ficha_codigo')) and re.search(r'\d', txt('ficha_codigo')):
        fila['ficha_codigo'] = ' '.join(txt('ficha_codigo').split())[:40]   # une la fila con la hoja de su ficha
    if 'imei' in reg:
        _poner_imei(fila, reg['imei'][0][1])
    if 'extra.imei2' in reg:
        _poner_imei(fila, reg['extra.imei2'][0][1], 'IMEI 2')
    texto_specs = ' '.join(todos('especificaciones') + ([descripcion] if descripcion else []))
    if texto_specs:
        cols, exts = specs_de_texto(texto_specs)
        for k, v in cols.items():
            e[k] = e.get(k) or v
        for k, v in exts.items():
            e['extras'].setdefault(k, v)
        for v in todos('especificaciones'):
            _agregar_nota(fila, f'Características: {v}')
    for destino, valores in reg.items():
        if destino.startswith('extra.') and destino != 'extra.imei2':
            v = valores[0][1]
            e['extras'][destino[6:]] = _fecha_valor(v).isoformat() if isinstance(v, (datetime, date)) else _celda(v)
    # Cliente
    c['nombre'] = txt('cliente_nombre') or None
    c['empresa'] = txt('cliente_empresa') or None
    for etiqueta, v in reg.get('cliente_telefono', []) + reg.get('cliente_whatsapp', []):
        tel = _telefono(v)
        if tel and not c.get('telefono'):
            c['telefono'] = tel
        elif tel and tel != c.get('telefono') and not c.get('whatsapp'):
            c['whatsapp'] = tel
        elif _celda(v) and not tel:
            _agregar_nota(fila, f'{etiqueta}: {_celda(v)}')
    for etiqueta, v in reg.get('cliente_email', []):
        correo = _correo(v)
        if correo and not c.get('email'):
            c['email'] = correo
        elif not correo:
            _agregar_nota(fila, f'{etiqueta}: {_celda(v)}')
    c['ciudad'] = txt('cliente_ciudad') or None
    c['direccion'] = txt('cliente_direccion') or None
    c['documento'] = txt('cliente_documento') or None
    # Mantenimiento de esta fila (un historial: una fila por servicio)
    que = ' '.join(todos('mant_descripcion'))
    falla = ' '.join(todos('mant_falla'))
    fecha = None
    for d in ('mant_fecha', 'fecha'):
        if d in reg:
            fecha = _fecha_valor(reg[d][0][1])
            if fecha:
                break
    tipo_txt = txt('mant_tipo')
    if fecha and (que or falla or tipo_txt or 'mant_fecha' in reg):
        tipo_m = tipo_mantenimiento(tipo_txt) or tipo_mantenimiento(f'{falla} {que}') or 'preventivo'
        partes = []
        if falla:
            partes.append(f'Falla: {falla}' if que else falla)
        if que:
            partes.append(f'Se hizo: {que}' if falla else que)
        if txt('mant_tecnico'):
            partes.append(f"Técnico: {txt('mant_tecnico')}")
        descripcion_m = '. '.join(partes) or f'Mantenimiento {tipo_m}'
        costo = _numero(reg['mant_costo'][0][1]) if 'mant_costo' in reg else None
        if fecha <= date.today():
            fila['historial'].append({'fecha': fecha.isoformat(), 'tipo': tipo_m, 'descripcion': descripcion_m[:1000],
                                      'costo': costo if costo and costo < 1e11 else None})
        elif not fila['plan']['proximo']:
            fila['plan']['proximo'] = fecha.isoformat()
    else:
        if 'fecha' in reg:
            _agregar_nota(fila, f"{reg['fecha'][0][0]}: {txt('fecha')}")
        for etiqueta, v in reg.get('mant_descripcion', []) + reg.get('mant_falla', []) + reg.get('mant_tecnico', []):
            _agregar_nota(fila, f'{etiqueta}: {_celda(v)}')
        if 'mant_costo' in reg:
            _agregar_nota(fila, f"{reg['mant_costo'][0][0]}: {txt('mant_costo')}")
    if 'mant_proximo' in reg:
        f = _fecha_valor(reg['mant_proximo'][0][1])
        if f:
            fila['plan']['proximo'] = f.isoformat()
        else:
            _agregar_nota(fila, f"{reg['mant_proximo'][0][0]}: {txt('mant_proximo')}")
    if 'mant_frecuencia' in reg:
        meses = meses_de(txt('mant_frecuencia'))
        if meses:
            fila['plan']['cada_meses'] = meses
        else:
            _agregar_nota(fila, f"{reg['mant_frecuencia'][0][0]}: {txt('mant_frecuencia')}")
    for etiqueta, v in reg.get('notas', []):
        _agregar_nota(fila, f'{etiqueta}: {_celda(v)}' if etiqueta else _celda(v))
    if not e.get('imei'):
        for _etq, v in reg.get('notas', []) + reg.get('especificaciones', []) + reg.get('descripcion', []):
            m = re.search(r'(?i)imei\D{0,6}(\d[\d\s-]{13,18}\d)', _celda(v))
            if m and lector.imei_luhn(re.sub(r'\D', '', m.group(1))):
                e['imei'] = re.sub(r'\D', '', m.group(1))
                break
    # f) Tipo de disco a partir del almacenamiento («1TB HDD» → HDD).
    if e.get('almacenamiento') and not e['extras'].get('tipo_disco'):
        _cols, exts = specs_de_texto(e['almacenamiento'])
        if exts.get('tipo_disco'):
            e['extras']['tipo_disco'] = exts['tipo_disco']
    return fila


def _sin_palabras_de_tipo(texto):
    """El texto sin las palabras que dicen el tipo («Portátil HP gris» → «HP gris»)."""
    palabras = {p for _c, lista in tipos._PALABRAS_TIPO for p in lista}
    resto = [w for w in str(texto or '').split() if _norm(w) not in palabras and _norm(w) not in ('de', 'del')]
    resto = ' '.join(resto).strip(' -:,')
    return resto if re.search(r'[A-Za-z0-9]{2}', resto) else None


def _filas_de_tabla(parte, mapeo_ia=None):
    columnas, _ = _columnas_de_tabla(parte, mapeo_ia)
    titulos = {int(j): t for j, t in parte['titulos'].items()}
    filas, por_clave = [], {}
    for numero, fila_excel in parte['filas']:
        fila = _cerrar(_fila_de_tabla(fila_excel, columnas, titulos))
        if not tiene_datos_fila(fila) and not fila['historial']:
            continue
        fila['filas_excel'] = [numero]
        claves = claves_de(fila['equipo'])
        destino = next((por_clave[k] for k in claves if k in por_clave), None)
        if destino is not None:                       # mismo equipo en otra fila (su historial)
            _unir(destino, fila)
            destino['filas_excel'].append(numero)
            for k in claves_de(destino['equipo']):
                por_clave.setdefault(k, destino)
            continue
        filas.append(fila)
        for k in claves:
            por_clave[k] = fila
    for f in filas:
        _cerrar(f)
        numeros = f.pop('filas_excel')
        if len(numeros) == 1:
            f['origen'] = f'{parte["titulo"]}, fila {numeros[0]}'
        else:
            lista = ', '.join(str(x) for x in numeros[:8]) + ('…' if len(numeros) > 8 else '')
            f['origen'] = f'{parte["titulo"]}, filas {lista}'
    return filas


# ── Documentos (PDF, texto, hojas tipo ficha) ───────────────────
_ETIQ_NOMBRE = re.compile(r'^[ \t]*(?:nombres?\s+(?:del\s+)?cliente|cliente|propietario|due[ñn]o|'
                          r'raz[oó]n\s+social|se[ñn]or(?:\(a\)|a)?|sr\.?|sra\.?|contratante|'
                          r'nombres?(?:\s+y\s+apellidos)?)[ \t]*[:\-][ \t]*(.+)$', re.I | re.M)
_ETIQ_EMPRESA = re.compile(r'^[ \t]*(?:empresa|entidad|compa[ñn][ií]a|instituci[oó]n)[ \t]*[:\-][ \t]*(.+)$',
                           re.I | re.M)
_ETIQ_TEL = re.compile(r'\b(?:tel[eé]fonos?|tel\.?|celular|cel\.?|m[oó]vil|whatsapp|contacto)[^:\n]{0,15}[:\-]?'
                       r'[ \t]*(\+?[\d][\d \-().]{6,20})', re.I)
_ETIQ_CORREO = re.compile(r'(?:correo|e-?mail|mail)[^:\n@]{0,15}[:\-]?[ \t]*([\w.+-]+@[\w-]+(?:\.[\w-]+)+)', re.I)
_ETIQ_CIUDAD = re.compile(r'^[ \t]*(?:ciudad|municipio)[ \t]*[:\-][ \t]*(.+)$', re.I | re.M)
_ETIQ_DIRECCION = re.compile(r'^[ \t]*direcci[oó]n[ \t]*[:\-][ \t]*(.+)$', re.I | re.M)
_ETIQ_DOC = re.compile(r'\b(?:c\.?\s?c\.?|c[eé]dula|nit|documento)\s*(?:n[°ºo.]?\s*)?[:#]?\s*(\d[\d.\s-]{5,14}\d)',
                       re.I)
_ETIQ_TIPO = re.compile(r'^[ \t]*(?:tipo\s+de\s+equipo|tipo|clase\s+de\s+equipo|equipo|dispositivo)[ \t]*[:\-]'
                        r'[ \t]*(.+)$', re.I | re.M)
_ETIQ_NOTAS = re.compile(r'^[ \t]*(c[oó]digo(?:\s+de)?\s+inventario|c[oó]digo|placa(?:\s+de\s+inventario)?|'
                         r'activo\s+fijo|n[uú]mero\s+de\s+inventario|[aá]rea|dependencia|ubicaci[oó]n|sede|oficina|'
                         r'piso|usuario(?:\s+asignado|\s+responsable)?|responsable|funcionario|cargo|'
                         r'nombre\s+del\s+equipo|hostname|antivirus|office|garant[ií]a|fecha\s+de\s+compra|'
                         r'proveedor|accesorios|estado\s+f[ií]sico|observaciones?)[ \t]*[:\-][ \t]*(.{1,160})$',
                         re.I | re.M)
_CORTE = re.compile(r'\s{2,}|\t| \| |\s+(?:c\.?\s?c\.?|nit|tel[eé]fono|tel\.|cel\.|celular|correo|e-?mail|'
                    r'direcci[oó]n|ciudad)\b', re.I)
_LINEA_NO_MANT = re.compile(r'compra|adquisici|factura|garant|nacimiento|elaboraci|expedici|impres[oa]|generad|'
                            r'vencimiento|ingreso al inventario', re.I)
_PALABRAS_MANT = re.compile(r'mantenimiento|mtto|limpieza|preventiv|correctiv|repar|cambio|reemplaz|instal|'
                            r'formate|revisi|diagn|actualiz|servicio|soporte|backup|copia de seguridad|pasta|ajuste|'
                            r'optimiz|soplad|aspirad|desinfec|configur|arregl|falla|da[ñn]o', re.I)
_SECCION_HIST = re.compile(r'historial|hist[oó]rico|mantenimientos|registro de (?:servicios|mantenimiento)|'
                           r'bit[aá]cora|servicios realizados|control de mantenimiento', re.I)
_PROXIMO = re.compile(r'pr[oó]xim|siguiente|programad', re.I)


_LINEA_SECRETA = re.compile(r'(?im)^.*\b(?:contrase[ñn]as?|claves?|password|pass|pin|puk|patr[oó]n|'
                            r'c[oó]digo\s+de\s+desbloqueo)\b\s*(?:de\s+\w+\s*)?[:=].*$')


def sin_secretos(texto):
    """Quita las líneas de claves, contraseñas y patrones de desbloqueo."""
    return _LINEA_SECRETA.sub('', texto or '')


def contacto_negocio():
    """Teléfonos, correos y nombre del propio negocio: el membrete de sus
    documentos no se toma por datos del cliente."""
    try:
        from services.public_site_service import get_brand_config
        marca = get_brand_config() or {}
    except Exception:  # noqa: BLE001
        marca = {}
    return {'telefonos': sorted({t for t in (_telefono(marca.get('empresa_telefono')),
                                             _telefono(marca.get('empresa_whatsapp'))) if t}),
            'correos': sorted({(marca.get('empresa_email') or '').strip().lower()} - {''}),
            'nombre': _norm(marca.get('empresa_nombre'))}


def _valor_limpio(valor):
    valor = _CORTE.split(str(valor or ''))[0].strip(' .,;:-')
    return valor if re.search(r'[A-Za-zÁÉÍÓÚÑáéíóúñ]{2}', valor) and _norm(valor) not in _SIN_SERIAL else None


def cliente_de_texto(texto, negocio=None):
    """Cliente que dice el documento, SOLO con reglas: nombre, empresa,
    teléfono y correo con su etiqueta, ciudad, dirección y documento.

    Si el documento dice quién es el cliente («Cliente: …»), el teléfono, el
    correo y el documento se toman de ahí en adelante: lo que va antes suele
    ser el membrete del taller. Entre varios teléfonos se prefiere el celular."""
    texto = texto or ''
    negocio = negocio or {'telefonos': [], 'correos': [], 'nombre': ''}
    c, inicio = {}, None
    for patron, clave in ((_ETIQ_NOMBRE, 'nombre'), (_ETIQ_EMPRESA, 'empresa'), (_ETIQ_CIUDAD, 'ciudad'),
                          (_ETIQ_DIRECCION, 'direccion')):
        for m in patron.finditer(texto):
            valor = _valor_limpio(m.group(1))
            if valor and _norm(valor) != negocio['nombre'] and not (clave == 'nombre' and len(valor.split()) > 8):
                c[clave] = valor[:150]
                if clave in ('nombre', 'empresa'):
                    inicio = m.start() if inicio is None else min(inicio, m.start())
                break

    def del_cliente(encontrados):
        """Los de después del nombre del cliente (o todos si no dice quién es)."""
        return [v for pos, v in encontrados if inicio is None or pos >= inicio]

    tels = del_cliente([(m.start(), _telefono(m.group(1))) for m in _ETIQ_TEL.finditer(texto)])
    tels = [t for t in tels if t and t not in negocio['telefonos']]
    if tels:
        celulares = [t for t in tels if t.startswith('3')]
        c['telefono'] = celulares[0] if celulares else tels[0]
        fijos = [t for t in tels if not t.startswith('3')]
        if celulares and fijos:
            c['whatsapp'], c['telefono'] = celulares[0], fijos[0]
    correos = [x for x in del_cliente([(m.start(), m.group(1).lower()) for m in _ETIQ_CORREO.finditer(texto)])
               if x not in negocio['correos']]
    if correos:
        c['email'] = correos[0][:150]
    documentos = del_cliente([(m.start(), re.sub(r'\s', '', m.group(1))) for m in _ETIQ_DOC.finditer(texto)])
    if documentos:
        c['documento'] = documentos[0]
    return c


def _tipo_de_documento(texto):
    m = _ETIQ_TIPO.search(texto or '')
    if m:
        tipo = tipos.tipo_desde_texto(m.group(1))
        if tipo:
            return tipo
    titulo = ' '.join(l for l in (texto or '').split('\n')[:6] if len(l) < 90)
    return tipos.tipo_desde_texto(re.sub(r'(?i)hoja de vida|ficha t[eé]cnica|de equipos?|de c[oó]mputo', ' ', titulo))


def historial_de_texto(texto, hoy=None):
    """Mantenimientos con fecha del documento (líneas con fecha y palabras de
    mantenimiento, o dentro de una sección «Historial»)."""
    hoy = hoy or date.today()
    entradas, en_seccion = [], False
    for linea in (texto or '').replace('\r', '').split('\n'):
        linea = linea.strip()
        if not linea:
            continue
        fechas = fechas_en(linea)
        if not fechas:
            if _SECCION_HIST.search(linea) and len(linea) < 80:
                en_seccion = True
            continue
        if _LINEA_NO_MANT.search(linea) or _PROXIMO.search(linea):
            continue
        f, ini, fin = fechas[0]
        if f > hoy or not (_PALABRAS_MANT.search(linea) or en_seccion):
            continue
        desc = linea[:ini] + ' ' + linea[fin:]
        desc = re.sub(r'(?i)^\s*(?:fecha(?:\s+(?:de|del)\s+\w+)*|f\.)\s*[:\-]?', '', desc)
        desc = ' '.join(re.sub(r'\s*[|;\t]\s*', ', ', desc).split()).strip(' ,:-–')
        tipo_m = tipo_mantenimiento(desc) or 'preventivo'
        desc = _sin_rotulo(desc)
        entradas.append({'fecha': f.isoformat(), 'tipo': tipo_m,
                         'descripcion': (desc if len(desc) >= 4 else f'Mantenimiento {tipo_m}')[:1000],
                         'costo': None})
    return _sin_repetidos(entradas)[:HISTORIAL_MAX]


def plan_de_texto(texto):
    plan = {'cada_meses': None, 'proximo': None}
    for linea in (texto or '').split('\n'):
        if _PROXIMO.search(linea) and re.search(r'(?i)mantenimiento|mtto|revisi|servicio|visita', linea):
            fechas = fechas_en(linea)
            if fechas and not plan['proximo']:
                plan['proximo'] = fechas[0][0].isoformat()
        if re.search(r'(?i)frecuencia|periodicidad|mantenimiento|mtto', linea) and not re.search(r'(?i)garant', linea):
            meses = meses_de(re.sub(r'\d{1,4}[/\-.]\d{1,2}[/\-.]\d{2,4}', ' ', linea))
            if meses and not plan['cada_meses']:
                plan['cada_meses'] = meses
    return plan


def _bloques(paginas):
    """Si cada página es la ficha de un equipo distinto (seriales distintos),
    un bloque por equipo; si no, None (el documento es un solo equipo)."""
    if len(paginas) < 2:
        return None
    bloques, vistas, equipos = [], set(), 0
    for pagina in paginas:
        claves = claves_de(lector.leer(lector.limpiar_personales(pagina))['columnas'])
        if claves and not claves & vistas:
            vistas |= claves
            equipos += 1
            bloques.append(pagina)
        elif bloques:
            bloques[-1] += '\n' + pagina
        else:
            bloques.append(pagina)
    return bloques if equipos >= 2 else None


def _fila_de_texto(texto, cliente):
    fila = _fila_vacia()
    fila['cliente'] = dict(cliente)
    limpio = lector.limpiar_personales(texto)
    tipo = _tipo_de_documento(texto)
    leido = lector.leer(limpio, tipo)
    e = fila['equipo']
    e.update(leido['columnas'])
    e['extras'] = dict(leido['extras'])
    if e.get('imei'):
        imei = e.pop('imei')
        _poner_imei(fila, imei)
    cols, exts = specs_de_texto(limpio)
    for k, v in cols.items():
        e[k] = e.get(k) or v
    for k, v in exts.items():
        e['extras'].setdefault(k, v)
    if exts.get('tipo_disco') == 'SSD NVMe' and e['extras'].get('tipo_disco') == 'SSD SATA' \
            and re.search(r'(?i)nvme|m\.2', limpio):
        e['extras']['tipo_disco'] = 'SSD NVMe'
    if not e.get('marca') or not e.get('modelo'):
        m = re.search(r'(?im)^[ \t]*(?:equipo|descripci[oó]n|referencia)[ \t]*[:\-][ \t]*(.+)$', texto)
        if m:
            marca, modelo = _marca_modelo(m.group(1))
            e['marca'] = e.get('marca') or marca
            e['modelo'] = e.get('modelo') or modelo
    e['tipo'] = tipo
    fila['historial'] = historial_de_texto(texto)
    fila['plan'] = plan_de_texto(texto)
    for m in _ETIQ_NOTAS.finditer(texto):
        _agregar_nota(fila, f'{m.group(1).strip().capitalize()}: {m.group(2).strip()}')
    return fila


def _filas_de_texto(parte, negocio):
    texto = parte['texto']
    cliente = cliente_de_texto(texto, negocio)
    # Ficha del taller en Excel (o texto): «EQUIPO ASIGNADO A PERSONA - EMPRESA».
    asignado, empresa = emp.titular_de_ficha(texto)
    if empresa and not cliente.get('empresa'):
        cliente['empresa'] = empresa[:150]
    if asignado and not cliente.get('nombre') and not cliente.get('empresa'):
        cliente['nombre'] = asignado[:150]
    bloques = _bloques(parte.get('paginas') or []) or [texto]
    filas = []
    for k, bloque in enumerate(bloques):
        fila = _cerrar(_fila_de_texto(bloque, cliente))
        fila['origen'] = parte['titulo'] + (f' (equipo {k + 1} de {len(bloques)})' if len(bloques) > 1 else '')
        if asignado and empresa:
            _agregar_nota(fila, f'Asignado a: {asignado}')
        filas.append(fila)
    return filas


def _nombres_para_tapar(parte, filas):
    """Nombres de personas que hay que tapar antes de mandar el texto a la IA."""
    nombres = set()
    for f in filas:
        for k in ('nombre', 'empresa'):
            if f['cliente'].get(k):
                nombres.add(f['cliente'][k])
    for m in re.finditer(r'(?im)^[ \t]*(?:usuario(?:\s+asignado)?|responsable|funcionario|t[eé]cnico|elabor[oó]|'
                         r'revis[oó]|aprob[oó]|recibi[oó]|entreg[oó]|firma)[ \t]*[:\-][ \t]*(.+)$',
                         parte.get('texto') or ''):
        valor = _valor_limpio(m.group(1))
        if valor:
            nombres.add(valor)
    return tuple(nombres)


def _mezclar_ia(filas, ia):
    """Completa las filas leídas con reglas con lo que leyó la IA (lo de las
    reglas manda; la IA llena lo que falta)."""
    libres = list(ia.get('equipos') or [])
    pares, usadas = [], set()
    if len(filas) == 1 and len(libres) == 1:          # un documento, un equipo: se completan
        pares.append((filas[0], libres.pop()))
        usadas.add(id(filas[0]))
    # 1) Mismo serial o IMEI.
    for fila in filas:
        claves = claves_de(fila['equipo'])
        par = next((eq for eq in libres if claves & claves_de(eq['columnas'])), None)
        if par is not None:
            libres.remove(par)
            pares.append((fila, par))
            usadas.add(id(fila))
    # 2) En orden, con un equipo de la IA cuyo serial no choque.
    for fila in filas:
        if id(fila) in usadas:
            continue
        claves = claves_de(fila['equipo'])
        par = next((eq for eq in libres if not (claves and claves_de(eq['columnas']))), None)
        if par is not None:
            libres.remove(par)
            pares.append((fila, par))
            usadas.add(id(fila))
    # 3) Equipos que solo encontró la IA (el documento describe varios).
    base = filas[0] if filas else _fila_vacia()
    nuevas = []
    for k, eq in enumerate(libres):
        fila = _fila_vacia()
        fila['cliente'] = dict(base.get('cliente') or {})
        fila['origen'] = (base.get('origen') or '') + ' (otro equipo del mismo documento)'
        nuevas.append(fila)
        pares.append((fila, eq))
    filas = list(filas) + nuevas
    for fila, eq in pares:
        e = fila['equipo']
        if (fila['tipo_adivinado'] or not tipos.es_valido(e.get('tipo') or '') or e.get('tipo') == 'otro') \
                and eq['tipo'] != 'otro':
            e['tipo'], fila['tipo_adivinado'] = eq['tipo'], False
        for k, v in eq['columnas'].items():
            if k == 'imei':
                if not e.get('imei'):
                    _poner_imei(fila, v)
            elif not e.get(k):
                e[k] = v
        for k, v in eq['extras'].items():
            e['extras'].setdefault(k, v)
        fila['historial'] = _sin_repetidos(fila['historial'] + [dict(h, costo=None) for h in eq['historial']])
        fila['plan']['proximo'] = fila['plan'].get('proximo') or eq.get('proximo')
        fila['plan']['cada_meses'] = fila['plan'].get('cada_meses') or eq.get('cada_meses')
        fila['resumen'] = eq.get('resumen') or fila.get('resumen')
        fila['fuente'] = 'ia'
        _cerrar(fila)
    if ia.get('aviso'):
        for fila in filas:
            if ia['aviso'] not in fila['avisos']:
                fila['avisos'].append(ia['aviso'])
    return filas


# ── Archivos → partes ───────────────────────────────────────────
def partes_de_archivo(clase, datos):
    """Lista de partes del archivo: {'tipo': 'tabla'|'texto'|'escaneado', ...}.
    Mismo archivo → mismas partes en el mismo orden (la IA las relee por índice)."""
    if clase == 'excel':
        return _partes_excel(datos)
    if clase == 'csv':
        return _partes_csv(datos)
    if clase == 'texto':
        texto = _decodificar(datos)
        return [{'tipo': 'texto', 'titulo': 'Texto', 'texto': texto, 'paginas': [texto]}] if texto.strip() else []
    if clase == 'pdf':
        return _partes_pdf(datos)
    raise ErrorServicio('Formato no admitido.')


def _decodificar(datos):
    for codificacion in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode('utf-8', errors='replace')


def _es_rojo(color):
    try:
        rgb = color.rgb if color is not None and color.type == 'rgb' else None
    except Exception:  # noqa: BLE001
        return False
    if not isinstance(rgb, str) or not re.fullmatch(r'[0-9A-Fa-f]{6,8}', rgb):
        return False
    r, g, b = int(rgb[-6:-4], 16), int(rgb[-4:-2], 16), int(rgb[-2:], 16)
    return r >= 190 and g <= 120 and b <= 120


def _parte_ficha_excel(hoja, n_filas, datos):
    """Hoja con la ficha técnica del taller → parte 'ficha' (como el PDF): sus
    celdas con el color de fondo y de letra (rojo = presenta inconvenientes)."""
    rejilla = []
    for fila in hoja.iter_rows(max_row=min(n_filas, 200), max_col=COLUMNAS_MAX):
        celdas = []
        for j, celda in enumerate(fila):
            valor = getattr(celda, 'value', None)
            if valor in (None, ''):
                continue
            texto = valor.date().isoformat() if isinstance(valor, datetime) else (
                valor.isoformat() if isinstance(valor, date) else _celda(valor))
            fondo = letra = False
            try:
                fondo = celda.fill is not None and celda.fill.fill_type == 'solid' and _es_rojo(celda.fill.fgColor)
                letra = celda.font is not None and _es_rojo(celda.font.color)
            except Exception:  # noqa: BLE001
                pass
            celdas.append((j, texto, fondo, letra))
        rejilla.append(celdas)
    ficha = fichas.leer_hoja(rejilla)
    if not ficha:
        return None
    ficha['fotos'] = len(fichas.fotos_de_excel(datos, hoja.title))
    texto = '\n'.join(' | '.join(c[1] for c in fila) for fila in rejilla if fila)
    return {'tipo': 'ficha', 'titulo': f'Hoja «{hoja.title}»', 'texto': texto, 'ficha': ficha, 'hoja': hoja.title}


def _partes_excel(datos):
    try:
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise ErrorServicio('No se pudo abrir el Excel. Si es de una versión vieja (.xls), ábrelo y guárdalo como '
                            '«Libro de Excel (.xlsx)».')
    partes = []
    try:
        for hoja in libro.worksheets[:HOJAS_MAX]:
            if getattr(hoja, 'sheet_state', 'visible') != 'visible' or _HOJAS_DE_AYUDA.match(_norm(hoja.title)):
                continue
            filas = []
            for fila in hoja.iter_rows(values_only=True, max_col=COLUMNAS_MAX):
                filas.append(list(fila))
                if len(filas) > FILAS_LEER_MAX + 30:
                    break
            if fichas.es_hoja_ficha([_celda(v) for fila in filas[:80] for v in fila]):
                parte = _parte_ficha_excel(hoja, len(filas), datos)
                if parte:                            # la plantilla en blanco no trae equipo
                    partes.append(parte)
                continue
            parte = _parte_de_hoja(filas, f'Hoja «{hoja.title}»')
            if parte:
                partes.append(parte)
    finally:
        libro.close()
    return partes


def _partes_csv(datos):
    texto = _decodificar(datos)
    muestra = texto[:5000]
    separador = max((';', ',', '\t', '|'), key=lambda s: muestra.count(s))
    filas = [f for f in csv.reader(io.StringIO(texto), delimiter=separador)][:FILAS_LEER_MAX + 30]
    parte = _parte_de_hoja(filas, 'Tabla')
    return [parte] if parte else []


def _partes_pdf(datos):
    try:
        from pypdf import PdfReader
        lector_pdf = PdfReader(io.BytesIO(datos))
        if lector_pdf.is_encrypted and not lector_pdf.decrypt(''):
            raise ErrorServicio('El PDF tiene clave: ábrelo, guárdalo sin clave y vuelve a subirlo.')
        paginas = []
        for pagina in lector_pdf.pages[:PAGINAS_MAX]:
            try:
                paginas.append(pagina.extract_text() or '')
            except Exception:  # noqa: BLE001
                paginas.append('')
    except ErrorServicio:
        raise
    except Exception:  # noqa: BLE001
        raise ErrorServicio('No se pudo abrir el PDF (puede estar dañado).')
    texto = '\n'.join(paginas)
    if len(re.sub(r'\s', '', texto)) < 40:
        return [{'tipo': 'escaneado', 'titulo': 'PDF', 'texto': ''}]
    from services import servicio_tecnico_ficha as fichas
    ficha = fichas.leer_pdf(datos)
    if ficha:                                        # «FICHA TECNICA» del taller: se lee completa, sin IA
        return [{'tipo': 'ficha', 'titulo': 'Ficha técnica', 'texto': texto, 'ficha': ficha}]
    return [{'tipo': 'texto', 'titulo': 'PDF', 'texto': texto, 'paginas': paginas}]


def _fila_de_ficha(f):
    """Fila del lote desde una ficha técnica del taller (formato reconocido):
    equipo, cliente (la empresa; la persona queda como «asignado a»), último
    mantenimiento, próximo y la ficha completa para volver a generarla igual."""
    fila = _fila_vacia()
    datos = dict(f.get('filas') or [])
    e = fila['equipo']
    if _descruzar_cpu_gpu(datos):
        fila['avisos'].append('En la ficha, el procesador y la tarjeta gráfica venían cruzados: se pusieron en su sitio.')
    descripcion = f.get('descripcion') or ''
    modelo_txt = datos.get('Modelo') or ''
    if any(re.search(r'(?i)impresi[oó]n|tinta|cartucho|t[oó]ner|tolva', etq) for etq in datos):
        e['tipo'] = 'impresora'                      # sus filas son de impresora, diga lo que diga la descripción
    elif re.search(r'(?i)port[aá]til|laptop|notebook', f'{descripcion} {modelo_txt}'):
        e['tipo'] = 'portatil'
    elif re.search(r'(?i)escritorio|todo en uno|all in one|torre|desktop', descripcion):
        e['tipo'] = 'computador'
    else:
        e['tipo'] = tipos.tipo_desde_texto(descripcion) or tipos.tipo_desde_texto(modelo_txt)
    fabricante = (datos.get('Fabricante del Sistema') or datos.get('Fabricante') or datos.get('Marca')
                  or f.get('fabricante') or '')
    if fabricante:
        e['marca'] = _marca_canonica(fabricante)
    sku = re.search(r'SKU[:\s]+([\w#._/-]+)', modelo_txt)
    nombre_equipo = re.search(r'Equipo:\s*([\w.-]+)', modelo_txt)
    modelo = re.split(r'\s*\(\s*SKU\b|\s+-\s*SKU\b|\s+-\s*Equipo:', modelo_txt)[0].strip(' -')
    if modelo and e.get('marca'):                    # «Hp Pavilion 24» → «Pavilion 24» (la marca va aparte)
        modelo = re.sub(r'(?i)^' + re.escape(e['marca']) + r'\s+', '', modelo).strip() or modelo
    e['modelo'] = modelo or None
    e['sistema_operativo'] = re.sub(r'(?i)^microsoft\s+', '', datos.get('Sistema Operativo') or '') or None
    e['procesador'] = datos.get('Procesador') or None
    ram_txt = datos.get('Memoria RAM') or ''
    ram = lector._gb(ram_txt)
    if ram:
        ddr = re.search(r'\b(LP)?DDR\d\w?\b', ram_txt)
        e['ram'] = f'{ram} {ddr.group(0)}' if ddr else ram
    usadas = re.search(r'(?i)(\d)\s+de\s+(\d)\s+(?:ranuras|slots|sockets)', ram_txt)
    ocupadas = re.search(r'(?i)(\d)\s+(?:ranuras|slots|sockets)[^.;]*?ocupad', ram_txt)
    sin_libres = re.search(r'(?i)sin\s+(?:slots?|ranuras?|sockets?|slock)\b', ram_txt)
    if usadas:
        e['extras']['ram_ranuras_total'] = usadas.group(2)
        e['extras']['ram_ranuras_libres'] = str(max(0, int(usadas.group(2)) - int(usadas.group(1))))
    elif ocupadas:
        e['extras']['ram_ranuras_total'] = ocupadas.group(1)
    if sin_libres and not usadas:
        e['extras']['ram_ranuras_libres'] = '0'
    discos = discos_de_texto(datos.get('Unidad de Disco'))
    if discos:                                       # «SSD 239 GB + HDD 932 GB»
        e['almacenamiento'] = ' + '.join(' '.join(x for x in ('SSD' if (t or '').startswith('SSD') else t, cap) if x)
                                         for t, cap in discos)[:120]
        clases = {t for t, _cap in discos if t}
        if 'HDD' in clases and any(t.startswith('SSD') for t in clases):
            e['extras']['tipo_disco'] = 'Mixto (SSD + HDD)'
        elif len(clases) == 1:
            e['extras']['tipo_disco'] = clases.pop()
    else:
        cols, exts = specs_de_texto(datos.get('Unidad de Disco') or '')
        e['almacenamiento'] = cols.get('almacenamiento') or (datos.get('Unidad de Disco') or '')[:120] or None
        e['extras'].update(exts)
    # La sección «Datos del fabricante» copiada de la ficha de otro equipo (p. ej.
    # «Dell» en la de una impresora HP): se deja el soporte de la marca del equipo.
    fab_seccion = _marca_canonica(f.get('fabricante'))
    fab_ajeno = bool(fab_seccion and e.get('marca') and _norm(fab_seccion) != _norm(e['marca'])
                     and _norm(fab_seccion) in {_norm(m) for m in MARCAS}
                     and _norm(e['marca']) in {_norm(m) for m in MARCAS})
    if fab_ajeno:
        fila['avisos'].append(f"La ficha dice que el fabricante es «{f['fabricante']}», pero el equipo es "
                              f"{e['marca']}: se deja el soporte de {e['marca']}.")
    if e['tipo'] == 'computador' and datos.get('Placa Base'):
        e['extras']['board'] = datos['Placa Base']
    if f.get('empresa'):
        fila['cliente'] = {'empresa': f['empresa']}
    elif f.get('asignado_a'):
        fila['cliente'] = {'nombre': f['asignado_a']}
    fecha, proximo = _iso(f.get('ultimo') or f.get('fecha')), _iso(f.get('proximo'))
    if fecha and fecha <= date.today():
        fila['historial'].append({'fecha': fecha.isoformat(), 'tipo': 'preventivo', 'costo': None,
                                  'descripcion': ('Mantenimiento registrado en la ficha técnica '
                                                  + (f.get('codigo') or '')).strip()})
    if proximo:
        fila['plan']['proximo'] = proximo.isoformat()
        if fecha:
            meses = (proximo.year - fecha.year) * 12 + proximo.month - fecha.month
            if 1 <= meses <= 36:
                fila['plan']['cada_meses'] = meses
    fila['resumen'] = descripcion or None
    fila['ficha'] = {k: v for k, v in {
        'estado': f.get('estado'), 'version': f.get('version'), 'fecha': f.get('fecha'),
        'asignado_a': f.get('asignado_a'), 'empresa': f.get('empresa'), 'descripcion': descripcion, 'filas': datos,
        'orden': list(datos),                        # el orden del documento (la base de datos no lo guarda)
        'destacar': f.get('destacar'), 'vida_util': f.get('vida_util'), 'recomendaciones': f.get('recomendaciones'),
        'fabricante': None if fab_ajeno else f.get('fabricante'),
        'fabricante_detalle': None if fab_ajeno else f.get('fabricante_detalle'),
        'fabricante_url': None if fab_ajeno else f.get('fabricante_url'), 'ultimo': f.get('ultimo'),
        'problema': f.get('problema'), 'requerimientos': f.get('requerimientos')}.items()
        if v not in (None, '', [], {})}
    fila['ficha_codigo'] = (f.get('codigo') or '')[:40] or None
    fila['fotos_pdf'] = bool(f.get('fotos'))
    if f.get('asignado_a') and f.get('empresa'):
        _agregar_nota(fila, f"Asignado a: {f['asignado_a']}")
    if sku:
        _agregar_nota(fila, f'SKU: {sku.group(1)}')
    if nombre_equipo:
        _agregar_nota(fila, f'Nombre del equipo: {nombre_equipo.group(1)}')
    fila['fuente'] = 'ficha'
    fila['origen'] = 'Ficha técnica' + (f" {f['codigo']}" if f.get('codigo') else '')
    return _cerrar(fila)


def filas_de_parte(parte, archivo, indice, negocio, ia=None, mapeo_ia=None):
    """Filas del lote que salen de una parte (con reglas, más la IA si la hay)."""
    if parte['tipo'] == 'tabla':
        filas = _filas_de_tabla(parte, mapeo_ia)
    elif parte['tipo'] == 'escaneado':
        fila = _fila_vacia()
        fila.update(clase='adjunto', origen=parte['titulo'])
        filas = [fila]
    elif parte['tipo'] == 'ficha':
        filas = [_fila_de_ficha(parte['ficha'])]
        if parte.get('hoja'):
            filas[0]['fotos_hoja'] = parte['hoja']     # las fotos salen de esa hoja del libro
    else:
        filas = _filas_de_texto(parte, negocio)
        if ia:
            filas = _mezclar_ia(filas, ia)
    for fila in filas:
        fila['archivo_id'], fila['parte'] = archivo['id'], indice
        fila['archivo'] = archivo['nombre']
    if archivo.get('clase') in ('pdf', 'texto') and filas:
        filas[0]['documento_id'] = archivo['id']
        for fila in filas[1:]:
            _agregar_nota(fila, f"Sale del documento «{archivo['nombre']}» (guardado en el primer equipo de ese archivo).")
    return filas


# ── Lote: crear, leer con IA, ajustes ───────────────────────────
def _guardar(cur, lote_id, **campos):
    sets, params = ['actualizado_en = NOW()'], []
    for clave, valor in campos.items():
        if clave in ('archivos', 'filas', 'ajustes', 'resultado'):
            sets.append(f'{clave} = %s::jsonb')
            params.append(json.dumps(valor, ensure_ascii=False, default=str))
        else:
            sets.append(f'{clave} = %s')
            params.append(valor)
    cur.execute(f"UPDATE st_importaciones SET {', '.join(sets)} WHERE id = %s", (*params, lote_id))


def obtener_lote(lote_id, cur=None, bloquear=False):
    def _leer(c):
        c.execute('SELECT * FROM st_importaciones WHERE id = %s' + (' FOR UPDATE' if bloquear else ''), (lote_id,))
        fila = c.fetchone()
        return dict(fila) if fila else None
    if cur is not None:
        return _leer(cur)
    with get_db_cursor(dict_cursor=True) as c:
        return _leer(c)


def lotes_recientes(limite=10):
    with get_db_cursor(dict_cursor=True) as cur:
        cur.execute("""SELECT id, estado, creado_en, importado_en, jsonb_array_length(archivos) AS archivos,
                              jsonb_array_length(filas) AS filas, resultado
                       FROM st_importaciones ORDER BY id DESC LIMIT %s""", (limite,))
        return [dict(r) for r in cur.fetchall()]


def _numerar(filas, desde=1):
    for k, fila in enumerate(filas):
        fila['n'] = desde + k
    return filas


def _necesita_ia(parte, filas):
    if parte['tipo'] == 'texto':
        return True
    if parte['tipo'] == 'tabla':
        return bool(_columnas_de_tabla(parte)[1])
    return False


# ── Carpeta de cada equipo: ¿de qué empresa es? ────────────────
_ETIQ_PERSONA = re.compile(r'(?im)^[ \t]*(?:usuario(?:\s+asignado|\s+responsable)?|responsable|funcionario|'
                           r't[eé]cnico|elabor[oó]|revis[oó]|aprob[oó]|recibi[oó]|entreg[oó]|firma|'
                           r'asignado\s+a)[ \t]*[:\-][ \t]*(.+)$')


def _candidatos(fila):
    """(nombre, puesto) que dicen de quién es el equipo, en orden de confianza:
    la empresa de la ficha, la empresa del cliente y el nombre del cliente."""
    salida, vistas = [], set()
    fx = fila.get('ficha') or {}
    for nombre, puesto in ((fx.get('empresa'), 'empresa'), (fila['cliente'].get('empresa'), 'empresa'),
                           (fila['cliente'].get('nombre'), 'cliente')):
        nombre = ' '.join(str(nombre or '').split())
        c = emp.clave(nombre)
        if c and c not in vistas:
            vistas.add(c)
            salida.append((nombre, puesto))
    return salida


def _encabezado(partes, personas=()):
    """Lo que dice de quién es el documento (nombre de la hoja, títulos sobre la
    tabla, primeras líneas): sin claves, correos, teléfonos ni documentos, y con
    los nombres de persona tapados. Las fichas del taller no lo necesitan."""
    lineas, nombres = [], {p for p in personas if p}
    for p in partes:
        if p['tipo'] == 'tabla':
            hoja = re.sub(r'^Hoja «(.*)»$', r'\1', p.get('titulo') or '').strip()
            if hoja and not re.fullmatch(r'(?i)(hoja|sheet|tabla|libro)\s*\d*', hoja):
                lineas.append(f'Hoja: {hoja}')
            lineas += p.get('encabezado') or []
        elif p['tipo'] == 'texto':
            texto = p.get('texto') or ''
            lineas += [l.strip() for l in texto.split('\n') if l.strip()][:12]
            asignado, _empresa = emp.titular_de_ficha(texto)
            if asignado:
                nombres.add(asignado)
        if len(lineas) >= 20:
            break
    texto = sin_secretos('\n'.join(lineas[:20]))
    for m in _ETIQ_PERSONA.finditer(texto):
        valor = _valor_limpio(m.group(1))
        if valor:
            nombres.add(valor)
    return lector.sin_datos_personales(texto, tuple(nombres))[:3000].strip()


def _nombres_de_archivo(filas, negocio_nombre):
    """(dudosos, personas, decidida): nombres que la IA debe clasificar, nombres de
    persona (para taparlos) y si alguna fila ya tiene su empresa con reglas."""
    dudosos, personas, decidida = [], [], False
    for f in filas:
        if f['clase'] != 'equipo' or not f.get('incluir', True):
            continue
        for nombre, _puesto in _candidatos(f):
            tipo, _motivo = emp.clasificar(nombre, negocio_nombre)
            if tipo == 'dudoso' and nombre not in dudosos:
                dudosos.append(nombre)
            elif tipo == 'persona':
                personas.append(nombre)
            elif tipo == 'empresa':
                decidida = True
    return dudosos[:30], personas, decidida


_ETIQ_EMPRESA_CAB = re.compile(r'(?i)(?:^|·|\|)\s*(?:empresa|cliente|raz[oó]n social|entidad|compa[ñn][ií]a)'
                               r'\s*(?:·|:|-)\s*([^·|\n]{2,80})')


def _empresa_etiquetada(partes, negocio_nombre):
    """La empresa que dice el encabezado de una tabla con su etiqueta
    («EMPRESA · ACME» encima de la tabla de control de mantenimientos)."""
    for p in partes:
        for linea in (p.get('encabezado') or []) if p['tipo'] == 'tabla' else []:
            m = _ETIQ_EMPRESA_CAB.search(linea)
            if m:
                nombre = emp.nombre_visible(m.group(1))
                if emp.clasificar(nombre, negocio_nombre)[0] in ('empresa', 'dudoso'):
                    return nombre
    return None


def _paso_empresa(partes, filas, negocio_nombre, ia_ok):
    """Estado inicial del paso «¿de qué empresa es?» de un archivo: hace falta la
    IA si hay nombres dudosos, o si ninguna fila dice su empresa (ni el
    encabezado con su etiqueta) y el documento tiene encabezado donde buscarla."""
    dudosos, personas, decidida = _nombres_de_archivo(filas, negocio_nombre)
    etiqueta = _empresa_etiquetada(partes, negocio_nombre)
    necesita = bool(dudosos) or (not decidida and not etiqueta and bool(_encabezado(partes, personas)))
    return {'ia': ('pendiente' if ia_ok else 'sin_ia') if necesita else 'no_aplica', 'intentos': 0,
            'nombre': None, 'clasificacion': {}, 'aviso': None, 'etiqueta': etiqueta}


def _carpeta_de_fila(fila, info, negocio_nombre):
    """{'tipo': 'empresa'|'particular', 'nombre', 'fuente', 'motivo'} de una fila.
    fuente: 'manual' (la escogió la persona), 'reglas', 'ia' o 'supuesto' (dudoso
    sin IA: se toma por empresa solo si el documento lo ponía como empresa)."""
    manual = fila.get('carpeta_manual')
    if manual:
        if manual.get('tipo') == 'empresa' and manual.get('nombre'):
            return {'tipo': 'empresa', 'nombre': manual['nombre'], 'fuente': 'manual', 'motivo': 'La escogiste tú.'}
        return {'tipo': 'particular', 'nombre': emp.PARTICULARES, 'fuente': 'manual', 'motivo': 'La escogiste tú.'}
    info = info or {}
    clasificacion = info.get('clasificacion') or {}
    for nombre, puesto in _candidatos(fila):
        tipo, motivo = emp.clasificar(nombre, negocio_nombre)
        fuente = 'reglas'
        if tipo == 'dudoso':
            dijo_ia = clasificacion.get(emp.clave(nombre))
            if dijo_ia:
                tipo, fuente = dijo_ia, 'ia'
                motivo = 'La IA dijo que es ' + ('una empresa.' if dijo_ia == 'empresa' else 'una persona.')
            elif puesto == 'empresa':
                tipo, fuente = 'empresa', 'supuesto'
                motivo = (f'No se pudo confirmar si «{nombre}» es una empresa: se tomó como empresa porque así '
                          'venía en el documento.')
        if tipo == 'empresa':
            return {'tipo': 'empresa', 'nombre': emp.nombre_visible(nombre), 'fuente': fuente, 'motivo': motivo}
    if info.get('etiqueta'):
        return {'tipo': 'empresa', 'nombre': info['etiqueta'], 'fuente': 'reglas',
                'motivo': f"El documento dice que es de la empresa «{info['etiqueta']}»."}
    if info.get('nombre'):
        return {'tipo': 'empresa', 'nombre': info['nombre'], 'fuente': 'ia',
                'motivo': 'La IA encontró la empresa en el documento.'}
    return {'tipo': 'particular', 'nombre': emp.PARTICULARES, 'fuente': 'reglas',
            'motivo': 'No dice que sea de una empresa: es de una persona.'}


def crear_lote(archivos, usuario_id=None):
    """Guarda y lee los archivos subidos. Devuelve (lote_id, errores)."""
    from services import servicio_tecnico_ia as st_ia
    archivos = [a for a in (archivos or []) if a and getattr(a, 'filename', '')]
    if not archivos:
        raise ErrorServicio('Escoge al menos un archivo.')
    if len(archivos) > ARCHIVOS_MAX:
        raise ErrorServicio(f'Sube máximo {ARCHIVOS_MAX} archivos a la vez.')
    leidos, errores = [], []
    for archivo in archivos:
        nombre = st.nombre_archivo(archivo.filename)
        ext = nombre.rsplit('.', 1)[-1].lower() if '.' in nombre else ''
        datos = archivo.read(ARCHIVO_MAX_BYTES + 1)
        if ext == 'xls':
            errores.append(f'{nombre}: es un Excel viejo (.xls). Ábrelo y guárdalo como «Libro de Excel (.xlsx)».')
        elif ext not in EXTENSIONES:
            errores.append(f'{nombre}: formato no admitido (usa Excel .xlsx, CSV, PDF o .txt).')
        elif not datos:
            errores.append(f'{nombre}: el archivo llegó vacío.')
        elif len(datos) > ARCHIVO_MAX_BYTES:
            errores.append(f'{nombre}: pesa más de 20 MB.')
        elif ext == 'pdf' and not datos.startswith(b'%PDF-'):
            errores.append(f'{nombre}: no es un PDF válido.')
        elif ext in ('xlsx', 'xlsm') and not datos.startswith(b'PK\x03\x04'):
            errores.append(f'{nombre}: no es un Excel válido.')
        else:
            leidos.append((nombre, ext, datos))
    if not leidos:
        raise ErrorServicio(' '.join(errores) or 'Escoge al menos un archivo.')
    if not st.asegurar_tablas():
        raise ErrorServicio('El módulo de Servicio Técnico está pendiente de actualizar en esta tienda.')
    ia_ok = st_ia.estado()[0]
    negocio = contacto_negocio()
    with get_db_cursor(dict_cursor=True) as cur:
        cur.execute("INSERT INTO st_importaciones (estado, creado_por) VALUES ('leyendo', %s) RETURNING id",
                    (usuario_id,))
        lote_id = cur.fetchone()['id']
        info, filas = [], []
        for nombre, ext, datos in leidos:
            clase = EXTENSIONES[ext]
            cur.execute("""SELECT d.importacion_id, i.estado, i.creado_en FROM st_documentos d
                           JOIN st_importaciones i ON i.id = d.importacion_id
                           WHERE d.huella = %s AND i.estado IN ('revision', 'importada') ORDER BY d.id DESC LIMIT 1""",
                        (hashlib.sha256(datos).hexdigest(),))
            previo = cur.fetchone()
            try:
                partes, error = partes_de_archivo(clase, datos), None
            except ErrorServicio as exc:
                partes, error = [], str(exc)
            texto = '\n\n'.join(p.get('texto') or '' for p in partes if p['tipo'] in ('texto', 'ficha')) or None
            doc_id = st.guardar_documento(cur, nombre, st.TIPOS_DOCUMENTO.get(ext, ('text/plain',))[0], datos,
                                          texto=texto if clase in ('pdf', 'texto') else None,
                                          importacion_id=lote_id, usuario_id=usuario_id)
            archivo = {'id': doc_id, 'nombre': nombre, 'clase': clase, 'error': error, 'partes': [],
                       'repetido': ({'lote': previo['importacion_id'], 'estado': previo['estado'],
                                     'fecha': previo['creado_en'].strftime('%d/%m/%Y')} if previo else None)}
            filas_archivo = []
            for i, parte in enumerate(partes):
                filas_parte = filas_de_parte(parte, archivo, i, negocio)
                necesita = _necesita_ia(parte, filas_parte)
                archivo['partes'].append({
                    'i': i, 'tipo': parte['tipo'], 'titulo': parte['titulo'],
                    'ia': ('pendiente' if ia_ok else 'sin_ia') if necesita else 'no_aplica',
                    'intentos': 0, 'aviso': None, 'mapeo_ia': {},
                    'sin_mapear': [parte['titulos'][j] or f'Columna {_letra(int(j))}'
                                   for j in _columnas_de_tabla(parte)[1]] if parte['tipo'] == 'tabla' else [],
                    'filas_leidas': len(parte.get('filas') or [])})
                filas_archivo.extend(filas_parte)
            archivo['empresa'] = _paso_empresa(partes, filas_archivo, negocio['nombre'], ia_ok)
            filas.extend(filas_archivo)
            info.append(archivo)
        equipos = sum(1 for f in filas if f['clase'] == 'equipo')
        if equipos > EQUIPOS_MAX:
            raise ErrorServicio(f'Los archivos traen {equipos} equipos y el máximo por vez es {EQUIPOS_MAX}. '
                                'Divídelos en varios archivos y súbelos por partes.')
        _numerar(filas)
        pendiente = any(p['ia'] == 'pendiente' for a in info for p in a['partes']) \
            or any((a.get('empresa') or {}).get('ia') == 'pendiente' for a in info)
        _guardar(cur, lote_id, archivos=info, filas=filas, estado='leyendo' if pendiente else 'revision')
    return lote_id, errores


def _contenido(doc_id):
    with get_db_cursor() as cur:
        cur.execute('SELECT contenido FROM st_documentos WHERE id = %s', (doc_id,))
        fila = cur.fetchone()
    return bytes(fila[0]) if fila else None


def _ejemplos_seguros(parte, j, nombres):
    """Hasta 4 valores de la columna sin datos personales (para la IA)."""
    salida = []
    for _, fila in parte['filas']:
        if j < len(fila) and _celda(fila[j]):
            valor = lector.sin_datos_personales(_celda(fila[j])[:80], nombres)
            if _parece_nombre(valor):
                valor = '[nombre]'
            if valor not in salida:
                salida.append(valor)
        if len(salida) >= 4:
            break
    return salida


_PALABRAS_TECNICAS = None


def _parece_nombre(valor):
    """«Juan Pérez», «CARLOS GOMEZ»: 2 a 5 palabras de solo letras que no son
    marcas ni palabras técnicas."""
    global _PALABRAS_TECNICAS
    if _PALABRAS_TECNICAS is None:
        _PALABRAS_TECNICAS = ({_norm(m) for m in MARCAS} | set(_EXACTOS) | {
            'windows', 'pro', 'home', 'intel', 'core', 'amd', 'ryzen', 'negro', 'blanco', 'gris', 'azul', 'rojo',
            'plata', 'plateado', 'dorado', 'verde', 'bueno', 'buena', 'malo', 'mala', 'regular', 'nuevo', 'usado',
            'si', 'no', 'preventivo', 'correctivo', 'limpieza', 'general', 'cambio', 'pantalla', 'bateria',
            'teclado', 'disco', 'estado', 'funciona', 'funcional', 'equipo', 'oficina', 'bodega', 'recepcion',
            'gerencia', 'contabilidad', 'ventas', 'sistemas', 'soporte', 'galaxy', 'iphone', 'thinkpad', 'ideapad'})
    palabras = _norm(valor).split()
    if not 2 <= len(palabras) <= 5 or re.search(r'\d', valor) or not re.fullmatch(r"[A-Za-zÁÉÍÓÚÑáéíóúñü' .]+", valor):
        return False
    return not any(p in _PALABRAS_TECNICAS for p in palabras)


def leer_siguiente(lote_id):
    """Lee con IA la siguiente parte pendiente del lote (una por llamada).
    Devuelve {'listo', 'pendientes', 'total', 'archivo', 'parte', 'estado',
    'aviso', 'esperar'}."""
    from services import servicio_tecnico_ia as st_ia
    ahora = datetime.now(timezone.utc)
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote:
            raise ErrorServicio('La importación no existe.')
        if lote['estado'] != 'leyendo':
            return {'listo': True}
        archivos = lote['archivos']
        objetivo, en_curso = None, 0
        for a in archivos:
            for p in a['partes']:
                viejo = p['ia'] == 'leyendo' and (
                    not p.get('desde') or (ahora - datetime.fromisoformat(p['desde'])).total_seconds() > LEYENDO_SEGUNDOS)
                if (p['ia'] == 'pendiente' or viejo) and int(p.get('intentos') or 0) >= INTENTOS_MAX + 2:
                    # Se cayó varias veces a mitad de la lectura: queda lo de las reglas.
                    p.update(ia='fallo', aviso='La IA no pudo leer esta parte: quedó lo que leyeron las reglas.')
                    p.pop('desde', None)
                elif objetivo is None and (p['ia'] == 'pendiente' or viejo):
                    objetivo = (a, p)
                elif p['ia'] == 'leyendo' and not viejo:
                    en_curso += 1
        cuenta = _conteo_lectura(archivos)
        archivo_empresa = None
        if objetivo is None:
            if en_curso:
                _guardar(cur, lote_id, archivos=archivos)
                return {'listo': False, 'esperar': 8, **cuenta}
            # Leídas todas las partes: ¿de qué empresa es cada archivo? (uno por llamada)
            siguiente = _siguiente_empresa(archivos, ahora)
            if siguiente is None:
                _guardar(cur, lote_id, archivos=archivos, estado='revision')
                return {'listo': True, **_conteo_lectura(archivos)}
            if siguiente == 'esperar':
                _guardar(cur, lote_id, archivos=archivos)
                return {'listo': False, 'esperar': 8, **_conteo_lectura(archivos)}
            ea = siguiente['empresa']
            ea.update(ia='leyendo', desde=ahora.isoformat(), intentos=int(ea.get('intentos') or 0) + 1)
            filas_archivo = [f for f in lote['filas'] if f.get('archivo_id') == siguiente['id']]
            _guardar(cur, lote_id, archivos=archivos)
            archivo_empresa = dict(siguiente)
        else:
            a, p = objetivo
            p.update(ia='leyendo', desde=ahora.isoformat(), intentos=int(p.get('intentos') or 0) + 1)
            _guardar(cur, lote_id, archivos=archivos)
            archivo, indice = dict(a), p['i']
    if archivo_empresa is not None:
        return _leer_empresa(lote_id, archivo_empresa, filas_archivo)

    negocio = contacto_negocio()
    resultado = {'ok': False, 'aviso': None, 'reintentar': False, 'ia': None, 'mapeo': None}
    parte = None
    try:
        datos = _contenido(archivo['id'])
        partes = partes_de_archivo(archivo['clase'], datos) if datos else []
        parte = partes[indice] if indice < len(partes) else None
    except ErrorServicio as exc:
        resultado['aviso'] = str(exc)
    if parte is not None and parte['tipo'] == 'texto':
        filas_reglas = _filas_de_texto(parte, negocio)
        tipo_sugerido = filas_reglas[0]['equipo']['tipo'] if len(filas_reglas) == 1 else None
        ia = st_ia.interpretar_documento(sin_secretos(parte['texto']), tipo_sugerido=tipo_sugerido,
                                         nombres=_nombres_para_tapar(parte, filas_reglas))
        resultado.update(ok=ia['ia'], aviso=ia['aviso'], reintentar=ia['reintentar'], ia=ia)
    elif parte is not None and parte['tipo'] == 'tabla':
        _, sin_mapear = _columnas_de_tabla(parte)
        nombres = tuple({_celda(f[j]) for _, f in parte['filas'] for j, d in _columnas_de_tabla(parte)[0].items()
                         if d in ('cliente_nombre', 'cliente_empresa') and j < len(f) and _celda(f[j])})[:400]
        columnas = [{'col': j, 'encabezado': lector.sin_datos_personales(parte['titulos'].get(j) or '', nombres),
                     'ejemplos': _ejemplos_seguros(parte, j, nombres)} for j in sin_mapear]
        destinos = {k: v for k, v in DESTINOS.items() if k not in ('cliente_documento',)}
        mapeo = st_ia.mapear_columnas(columnas, destinos)
        resultado.update(ok=mapeo['ia'], aviso=mapeo['aviso'], reintentar=mapeo['reintentar'], mapeo=mapeo['mapeo'])
    elif parte is not None:
        resultado['ok'] = True

    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote or lote['estado'] != 'leyendo':
            return {'listo': True}
        archivos, filas = lote['archivos'], lote['filas']
        a = next(x for x in archivos if x['id'] == archivo['id'])
        p = next(x for x in a['partes'] if x['i'] == indice)
        esperar = 0
        if resultado['reintentar'] and int(p.get('intentos') or 0) < INTENTOS_MAX:
            p.update(ia='pendiente', aviso='La IA se está preparando; se vuelve a intentar.')
            esperar = 20
        elif resultado['ok'] and parte is not None:
            p.update(ia='hecha', aviso=resultado['aviso'])
            if resultado['mapeo']:
                p['mapeo_ia'] = {str(k): v for k, v in resultado['mapeo'].items()}
            propias = [f for f in filas if f.get('archivo_id') == a['id'] and f.get('parte') == indice]
            if not any(f.get('editada') for f in propias):
                nuevas = filas_de_parte(parte, a, indice, negocio, ia=resultado['ia'],
                                        mapeo_ia={int(k): v for k, v in (p.get('mapeo_ia') or {}).items()})
                total = sum(1 for f in filas if f['clase'] == 'equipo') - sum(1 for f in propias if f['clase'] == 'equipo')
                if total + sum(1 for f in nuevas if f['clase'] == 'equipo') <= EQUIPOS_MAX:
                    filas = _reemplazar(filas, propias, nuevas)
        else:
            p.update(ia='fallo', aviso=resultado['aviso'] or 'La IA no pudo leer esta parte: quedó lo que leyeron las reglas.')
        p.pop('desde', None)
        _guardar(cur, lote_id, archivos=archivos, filas=filas)
        cuenta = _conteo_lectura(archivos)
    return {'listo': False, 'archivo': archivo['nombre'], 'parte': p['titulo'], 'estado': p['ia'],
            'aviso': p.get('aviso'), 'esperar': esperar, **cuenta}


def _siguiente_empresa(archivos, ahora):
    """Archivo cuyo paso «¿de qué empresa es?» toca ahora; 'esperar' si hay uno
    en curso; None si no queda ninguno."""
    objetivo, en_curso = None, False
    for a in archivos:
        ea = a.get('empresa')
        if not ea:
            continue
        viejo = ea.get('ia') == 'leyendo' and (
            not ea.get('desde') or (ahora - datetime.fromisoformat(ea['desde'])).total_seconds() > LEYENDO_SEGUNDOS)
        if (ea.get('ia') == 'pendiente' or viejo) and int(ea.get('intentos') or 0) >= INTENTOS_MAX + 2:
            ea.update(ia='fallo', aviso='La IA no pudo decir de qué empresa es: quedó lo que dicen las reglas.')
            ea.pop('desde', None)
        elif objetivo is None and (ea.get('ia') == 'pendiente' or viejo):
            objetivo = a
        elif ea.get('ia') == 'leyendo' and not viejo:
            en_curso = True
    if objetivo is not None:
        return objetivo
    return 'esperar' if en_curso else None


def _leer_empresa(lote_id, archivo, filas):
    """Paso de la lectura: la IA local dice de qué empresa es el archivo y si sus
    nombres dudosos son de empresa o de persona (servicio_tecnico_ia.identificar_empresa)."""
    from services import servicio_tecnico_ia as st_ia
    negocio = contacto_negocio()
    dudosos, personas, decidida = _nombres_de_archivo(filas, negocio['nombre'])
    decidida = decidida or bool((archivo.get('empresa') or {}).get('etiqueta'))
    resultado, hace_falta = {'ia': False, 'aviso': None, 'reintentar': False}, True
    encabezado = ''
    try:
        datos = _contenido(archivo['id'])
        encabezado = _encabezado(partes_de_archivo(archivo['clase'], datos) if datos else [], personas)
    except ErrorServicio as exc:
        resultado['aviso'] = str(exc)
    if not dudosos and (decidida or not encabezado):
        hace_falta = False                    # la lectura con IA ya dejó cada fila con su empresa
    elif encabezado or dudosos:
        resultado = st_ia.identificar_empresa(encabezado, dudosos, negocio=negocio['nombre'],
                                              archivo=archivo['nombre'])
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote or lote['estado'] != 'leyendo':
            return {'listo': True}
        archivos = lote['archivos']
        a = next(x for x in archivos if x['id'] == archivo['id'])
        ea, esperar = a['empresa'], 0
        if not hace_falta:
            ea.update(ia='no_aplica', aviso=None)
        elif resultado.get('reintentar') and int(ea.get('intentos') or 0) < INTENTOS_MAX:
            ea.update(ia='pendiente', aviso='La IA se está preparando; se vuelve a intentar.')
            esperar = 20
        elif resultado.get('ia'):
            ea.update(ia='hecha', nombre=resultado.get('empresa'), clasificacion=resultado.get('clasificacion') or {},
                      aviso=None)
        else:
            ea.update(ia='fallo', aviso=resultado.get('aviso')
                      or 'La IA no pudo decir de qué empresa es: quedó lo que dicen las reglas.')
        ea.pop('desde', None)
        _guardar(cur, lote_id, archivos=archivos)
        cuenta = _conteo_lectura(archivos)
    return {'listo': False, 'archivo': archivo['nombre'], 'parte': 'Empresa', 'estado': ea['ia'],
            'aviso': ea.get('aviso'), 'esperar': esperar, **cuenta}


def _conteo_lectura(archivos):
    pasos = [p['ia'] for a in archivos for p in a['partes']] \
        + [a['empresa'].get('ia') for a in archivos if a.get('empresa')]
    pasos = [x for x in pasos if x not in ('no_aplica', 'sin_ia', None)]
    return {'total': len(pasos), 'pendientes': sum(1 for x in pasos if x in ('pendiente', 'leyendo')),
            'hechas': sum(1 for x in pasos if x == 'hecha'), 'fallos': sum(1 for x in pasos if x == 'fallo')}


def _reemplazar(filas, viejas, nuevas):
    """Cambia las filas de una parte por las nuevas conservando sus números."""
    numeros = [f['n'] for f in viejas]
    siguiente = max([f['n'] for f in filas] or [0]) + 1
    for k, fila in enumerate(nuevas):
        if k < len(numeros):
            fila['n'] = numeros[k]
        else:
            fila['n'] = siguiente
            siguiente += 1
    quitados = set(numeros) - {f['n'] for f in nuevas}
    ids = {id(f) for f in viejas}
    salida, puesto = [], False
    for fila in filas:
        if id(fila) in ids:
            if not puesto:
                salida.extend(nuevas)
                puesto = True
            continue
        if fila.get('unir_con') in quitados:
            fila['unir_con'] = None
        if fila.get('adjuntar_a') in quitados:
            fila['adjuntar_a'] = None
        salida.append(fila)
    return salida


def terminar_lectura(lote_id):
    """«Seguir sin esperar a la IA»: lo pendiente queda como lo leyeron las reglas."""
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote:
            raise ErrorServicio('La importación no existe.')
        if lote['estado'] != 'leyendo':
            return
        for a in lote['archivos']:
            for p in a['partes']:
                if p['ia'] in ('pendiente', 'leyendo'):
                    p.update(ia='omitida', aviso='Se siguió sin esperar a la IA: quedó lo que leyeron las reglas.')
                    p.pop('desde', None)
            ea = a.get('empresa') or {}
            if ea.get('ia') in ('pendiente', 'leyendo'):
                ea.update(ia='omitida', aviso='Se siguió sin esperar a la IA: la empresa la dicen las reglas.')
                ea.pop('desde', None)
        _guardar(cur, lote_id, archivos=lote['archivos'], estado='revision')


def volver_a_leer(lote_id):
    """Desde la revisión: leer con IA lo que se leyó sin ella (si no se ha
    corregido nada a mano en esas filas)."""
    from services import servicio_tecnico_ia as st_ia
    if not st_ia.estado()[0]:
        raise ErrorServicio('La IA no está disponible en este momento.')
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote or lote['estado'] != 'revision':
            raise ErrorServicio('Esta importación ya no se puede volver a leer.')
        editadas = {(f.get('archivo_id'), f.get('parte')) for f in lote['filas'] if f.get('editada')}
        cambiadas = 0
        for a in lote['archivos']:
            for p in a['partes']:
                if p['ia'] in ('sin_ia', 'fallo', 'omitida') and (a['id'], p['i']) not in editadas:
                    if p['tipo'] == 'texto' or p.get('sin_mapear'):
                        p.update(ia='pendiente', intentos=0, aviso=None)
                        cambiadas += 1
            ea = a.get('empresa') or {}
            if ea.get('ia') in ('sin_ia', 'fallo', 'omitida'):        # la carpeta escogida a mano siempre gana
                ea.update(ia='pendiente', intentos=0, aviso=None)
                cambiadas += 1
        if not cambiadas:
            raise ErrorServicio('No hay nada pendiente de leer con IA (o ya lo corregiste a mano).')
        _guardar(cur, lote_id, archivos=lote['archivos'], estado='leyendo')


def guardar_ajustes(lote_id, datos):
    """Cliente para los equipos que no dicen de quién son, plan por defecto y
    qué hacer con los mantenimientos vencidos."""
    ajustes = {}
    contacto_id = datos.get('cliente_id')
    if contacto_id:
        cliente = st.obtener_cliente(int(contacto_id))
        if not cliente:
            raise ErrorServicio('El cliente escogido no existe.')
        ajustes['cliente_id'] = cliente['id']
    elif (datos.get('cliente_nombre') or '').strip():
        ajustes['cliente_nuevo'] = {'nombre': ' '.join(datos['cliente_nombre'].split())[:150],
                                    'telefono': _telefono(datos.get('cliente_telefono')),
                                    'email': _correo(datos.get('cliente_email'))}
    meses = datos.get('plan_meses')
    if meses not in (None, ''):
        try:
            meses = int(meses)
        except (TypeError, ValueError):
            raise ErrorServicio('La frecuencia debe ser un número de meses.')
        if not 0 <= meses <= 36:
            raise ErrorServicio('La frecuencia va de 1 a 36 meses.')
        ajustes['plan_meses'] = meses or None
    ajustes['vencidos'] = datos.get('vencidos') if datos.get('vencidos') in dict(VENCIDOS) else 'siguiente'
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote or lote['estado'] != 'revision':
            raise ErrorServicio('Esta importación ya no se puede cambiar.')
        _guardar(cur, lote_id, ajustes=ajustes)


def editar_fila(lote_id, n, datos):
    """Corrige una fila desde la revisión: incluirla o no, tipo, marca, modelo,
    serial, IMEI, cliente; unirla con otra o escoger a qué equipo va un PDF."""
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote or lote['estado'] != 'revision':
            raise ErrorServicio('Esta importación ya no se puede cambiar.')
        filas = lote['filas']
        fila = next((f for f in filas if f['n'] == n), None)
        if not fila:
            raise ErrorServicio('Esa fila no existe.')
        numeros = {f['n'] for f in filas if f['clase'] == 'equipo'}
        accion = datos.get('accion')
        if accion in ('quitar', 'incluir'):
            fila['incluir'] = accion == 'incluir'
        elif accion == 'unir':
            otro = int(datos.get('con') or 0)
            if otro not in numeros or otro == n:
                raise ErrorServicio('Escoge con cuál fila unirla.')
            fila['unir_con'] = otro
        elif accion == 'separar':
            fila['unir_con'] = None
        elif accion == 'otro_equipo':                # no es el equipo ya registrado que se le parece
            fila['es_otro_equipo'] = True
            fila.pop('usar_equipo', None)
        elif accion == 'mismo_registrado':
            fila.pop('es_otro_equipo', None)
        elif accion == 'usar_equipo':                # «ya estaba registrado: es este» (se actualiza ese)
            eid = str(datos.get('equipo_id') or '')
            if not eid.isdigit():
                raise ErrorServicio('Escoge cuál de los equipos registrados es.')
            cur.execute('SELECT 1 FROM st_equipos WHERE id = %s AND activo', (int(eid),))
            if not cur.fetchone():
                raise ErrorServicio('Ese equipo ya no está registrado.')
            fila['usar_equipo'] = int(eid)
            fila.pop('es_otro_equipo', None)
        elif accion == 'no_usar_equipo':
            fila.pop('usar_equipo', None)
        elif accion == 'adjuntar':
            otro = int(datos.get('con') or 0)
            fila['adjuntar_a'] = otro if otro in numeros else None
        elif accion == 'corregir':
            e, c = fila['equipo'], fila['cliente']
            tipo = (datos.get('tipo') or '').strip()
            if tipo:
                if not tipos.es_valido(tipo):
                    raise ErrorServicio('Escoge el tipo de equipo.')
                e['tipo'], fila['tipo_adivinado'] = tipo, False
            for campo in ('marca', 'modelo', 'serial'):
                if campo in datos:
                    valor = ' '.join(str(datos.get(campo) or '').split())[:st._LARGOS[campo]]
                    if valor:
                        e[campo] = valor
                    else:
                        e.pop(campo, None)
            if 'imei' in datos:
                digitos = re.sub(r'\D', '', datos.get('imei') or '')
                if digitos and not lector.imei_luhn(digitos):
                    raise ErrorServicio('El IMEI no es válido: debe tener 15 dígitos y su dígito de control.')
                if digitos:
                    e['imei'] = digitos
                else:
                    e.pop('imei', None)
            for campo in ('nombre', 'telefono', 'email'):
                if campo in datos:
                    valor = ' '.join(str(datos.get(campo) or '').split())
                    if campo == 'telefono':
                        valor = _telefono(valor) or (valor or None)
                    elif campo == 'email':
                        if valor and not _correo(valor):
                            raise ErrorServicio('El correo no es válido.')
                        valor = _correo(valor)
                    c[campo] = valor or None
            contacto = datos.get('usar_contacto')
            if contacto == 'nuevo':
                fila['usar_contacto'] = 'nuevo'
            elif contacto and str(contacto).isdigit():
                if not st.obtener_cliente(int(contacto)):
                    raise ErrorServicio('El cliente escogido no existe.')
                fila['usar_contacto'] = int(contacto)
            elif 'usar_contacto' in datos:
                fila['usar_contacto'] = None
            eleccion = datos.get('carpeta')
            if eleccion == 'auto':
                fila.pop('carpeta_manual', None)
            elif eleccion == 'particulares':
                fila['carpeta_manual'] = {'tipo': 'particular'}
            elif eleccion == 'empresa':
                nombre = emp.nombre_visible(datos.get('carpeta_nombre') or '')
                if not emp.clave(nombre):
                    raise ErrorServicio('Escribe el nombre de la empresa, o escoge «Particulares».')
                fila['carpeta_manual'] = {'tipo': 'empresa', 'nombre': nombre}
            _cerrar(fila)
        else:
            raise ErrorServicio('Acción no válida.')
        fila['editada'] = True
        _guardar(cur, lote_id, filas=filas)


def descartar(lote_id, usuario_id=None):
    """El lote no se importa: sus archivos quedan retirados (no se borran)."""
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote or lote['estado'] not in ('leyendo', 'revision'):
            raise ErrorServicio('Esta importación ya no se puede descartar.')
        cur.execute('UPDATE st_documentos SET activo = FALSE WHERE importacion_id = %s AND equipo_id IS NULL',
                    (lote_id,))
        _guardar(cur, lote_id, estado='descartada', resultado={'descartado_por': usuario_id})


# ── Revisión (lo que va a pasar) ────────────────────────────────
def _sin_tildes_sql(columna='nombre'):
    return (f"regexp_replace(lower(translate(COALESCE({columna}, ''), 'ÁÉÍÓÚÜÑáéíóúüñ', 'AEIOUUNaeiouun')), "
            "'\\s+', ' ', 'g')")


def _contactos_existentes(cur, correos, telefonos, nombres):
    por_correo, por_tel, por_nombre = {}, {}, {}
    if correos:
        cur.execute("""SELECT id, nombre, email, telefono, whatsapp FROM crm_contactos
                       WHERE COALESCE(activo, TRUE) AND lower(email) = ANY(%s) ORDER BY id""", (sorted(correos),))
        for r in cur.fetchall():
            por_correo.setdefault(r['email'].lower(), dict(r))
    if telefonos:
        cur.execute("""SELECT id, nombre, email, telefono, whatsapp,
                              right(regexp_replace(COALESCE(telefono, ''), '\\D', '', 'g'), 10) AS t1,
                              right(regexp_replace(COALESCE(whatsapp, ''), '\\D', '', 'g'), 10) AS t2
                       FROM crm_contactos
                       WHERE COALESCE(activo, TRUE)
                         AND (right(regexp_replace(COALESCE(telefono, ''), '\\D', '', 'g'), 10) = ANY(%s)
                              OR right(regexp_replace(COALESCE(whatsapp, ''), '\\D', '', 'g'), 10) = ANY(%s))
                       ORDER BY id""", (sorted(telefonos), sorted(telefonos)))
        for r in cur.fetchall():
            for t in (r['t1'], r['t2']):
                if t in telefonos:
                    por_tel.setdefault(t, {k: r[k] for k in ('id', 'nombre', 'email', 'telefono', 'whatsapp')})
    if nombres:
        # Por su nombre o por su campo «empresa» (una empresa puede ser una sola palabra: «ACME»).
        cur.execute(f"""SELECT id, nombre, email, telefono, whatsapp, {_sin_tildes_sql()} AS clave,
                               {_sin_tildes_sql('empresa')} AS clave_empresa
                        FROM crm_contactos
                        WHERE COALESCE(activo, TRUE)
                          AND ({_sin_tildes_sql()} = ANY(%s) OR {_sin_tildes_sql('empresa')} = ANY(%s))""",
                    (sorted(nombres), sorted(nombres)))
        vistos = {}
        for r in cur.fetchall():
            datos = {k: r[k] for k in ('id', 'nombre', 'email', 'telefono', 'whatsapp')}
            for clave in {r['clave'], r['clave_empresa']} & set(nombres):
                if datos['id'] not in [x['id'] for x in vistos.get(clave, [])]:
                    vistos.setdefault(clave, []).append(datos)
        por_nombre = {k: v[0] for k, v in vistos.items() if len(v) == 1}
    return por_correo, por_tel, por_nombre


def _palabras_de_nombre(nombre):
    return {w for w in _clave_nombre(nombre).split() if len(w) >= 3}


def _clave_nombre(nombre):
    t = unicodedata.normalize('NFD', str(nombre or '').lower())
    t = ''.join(ch for ch in t if unicodedata.category(ch) != 'Mn')
    return ' '.join(t.split())


def _equipos_existentes(cur, seriales, imeis, codigos=()):
    """Equipos activos por serial, IMEI o código de ficha técnica (el código
    de la ficha del taller identifica al equipo aunque no tenga serial)."""
    por_clave = {}
    if not seriales and not imeis and not codigos:
        return por_clave
    cur.execute("""SELECT e.id, e.tipo, e.marca, e.modelo, e.serial, e.imei, e.crm_contacto_id, e.mant_cada_meses,
                          e.mant_proximo, e.ficha_codigo, e.creado_en, c.nombre AS cliente_nombre,
                          x.id AS empresa_id, x.nombre AS empresa_nombre
                   FROM st_equipos e LEFT JOIN crm_contactos c ON c.id = e.crm_contacto_id
                   LEFT JOIN st_empresas x ON x.id = e.empresa_id AND x.activo
                   WHERE e.activo AND (regexp_replace(lower(COALESCE(e.serial, '')), '[\\s./-]', '', 'g') = ANY(%s)
                                       OR e.imei = ANY(%s) OR lower(e.ficha_codigo) = ANY(%s))
                   ORDER BY e.creado_en, e.id""", (sorted(seriales), sorted(imeis), sorted(codigos)))
    for r in cur.fetchall():
        r = dict(r)
        r['descripcion'] = st.descripcion_equipo(r)
        for clave in claves_de(r) | ({'ficha:' + r['ficha_codigo'].lower()} if r['ficha_codigo'] else set()):
            por_clave.setdefault(clave, r)
    return por_clave


_RELLENO_MODELO = {'de', 'del', 'la', 'el', 'con', 'para', 'y', 'impresora', 'multifuncional', 'computador',
                   'computadora', 'portatil', 'equipo', 'pc', 'touch', 'tactil'}


def _modelo_clave(marca, modelo):
    """El modelo para comparar: sin la marca, sin lo que va entre paréntesis ni
    palabras de relleno («IdeaCentre B320 (Touch)» = «Lenovo IdeaCentre B320»).
    Vacío si no trae un número de modelo (un «Galaxy Tab» solo no identifica nada)."""
    marca_p = set(_norm(marca).split()) | set(_norm(_marca_canonica(marca)).split()) if marca else set()
    palabras = [p for p in _norm(re.sub(r'\([^)]*\)', ' ', str(modelo or ''))).split()
                if p not in marca_p and p not in _RELLENO_MODELO]
    clave = ' '.join(palabras)
    return clave if re.search(r'\d', clave) else ''


# ── El mismo equipo aunque el archivo no traiga su serial ───────
# Lo que nombra a un equipo dentro de su cliente (en sus notas: «Área
# correspondiente: Contabilidad», «Asignado a: …», «Nombre de usuario: …»).
_ETIQ_ASIGNADO = re.compile(
    r'(?im)^[ \t]*(?:asignado(?:\s+a)?|[aá]rea(?:\s+(?:correspondiente|asignada|responsable))?|'
    r'usuario(?:\s+(?:asignado|responsable|windows))?|nombre\s+de\s+usuario|nombre\s+(?:del\s+)?equipo|'
    r'hostname|nombre\s+de\s+host|ubicaci[oó]n|dependencia|oficina|responsable|funcionario)[ \t]*[:\-][ \t]*(.+)$')
_NOMBRES_VACIOS = {'del', 'las', 'los', 'con', 'para', 'equipo', 'equipos', 'area', 'computador', 'usuario', 'user',
                   'admin', 'administrador', 'hoja', 'sheet', 'ficha', 'tecnica', 'ninguno', 'sin', 'dato', 'historico'}
_GRUPO_TIPO = {'computador': 'pc', 'portatil': 'pc'}
_COLS_PARECIDOS = """e.id, e.tipo, e.marca, e.modelo, e.serial, e.imei, e.crm_contacto_id, e.mant_cada_meses,
                     e.mant_proximo, e.ficha_codigo, e.ficha, e.notas, e.creado_en, c.nombre AS cliente_nombre,
                     x.id AS empresa_id, x.nombre AS empresa_nombre"""


def _nombres_de(notas, ficha, cliente_nombre='', hoja=None):
    """Lo que nombra al equipo dentro de su cliente: la persona o el área a la que
    está asignado, su usuario y la hoja del libro donde vino (si no es «Hoja1»).
    Sin las palabras del nombre del cliente ni las de relleno."""
    valores = [m.group(1) for m in _ETIQ_ASIGNADO.finditer(notas or '')]
    ficha = ficha or {}
    valores += [ficha.get('asignado_a'), (ficha.get('filas') or {}).get('Nombre de Usuario')]
    if hoja and not re.match(r'(?i)^\s*(hoja|sheet|formato|plantilla|libro)\b', hoja):
        valores.append(hoja)
    vacias = _NOMBRES_VACIOS | set(_norm(cliente_nombre).split())
    salida = set()
    for v in valores:
        palabras = [w for w in _norm(v).split() if len(w) >= 3 and w not in vacias]
        if palabras:
            salida.add(' '.join(palabras))
    return salida


def _parecido_nombre(a, b):
    """2: el mismo nombre (persona, área o usuario); 1: comparten una palabra; 0: no."""
    if not a or not b:
        return 0
    if a & b:
        return 2
    return 1 if {w for x in a for w in x.split()} & {w for x in b for w in x.split()} else 0


def _parecido_modelo(marca_a, modelo_a, marca_b, modelo_b):
    """3: el mismo modelo; 2: parecido (con los mismos números: «OptiPlex AIO 7410»
    y «OptiPlex AIO 7410 65W»); 0: otro; None: no se puede comparar."""
    ka, kb = _modelo_clave(marca_a, modelo_a), _modelo_clave(marca_b, modelo_b)
    if not ka or not kb:
        return None
    if ka == kb:
        return 3
    ta, tb = set(ka.split()), set(kb.split())
    na, nb = {t for t in ta if re.search(r'\d', t)}, {t for t in tb if re.search(r'\d', t)}
    corto, largo = (na, nb) if len(na) <= len(nb) else (nb, na)
    if corto and corto <= largo and len(ta & tb) / len(ta | tb) >= 0.5:
        return 2
    return 0


def _serie_codigo(codigo):
    return re.sub(r'\d+\s*$', '', (codigo or '').strip().lower())


def _es_otro(equipo, codigo, candidato):
    """¿Algo dice que NO es el mismo equipo? Serial, IMEI o número de ficha (de la
    misma serie) distintos, otra marca o un tipo que no se parece."""
    s1 = re.sub(r'[\s\-./]', '', str(equipo.get('serial') or '')).lower()
    s2 = re.sub(r'[\s\-./]', '', str(candidato.get('serial') or '')).lower()
    if len(s1) >= 4 and len(s2) >= 4 and s1 != s2:
        return True
    i1, i2 = re.sub(r'\D', '', str(equipo.get('imei') or '')), re.sub(r'\D', '', str(candidato.get('imei') or ''))
    if i1 and i2 and i1 != i2:
        return True
    c2 = candidato.get('ficha_codigo')
    if codigo and c2 and codigo.lower() != c2.lower() and _serie_codigo(codigo) == _serie_codigo(c2):
        return True
    m1, m2 = _norm(_marca_canonica(equipo.get('marca'))), _norm(_marca_canonica(candidato.get('marca')))
    if m1 and m2 and m1 != m2:
        return True
    t1, t2 = equipo.get('tipo'), candidato.get('tipo')
    return bool(t1 and t2 and 'otro' not in (t1, t2) and _GRUPO_TIPO.get(t1, t1) != _GRUPO_TIPO.get(t2, t2))


def _equipos_por_id_o_cliente(cur, ids=(), contactos=()):
    """Equipos activos por id o de esos clientes, del registrado primero al último."""
    if not ids and not contactos:
        return []
    cur.execute(f"""SELECT {_COLS_PARECIDOS}
                    FROM st_equipos e LEFT JOIN crm_contactos c ON c.id = e.crm_contacto_id
                    LEFT JOIN st_empresas x ON x.id = e.empresa_id AND x.activo
                    WHERE e.activo AND (e.id = ANY(%s) OR e.crm_contacto_id = ANY(%s))
                    ORDER BY e.creado_en, e.id""", (sorted(ids), sorted(contactos)))
    salida = []
    for r in cur.fetchall():
        r = dict(r)
        r['descripcion'] = st.descripcion_equipo(r)
        salida.append(r)
    return salida


def _datos_del_grupo(filas, cliente_nombre):
    """Lo que dice del equipo el grupo de filas que se juntan (la tabla de control
    y su hoja): tipo, marca y modelo, serial, número de ficha y sus nombres."""
    equipo = {}
    con_modelo = next((f['equipo'] for f in filas if f['equipo'].get('marca') and f['equipo'].get('modelo')), None) \
        or next((f['equipo'] for f in filas if f['equipo'].get('modelo')), None) or {}
    for e in [con_modelo] + [f['equipo'] for f in filas]:
        for k in ('tipo', 'marca', 'modelo', 'serial', 'imei'):
            if e.get(k) and not equipo.get(k):
                equipo[k] = e[k]
    codigo = next((f.get('ficha_codigo') for f in filas if f.get('ficha_codigo')), None)
    nombres = set()
    for f in filas:
        nombres |= _nombres_de('\n'.join(f.get('notas') or []), f.get('ficha'), cliente_nombre, f.get('fotos_hoja'))
    return equipo, codigo, nombres


def _parecidos_del_cliente(cur, equipos, grupos, por_clave, cliente_de, buscadores, defecto):
    """El mismo equipo aunque el archivo no traiga su serial: del mismo cliente y
    parecido —el mismo modelo o, si el modelo no se puede comparar, la misma
    persona, área o usuario— y nada que diga que es otro (serial, IMEI, número de
    ficha, marca o tipo). Si varios se parecen se queda el más parecido y, entre
    iguales, el que se registró primero: se actualiza la ficha más antigua en vez
    de crear otra. Cada equipo registrado sirve a una sola fila. «Es otro equipo»
    lo deja aparte y «ya estaba registrado: es este» lo escoge a mano.

    Si el parecido es un equipo que otra fila del lote ya encontró por su serial o
    número de ficha (p. ej. la tabla de control sin número y la hoja con él), esa
    fila se junta con aquella: es el mismo equipo, no uno nuevo.

    Devuelve ({n: equipo}, {n: [otros parecidos]}, {n: [candidatos de la fila marcada
    «es otro equipo»]}, {n: [equipos del cliente para escoger a mano]}, {n: fila con
    la que se junta})."""
    por_correo, por_tel, por_nombre = buscadores
    fila_n = {f['n']: f for f in equipos}

    def contacto_probable(f):                  # el mismo que escogerá la revisión para la fila
        uso = f.get('usar_contacto')
        if isinstance(uso, int):
            return uso
        if uso == 'nuevo':
            return None
        c = cliente_de[f['n']]
        contacto = (por_correo.get(c.get('email') or '') or por_tel.get(c.get('telefono') or '')
                    or por_tel.get(c.get('whatsapp') or '')
                    or por_nombre.get(_clave_nombre(c.get('nombre') or c.get('empresa'))))
        if contacto:
            return contacto['id']
        if not (c.get('nombre') or c.get('empresa')) and defecto and defecto.get('existente'):
            return defecto['existente']['id']
        return None

    fuertes = {e['id']: r for r, e in por_clave.items() if e}    # ya los encontró una fila por serial o código
    tomados = set(fuertes)
    contacto_de = {r: contacto_probable(fila_n[r]) for r in grupos if not por_clave.get(r)}
    elegidos = {r: fila_n[r]['usar_equipo'] for r in contacto_de if isinstance(fila_n[r].get('usar_equipo'), int)}
    registrados = _equipos_por_id_o_cliente(cur, set(elegidos.values()), {c for c in contacto_de.values() if c})
    por_id = {e['id']: e for e in registrados}
    por_modelo, alternativas, candidatos_de, opciones, juntar = {}, {}, {}, {}, {}
    for r, eid in sorted(elegidos.items()):        # escogido a mano: ese es, si nadie más lo tomó
        if eid in por_id and eid not in tomados:
            por_modelo[r] = dict(por_id[eid], manual=True)
            tomados.add(eid)
    nombres_de = {e['id']: _nombres_de(e.get('notas'), e.get('ficha'), e.get('cliente_nombre')) for e in registrados}
    pares, posibles = [], {}
    for r, contacto in contacto_de.items():
        if not contacto or r in por_modelo:
            continue
        filas = [fila_n[n] for n in sorted(grupos[r])]
        cliente = cliente_de[r].get('nombre') or cliente_de[r].get('empresa') or ''
        equipo, codigo, nombres = _datos_del_grupo(filas, cliente)
        for e in registrados:
            if e['crm_contacto_id'] != contacto or (e['id'] in tomados and e['id'] not in fuertes) \
                    or _es_otro(equipo, codigo, e):
                continue
            if e['id'] not in fuertes:
                posibles.setdefault(r, []).append(e)
            modelo = _parecido_modelo(equipo.get('marca'), equipo.get('modelo'), e.get('marca'), e.get('modelo'))
            nombre = _parecido_nombre(nombres, nombres_de[e['id']])
            if (modelo is not None and modelo < 2) or (modelo is None and nombre < 2):
                continue
            puntos = (modelo or 0) * 10 + nombre * 3 + (1 if equipo.get('tipo') == e.get('tipo') else 0)
            pares.append((-puntos, e['creado_en'], e['id'], r, e))
    pares.sort(key=lambda x: (x[0], x[1], x[2], x[3]))
    juntados = set()
    for _p, _f, eid, r, e in pares:
        if r in por_modelo or r in juntar:
            continue
        if fila_n[r].get('es_otro_equipo'):        # dijiste que es otro: no se junta (se recuerda con cuál)
            if eid not in fuertes and eid not in {x['id'] for x in candidatos_de.get(r, [])}:
                candidatos_de.setdefault(r, []).append(e)
            continue
        if eid in fuertes:                         # otra fila del lote ya es ese equipo: van juntas
            if eid not in juntados:
                juntar[r] = fuertes[eid]
                juntados.add(eid)
            continue
        if eid in tomados:
            continue
        por_modelo[r] = e
        tomados.add(eid)
    for _p, _f, eid, r, e in pares:                # los otros parecidos que nadie tomó: para escoger a mano
        if r in por_modelo and eid not in tomados and eid not in {x['id'] for x in alternativas.get(r, [])}:
            alternativas.setdefault(r, []).append(e)
    for r, lista in posibles.items():              # para escoger a mano cuál era, si no se encontró solo
        libres = [e for e in lista if e['id'] not in tomados]
        if r not in por_modelo and r not in juntar and libres:
            opciones[r] = libres[:30]
    return por_modelo, alternativas, candidatos_de, opciones, juntar


def _plan_final(fila, ajustes, hoy):
    """(cada_meses, próximo, aviso) del plan de mantenimiento del equipo."""
    plan = fila.get('plan') or {}
    cada = plan.get('cada_meses') or ajustes.get('plan_meses') or None
    proximo = _iso(plan.get('proximo'))
    if proximo and proximo > hoy + timedelta(days=3 * 366):
        proximo = None
    from services import servicio_tecnico_seguimiento as seg
    ultimo = max((_iso(h['fecha']) for h in fila['historial'] if h['tipo'] == 'preventivo'), default=None)
    if not proximo and cada:
        proximo = seg._sumar_meses(ultimo or hoy, cada)
    aviso = None
    if proximo and proximo < hoy:
        if ajustes.get('vencidos') == 'hoy':
            aviso = f'El mantenimiento estaba para el {proximo:%d/%m/%Y}: queda para hoy, para llamar al cliente.'
            proximo = hoy
        elif cada:
            vencido = proximo
            while proximo < hoy:
                proximo = seg._sumar_meses(proximo, cada)
            aviso = (f'El mantenimiento estaba para el {vencido:%d/%m/%Y}: se programó el siguiente, '
                     f'{proximo:%d/%m/%Y}.')
        else:
            aviso = (f'El próximo mantenimiento era el {proximo:%d/%m/%Y} (ya pasó) y no dice cada cuánto: '
                     'no se programa.')
            proximo = None
    return cada, proximo, aviso


def _specs(e):
    return ' · '.join(str(e[c]) for c in ('procesador', 'ram', 'almacenamiento', 'sistema_operativo') if e.get(c))


class _Agenda:
    """Clientes nuevos del lote: uno solo por persona aunque aparezca en varias
    filas (mismo correo o teléfono; o mismo nombre sin datos que choquen)."""

    def __init__(self):
        self.personas = []

    def _buscar(self, c):
        correo = c.get('email')
        tels = {t for t in (c.get('telefono'), c.get('whatsapp')) if t}
        nombre = _clave_nombre(c.get('nombre') or c.get('empresa'))
        for p in self.personas:
            d = p['datos']
            if (correo and correo == d.get('email')) or tels & {d.get('telefono'), d.get('whatsapp')} - {None}:
                return p
        for p in self.personas:
            d = p['datos']
            if nombre and nombre == _clave_nombre(d.get('nombre') or d.get('empresa')):
                sus_tels = {t for t in (d.get('telefono'), d.get('whatsapp')) if t}
                choca = (correo and d.get('email') and correo != d['email']) or (tels and sus_tels and not tels & sus_tels)
                if not choca:
                    return p
        return None

    def persona(self, c):
        p = self._buscar(c)
        if p is None:
            p = {'datos': {k: v for k, v in c.items() if v}, 'id': None}
            self.personas.append(p)
        else:
            for k, v in c.items():
                if v and not p['datos'].get(k):
                    p['datos'][k] = v
        return p


def _raices(equipos, claves):
    """Con qué fila se junta cada una (mismo serial o IMEI, o «Unir» a mano).
    Devuelve {n: n_raíz}; la raíz es la primera fila de su grupo."""
    padre = {}

    def raiz(n):
        vistos = set()
        while padre.get(n, n) != n and n not in vistos:
            vistos.add(n)
            n = padre[n]
        return n
    incluidas = {f['n'] for f in equipos if f.get('incluir', True)}
    duenos = {}
    for f in equipos:
        n = f['n']
        if n not in incluidas:
            continue
        padre.setdefault(n, n)
        otra = f.get('unir_con')
        if otra in incluidas and otra != n:
            r = raiz(otra)
            if r != n and raiz(n) == n:
                padre[n] = r
        for k in claves[n]:
            if k in duenos:
                r, propia = raiz(duenos[k]), raiz(n)
                if r != propia:
                    padre[max(r, propia)] = min(r, propia)
            else:
                duenos[k] = n
    return {n: raiz(n) for n in incluidas}


def _actualizaciones_previstas(cur, destino, grupos, fila_n, hoy):
    """{raíz: versión que tiene, versión con la que queda y qué cambia} de los
    equipos que ya estaban, haciendo la importación sin escribir nada."""
    ids = sorted({e['id'] for e in destino.values() if e})
    if not ids:
        return {}
    cur.execute('SELECT * FROM st_equipos WHERE id = ANY(%s)', (ids,))
    completos = {r['id']: dict(r) for r in cur.fetchall()}
    salida = {}
    for raiz, e in destino.items():
        actual = completos.get(e['id']) if e else None
        if not actual:
            continue
        d, version_fija = actual, None
        for n in [raiz] + sorted(x for x in grupos[raiz] if x != raiz):   # como al importar: la raíz primero
            f = fila_n[n]
            nuevo = dict(f['equipo'], notas='\n'.join(f['notas'])[:2000] or None, resumen_ia=f.get('resumen'))
            d2 = _equipo_actualizado(d, nuevo, f, hoy, version_fija)
            v2 = (d2.get('ficha') or {}).get('version')
            if v2 and v2 != (d.get('ficha') or {}).get('version'):
                version_fija = v2
            d = d2
        antes, despues = ver.estado_de(actual), ver.estado_de(d)
        salida[raiz] = {'version_actual': ver.numero(antes), 'version_nueva': ver.numero(despues),
                        'cambios': ver.diferencias(antes, despues)}
    return salida


def _texto_actualizacion(prev):
    """Qué le pasa al equipo que ya estaba: su ficha pasa a una versión nueva y
    la anterior se guarda; se dice qué cambia."""
    if not prev or not prev['cambios']:
        return 'su ficha ya dice esto, así que no cambia (solo se le agrega el historial que le falte).'
    if prev['version_actual'] and prev['version_nueva'] != prev['version_actual']:
        texto = (f"su ficha pasa a la versión {prev['version_nueva']} y la versión {prev['version_actual']} "
                 'queda guardada.')
    elif prev['version_nueva'] and not prev['version_actual']:
        texto = f"se le pone su ficha técnica (versión {prev['version_nueva']}) y lo que tenía antes queda guardado."
    else:
        texto = 'se actualizan sus datos y lo que tenía antes queda guardado.'
    return (f"{texto} Cambia: {ver.resumen_cambios(prev['cambios'])}. Sus fotos, documentos, mantenimientos y "
            'piezas se conservan.')


def _revisar(cur, filas, ajustes, hoy, archivos=()):
    equipos = [f for f in filas if f['clase'] == 'equipo']
    info_empresa = {a['id']: a.get('empresa') or {} for a in archivos or ()}
    negocio_nombre = contacto_negocio()['nombre']
    # Cliente de cada fila: el que dice el documento o, si no dice ninguno y los
    # equipos son de una empresa (p. ej. su inventario de funcionarios), la empresa.
    cliente_de, cliente_es_empresa = {}, set()
    for f in equipos:
        c = f['cliente']
        if not (c.get('nombre') or c.get('empresa')) and not f.get('usar_contacto'):
            carpeta = _carpeta_de_fila(f, info_empresa.get(f.get('archivo_id')), negocio_nombre)
            if carpeta['tipo'] == 'empresa':
                c = {**c, 'empresa': carpeta['nombre']}
                cliente_es_empresa.add(f['n'])
        cliente_de[f['n']] = c
    claves = {f['n']: claves_de(f['equipo']) | ({'ficha:' + f['ficha_codigo'].lower()} if f.get('ficha_codigo')
                                                 else set()) for f in equipos}
    todas = set().union(*claves.values()) if claves else set()
    existentes = _equipos_existentes(cur, {k[7:] for k in todas if k.startswith('serial:')},
                                     {k[5:] for k in todas if k.startswith('imei:')},
                                     {k[6:] for k in todas if k.startswith('ficha:')})
    clientes = list(cliente_de.values())
    correos = {c['email'] for c in clientes if c.get('email')}
    telefonos = {c[k] for c in clientes for k in ('telefono', 'whatsapp') if c.get(k) and len(c[k]) == 10}
    nombres = {_clave_nombre(c.get('nombre') or c.get('empresa')) for c in clientes if c.get('nombre') or c.get('empresa')}
    empresas = {_clave_nombre(c['empresa']) for c in clientes if c.get('empresa') and not c.get('nombre')}
    # Personas: nombre y apellido (un «Juan» solo no basta). Empresas: también de una palabra.
    nombres = {x for x in nombres if len(x.split()) >= 2 or (x in empresas and len(x) >= 4)}
    por_correo, por_tel, por_nombre = _contactos_existentes(cur, correos, telefonos, nombres)
    elegidos = {f['usar_contacto'] for f in equipos if isinstance(f.get('usar_contacto'), int)}
    if ajustes.get('cliente_id'):
        elegidos.add(ajustes['cliente_id'])
    contactos_elegidos = {}
    if elegidos:
        cur.execute("""SELECT id, nombre, email, telefono, whatsapp FROM crm_contactos
                       WHERE id = ANY(%s) AND COALESCE(activo, TRUE)""", (sorted(elegidos),))
        contactos_elegidos = {r['id']: dict(r) for r in cur.fetchall()}
    defecto = None
    if ajustes.get('cliente_id') in contactos_elegidos:
        defecto = {'existente': contactos_elegidos[ajustes['cliente_id']]}
    elif ajustes.get('cliente_nuevo'):
        defecto = {'nuevo': ajustes['cliente_nuevo']}

    codigos = {(f.get('ficha_codigo') or '').lower() for f in equipos if f.get('ficha_codigo')}
    codigo_de = {}
    if codigos:
        cur.execute("""SELECT id, lower(ficha_codigo) AS codigo, tipo, marca, modelo FROM st_equipos
                       WHERE activo AND lower(ficha_codigo) = ANY(%s)""", (sorted(codigos),))
        codigo_de = {r['codigo']: dict(r) for r in cur.fetchall()}
    codigos_vistos = set()
    raices = _raices(equipos, claves)
    grupos = {}
    for n, r in raices.items():
        grupos.setdefault(r, []).append(n)
    # Mismo cliente, marca y modelo y sin serial: puede ser el mismo equipo
    # (solo si vienen de archivos distintos o las dos traen historial).
    posibles = {}
    for f in equipos:
        if f['n'] in raices and raices[f['n']] == f['n'] and not claves[f['n']]:
            sello = (_clave_nombre(cliente_de[f['n']].get('nombre') or cliente_de[f['n']].get('empresa')),
                     _norm(f['equipo'].get('marca')), _norm(f['equipo'].get('modelo')))
            if sello[1] and sello[2]:
                posibles.setdefault(sello, []).append(f)
    duplicado_de = {}
    for lista in posibles.values():
        for f in lista[1:]:
            primera = lista[0]
            if f.get('archivo_id') != primera.get('archivo_id') or (f['historial'] and primera['historial']):
                duplicado_de[f['n']] = primera['n']
    por_clave = {r: next((existentes[k] for n in grupos[r] for k in claves[n] if k in existentes), None)
                 for r in grupos}
    por_modelo, alternativas, candidatos_de, opciones, juntar = _parecidos_del_cliente(
        cur, equipos, grupos, por_clave, cliente_de, (por_correo, por_tel, por_nombre), defecto)
    for r, destino in juntar.items():              # p. ej. la fila de la tabla sin número, con la de su hoja
        for n in grupos.pop(r):
            raices[n] = destino
            grupos[destino].append(n)
        por_clave.pop(r, None)
    # Cómo quedará cada equipo que ya estaba: a qué versión pasa su ficha y qué cambia.
    previstos = _actualizaciones_previstas(cur, {r: por_clave.get(r) or por_modelo.get(r) for r in grupos}, grupos,
                                           {f['n']: f for f in equipos}, hoy)

    por_n = {f['n']: f for f in filas}
    vistas = {}
    for f in filas:
        e = f['equipo']
        v = {'n': f['n'], 'clase': f['clase'], 'incluir': f.get('incluir', True), 'fila': f, 'mensajes': [],
             'origen': f"{f.get('archivo', '')} · {f.get('origen', '')}".strip(' ·'), 'fuente': f.get('fuente'),
             'titulo': st.descripcion_equipo(e) if f['clase'] == 'equipo' else 'PDF sin texto',
             'icono': tipos.icono(e.get('tipo')) if f['clase'] == 'equipo' else 'file-pdf',
             'specs': _specs(e), 'contacto': None, 'cliente_accion': None, 'existente': None, 'mismo_que': None,
             'plan': (None, None), 'historial': f['historial'], 'notas': f['notas'], 'posible_duplicado': None,
             'estado': None, 'cliente': cliente_de.get(f['n'], f['cliente']), 'por_modelo': False,
             'candidato': None, 'cambios': [], 'version_actual': None, 'version_nueva': None, 'opciones': []}
        vistas[f['n']] = v
        if f['clase'] != 'equipo':
            continue
        if not v['incluir']:
            v['estado'] = 'quitada'
            continue
        raiz = raices[f['n']]
        if raiz != f['n']:
            v['mismo_que'] = raiz
            v['mensajes'].append(('ok', f'Es el mismo equipo de la fila {raiz}: se juntan en una sola ficha.'))
        else:
            existente = por_clave.get(raiz)
            prev = previstos.get(raiz)
            if existente:
                v['existente'] = existente
                v['mensajes'].append(('ok', f"Ya está registrado ({existente['descripcion']}, de "
                                            f"{existente['cliente_nombre'] or 'otro cliente'}): "
                                            + _texto_actualizacion(prev)))
            elif f['n'] in por_modelo:
                existente = por_modelo[f['n']]
                v['existente'], v['por_modelo'] = existente, True
                if existente.get('manual'):
                    texto = (f"Escogiste que es el que ya estaba ({existente['descripcion']}, de "
                             f"{existente['cliente_nombre'] or 'otro cliente'}): ")
                else:
                    texto = (f"Es el mismo cliente y un equipo parecido ({existente['descripcion']}, registrado el "
                             f"{existente['creado_en']:%d/%m/%Y}), así que no se crea otro: ")
                v['mensajes'].append(('ok', texto + _texto_actualizacion(prev)
                                      + ('' if existente.get('manual') else
                                         ' Si es un equipo distinto, pulsa «Es otro equipo».')))
                otros = alternativas.get(f['n']) or []
                if otros:
                    v['mensajes'].append(('aviso', f"También se parece a {len(otros)} equipo"
                                                   f"{'s' if len(otros) != 1 else ''} más de este cliente: se escogió "
                                                   'el más parecido y, entre iguales, el que se registró primero. Si '
                                                   'no es ese, escoge cuál es abajo.'))
            elif f.get('es_otro_equipo') and candidatos_de.get(f['n']):
                v['candidato'] = candidatos_de[f['n']][0]
                v['mensajes'].append(('ok', 'Marcaste que es otro equipo: se registra aparte del que ya estaba '
                                            f"({v['candidato']['descripcion']})."))
            elif opciones.get(f['n']):
                k = len(opciones[f['n']])
                v['mensajes'].append(('ok', f"Se registra como equipo nuevo. El cliente ya tiene {k} equipo"
                                            f"{'s' if k != 1 else ''} registrado{'s' if k != 1 else ''}: si es uno "
                                            'de ellos, escógelo abajo y se actualiza en vez de crear otro.'))
            if prev:
                v.update(cambios=prev['cambios'], version_actual=prev['version_actual'],
                         version_nueva=prev['version_nueva'])
            v['opciones'] = opciones.get(f['n']) or alternativas.get(f['n']) or []
        if f.get('ficha'):
            fx = f['ficha']
            v['mensajes'].append(('ok', 'Ficha técnica reconocida' + (f" {f['ficha_codigo']}" if f.get('ficha_codigo') else '')
                                  + (' · roja: presenta inconvenientes' if fx.get('estado') == 'inconveniente'
                                     else ' · azul: en buen estado')
                                  + (f" · asignado a {fx['asignado_a']}" if fx.get('asignado_a') else '') + '.'))
            codigo = (f.get('ficha_codigo') or '').lower()
            otro = codigo_de.get(codigo)
            if codigo and not v['mismo_que'] and ((otro and (not v['existente'] or otro['id'] != v['existente']['id']))
                                                  or codigo in codigos_vistos):
                v['mensajes'].append(('aviso', f"El código {f['ficha_codigo']} ya lo tiene otro equipo: esta ficha queda "
                                               'sin código y se le asigna uno nuevo al generarla.'))
            codigos_vistos.add(codigo)
        # Cliente
        c = cliente_de[f['n']]
        contacto, accion = None, None
        if v['existente']:
            accion = 'del_equipo'
        elif v['mismo_que']:
            accion = 'de_la_fila'
        elif isinstance(f.get('usar_contacto'), int):
            contacto = contactos_elegidos.get(f['usar_contacto'])
            accion = 'existente' if contacto else None
        elif f.get('usar_contacto') != 'nuevo':
            contacto = (por_correo.get(c.get('email') or '') or por_tel.get(c.get('telefono') or '')
                        or por_tel.get(c.get('whatsapp') or ''))
            if contacto:
                accion = 'existente'
                nombre_archivo = c.get('nombre') or c.get('empresa')
                if nombre_archivo and not _palabras_de_nombre(nombre_archivo) & _palabras_de_nombre(contacto['nombre']):
                    v['mensajes'].append(('aviso', f"En el CRM ese celular o correo es de «{contacto['nombre']}» y el "
                                                   f'archivo dice «{nombre_archivo}». Si no es la misma persona, '
                                                   'corrígelo.'))
            else:
                contacto = por_nombre.get(_clave_nombre(c.get('nombre') or c.get('empresa')))
                if contacto:
                    accion = 'existente'
                    v['mensajes'].append(('aviso', f"Se une al contacto «{contacto['nombre']}» del CRM porque tiene "
                                                   'el mismo nombre. Si es otra persona, corrígelo.'))
        if accion is None:
            if c.get('nombre') or c.get('empresa'):
                accion = 'nuevo'
                if f['n'] in cliente_es_empresa:
                    v['mensajes'].append(('ok', f"No dice el cliente: queda la empresa «{c['empresa']}», que es de "
                                                'quien es el documento.'))
            elif defecto:
                accion = 'defecto'
                contacto = defecto.get('existente')
        if accion is None:
            v['mensajes'].append(('error', 'No dice de quién es el equipo. Escribe el cliente en «Corregir» o escoge '
                                           'arriba un cliente para los equipos sin dueño.'))
        v['contacto'], v['cliente_accion'] = contacto, accion
        # Carpeta: la empresa del equipo o «Particulares». Un equipo que ya está en
        # una carpeta se deja ahí, salvo que se escoja otra a mano.
        if v['mismo_que']:
            v['carpeta'] = None                        # la de la fila con la que se junta
        else:
            carpeta = _carpeta_de_fila(f, info_empresa.get(f.get('archivo_id')), negocio_nombre)
            ya = v['existente'] or {}
            if ya.get('empresa_id') and carpeta['fuente'] != 'manual':
                if carpeta['tipo'] == 'empresa' and emp.clave(carpeta['nombre']) != emp.clave(ya['empresa_nombre']):
                    v['mensajes'].append(('aviso', f"El documento dice que es de «{carpeta['nombre']}», pero el equipo "
                                                   f"ya está en la carpeta «{ya['empresa_nombre']}»: se deja ahí. "
                                                   'Si cambió de empresa, escógela en «Corregir».'))
                carpeta = {'tipo': 'empresa', 'nombre': ya['empresa_nombre'], 'fuente': 'existente',
                           'motivo': 'Es la carpeta que ya tiene el equipo.'}
            v['carpeta'] = carpeta
        # Equipo
        if not tiene_datos_fila(f) and not v['existente'] and not v['mismo_que']:
            v['mensajes'].append(('error', 'No se encontraron datos del equipo (marca, modelo, serial…).'))
        if f.get('tipo_adivinado') and not v['mismo_que'] and not v['existente']:
            v['mensajes'].append(('aviso', 'El tipo de equipo se dedujo de los datos: revísalo.' if e['tipo'] != 'otro'
                                  else 'No dice qué tipo de equipo es: quedará como «Otro dispositivo».'))
        if f['n'] in duplicado_de:
            v['posible_duplicado'] = duplicado_de[f['n']]
            v['mensajes'].append(('aviso', f"Puede ser el mismo equipo de la fila {duplicado_de[f['n']]} (mismo "
                                           'cliente, marca y modelo, sin serial). Si es así, únelas.'))
        for aviso in f.get('avisos') or []:
            v['mensajes'].append(('aviso', aviso))
        cada, proximo, aviso_plan = _plan_final(f, ajustes, hoy)
        if v['existente'] and (v['existente']['mant_cada_meses'] or v['existente']['mant_proximo']):
            if cada or proximo:
                v['mensajes'].append(('ok', 'El equipo ya tiene su plan de mantenimiento: se deja el que tiene.'))
            cada, proximo, aviso_plan = None, None, None
        if aviso_plan:
            v['mensajes'].append(('aviso', aviso_plan))
            v['plan_vencido'] = True
            v['aviso_plan'] = aviso_plan
        v['plan'] = (cada, proximo)
        v['estado'] = 'error' if any(m[0] == 'error' for m in v['mensajes']) else (
            'aviso' if any(m[0] == 'aviso' for m in v['mensajes']) else 'ok')
    # Las filas que se juntan dependen de su raíz.
    for v in vistas.values():
        if v['mismo_que'] and v['estado'] in ('ok', 'aviso'):
            raiz = vistas[v['mismo_que']]
            if raiz['estado'] == 'error':
                v['mensajes'].append(('error', f"Se junta con la fila {raiz['n']}, que tiene un error: corrígela "
                                               'primero.'))
                v['estado'] = 'error'
            ya = raiz.get('existente') or {}
            if v['estado'] != 'error' and (ya.get('mant_cada_meses') or ya.get('mant_proximo')):
                # El equipo ya registrado tiene su plan: el de esta fila no se usa (ni su aviso).
                v['plan'] = (None, None)
                v.pop('plan_vencido', None)
                if v.get('aviso_plan'):
                    v['mensajes'] = [m for m in v['mensajes'] if m != ('aviso', v['aviso_plan'])]
                v['estado'] = 'aviso' if any(m[0] == 'aviso' for m in v['mensajes']) else 'ok'
    # PDF escaneados: ¿a qué equipo van?
    for v in vistas.values():
        if v['clase'] != 'adjunto':
            continue
        destino = vistas.get(v['fila'].get('adjuntar_a'))
        if destino and destino['clase'] == 'equipo' and destino['estado'] in ('ok', 'aviso'):
            v['estado'] = 'ok'
            v['mensajes'].append(('ok', f"Se guardará en el equipo de la fila {destino['n']} ({destino['titulo']})."))
        else:
            v['estado'] = 'aviso'
            v['mensajes'].append(('aviso', 'Es un PDF escaneado: no tiene texto que leer. Escoge a qué equipo de la '
                                           'lista pertenece, o súbelo después desde la ficha del equipo.'))
    # Carpetas que ya existen (se usa su nombre tal como quedó guardado).
    hechas = emp.por_claves(cur, {emp.clave(v['carpeta']['nombre']) for v in vistas.values()
                                  if v.get('carpeta') and v['carpeta']['tipo'] == 'empresa'})
    for v in vistas.values():
        c = v.get('carpeta')
        if c and c['tipo'] == 'empresa':
            hallada = hechas.get(emp.clave(c['nombre']))
            c['existe'] = bool(hallada)
            if hallada:
                c['nombre'], c['id'] = hallada['nombre'], hallada['id']
    # Resumen
    resumen = {'equipos': 0, 'incluidas': 0, 'ok': 0, 'aviso': 0, 'error': 0, 'quitadas': 0, 'nuevos': 0,
               'existentes': 0, 'unidas': 0, 'clientes_nuevos': 0, 'clientes_existentes': 0, 'mantenimientos': 0,
               'planes': 0, 'documentos': 0, 'vencidos': 0, 'adjuntos': 0, 'carpetas': 0, 'carpetas_nuevas': 0,
               'particulares': 0}
    carpetas_usadas, carpetas_nuevas = set(), set()
    agenda, contactos_usados, defecto_nuevo = _Agenda(), set(), False
    hojas_de_vida, con_plan = {}, set()
    for v in vistas.values():
        if v['clase'] == 'adjunto':
            resumen['adjuntos'] += 1
            resumen['documentos'] += 1 if v['estado'] == 'ok' else 0
            continue
        resumen['equipos'] += 1
        if v['estado'] == 'quitada':
            resumen['quitadas'] += 1
            continue
        resumen[v['estado']] += 1
        if v['estado'] == 'error':
            continue
        resumen['incluidas'] += 1
        f = v['fila']
        if v['mismo_que']:
            resumen['unidas'] += 1
        elif v['existente']:
            resumen['existentes'] += 1
        else:
            resumen['nuevos'] += 1
        c = v.get('carpeta')
        if c and c['tipo'] == 'empresa':
            carpetas_usadas.add(emp.clave(c['nombre']))
            if not c.get('existe'):
                carpetas_nuevas.add(emp.clave(c['nombre']))
        elif c:
            resumen['particulares'] += 1
        grupo = v['mismo_que'] or v['n']                 # las filas que se juntan son un solo equipo
        hojas_de_vida.setdefault(grupo, set()).update(
            (h['fecha'], h['tipo']) for h in f['historial'] if _iso(h['fecha']) and _iso(h['fecha']) <= hoy)
        if v['plan'][1]:
            con_plan.add(grupo)
        resumen['vencidos'] += 1 if v.get('plan_vencido') else 0
        resumen['documentos'] += 1 if f.get('documento_id') else 0
        if v['cliente_accion'] in ('existente', 'defecto') and v['contacto']:
            contactos_usados.add(v['contacto']['id'])
        elif v['cliente_accion'] == 'nuevo':
            agenda.persona(v['cliente'])
        elif v['cliente_accion'] == 'defecto':
            defecto_nuevo = True
    resumen['clientes_nuevos'] = len(agenda.personas) + (1 if defecto_nuevo else 0)
    resumen['clientes_existentes'] = len(contactos_usados)
    resumen['carpetas'], resumen['carpetas_nuevas'] = len(carpetas_usadas), len(carpetas_nuevas)
    # Mantenimientos que de verdad se agregan: los que el equipo ya registrado no tiene.
    ya_tiene, ids = {}, sorted({vistas[g]['existente']['id'] for g in hojas_de_vida if vistas[g].get('existente')})
    if ids:
        cur.execute('SELECT equipo_id, fecha, tipo FROM st_mantenimientos WHERE equipo_id = ANY(%s) AND activo', (ids,))
        for r in cur.fetchall():
            ya_tiene.setdefault(r['equipo_id'], set()).add((r['fecha'].isoformat(), r['tipo']))
    resumen['mantenimientos'] = sum(len(entradas - ya_tiene.get((vistas[g].get('existente') or {}).get('id'), set()))
                                    for g, entradas in hojas_de_vida.items())
    resumen['planes'] = len(con_plan)
    return {'filas': [vistas[f['n']] for f in filas], 'resumen': resumen, 'defecto': defecto,
            'por_n': {n: v for n, v in vistas.items()}}


def revisar(lote, hoy=None):
    """Vista previa del lote: estado de cada fila y lo que se hará. No escribe."""
    hoy = hoy or date.today()
    with get_db_cursor(dict_cursor=True) as cur:
        return _revisar(cur, lote['filas'], lote.get('ajustes') or {}, hoy, lote.get('archivos') or ())


# ── Importar (una sola transacción) ─────────────────────────────
def _crear_contacto(cur, c, lote_id):
    nombre = c.get('nombre') or c.get('empresa')
    notas = [f'Importado desde Servicio Técnico (importación #{lote_id}).']
    if c.get('documento'):
        notas.append(f"Documento: {c['documento']}")
    cur.execute("""INSERT INTO crm_contactos (tipo, nombre, empresa, email, telefono, whatsapp, direccion, ciudad,
                                              origen, notas)
                   VALUES ('cliente', %s, %s, %s, %s, %s, %s, %s, 'servicio_tecnico', %s) RETURNING id""",
                (nombre[:150], (c.get('empresa') or None) if c.get('nombre') else None, c.get('email'),
                 _formato_tel(c.get('telefono')), _formato_tel(c.get('whatsapp') or c.get('telefono')),
                 (c.get('direccion') or None) and c['direccion'][:200], (c.get('ciudad') or None) and c['ciudad'][:80],
                 '\n'.join(notas)))
    return cur.fetchone()[0]


def _completar_existente(cur, equipo_id, nuevo):
    """Llena SOLO lo que le falta al equipo que ya existía. Devuelve lo que
    cambió (para poder deshacerlo)."""
    cur.execute('SELECT * FROM st_equipos WHERE id = %s FOR UPDATE', (equipo_id,))
    actual = dict(cur.fetchone())
    columnas = {}
    for c in st._COLUMNAS_EQUIPO:
        valor = nuevo.get(c)
        if valor and not actual.get(c) and c in tipos.columnas(actual['tipo']) + ['serial', 'imei']:
            if c == 'imei' and not st.imei_valido(valor):
                continue
            columnas[c] = str(valor)[:st._LARGOS[c]]
    extras_actuales = actual.get('extras') or {}
    extras = {k: v for k, v in tipos.limpiar_extras(actual['tipo'], nuevo.get('extras') or {}).items()
              if not extras_actuales.get(k)}
    notas = None
    if nuevo.get('notas'):
        ya = {l.strip() for l in (actual.get('notas') or '').split('\n')}
        lineas = [l.strip() for l in nuevo['notas'].split('\n') if l.strip() and l.strip() not in ya]
        if lineas:                                   # solo lo que no estaba ya en las notas
            nueva = ((actual.get('notas') or '') + '\n' + '\n'.join(lineas)).strip()[:2000]
            notas = {'antes': actual.get('notas'), 'despues': nueva}
    sets, params = [], []
    for c, valor in columnas.items():
        sets.append(f'{c} = %s')
        params.append(valor)
    if extras:
        sets.append("extras = COALESCE(extras, '{}'::jsonb) || %s::jsonb")
        params.append(json.dumps(extras, ensure_ascii=False))
    if notas:
        sets.append('notas = %s')
        params.append(notas['despues'])
    if sets:
        cur.execute(f"UPDATE st_equipos SET {', '.join(sets)}, actualizado_en = NOW() WHERE id = %s",
                    (*params, equipo_id))
    return {'equipo_id': equipo_id, 'columnas': columnas, 'extras': extras, 'notas': notas} if sets else None


def _entero(valor):
    try:
        return int(valor or 0)
    except (TypeError, ValueError):
        return 0


def _unir_ficha(actual, llega):
    """La ficha que llega es la nueva versión, pero lo que ella no dice se conserva
    de la anterior: filas, textos y las fotos escogidas. El estado (azul o roja),
    las filas en rojo y el problema los dice el documento que llega."""
    if not actual:
        return {k: v for k, v in llega.items() if k != 'fotos'}
    nueva = dict(actual)
    for k, v in llega.items():
        if k == 'filas':
            filas = dict(actual.get('filas') or {})
            filas.update({etq: x for etq, x in (v or {}).items() if str(x or '').strip()})
            nueva['filas'] = filas
        elif k == 'orden':
            orden = [e for e in v or [] if isinstance(e, str)]
            nueva['orden'] = orden + [e for e in actual.get('orden') or [] if e not in orden]
        elif k in ('version', 'fecha', 'fotos'):
            continue
        elif v not in (None, '', [], {}):
            nueva[k] = v
    if llega.get('estado'):
        nueva['destacar'] = list(llega.get('destacar') or [])
        if llega['estado'] == 'bueno' and not llega.get('problema'):
            nueva.pop('problema', None)
    if llega.get('fabricante') and _norm(llega['fabricante']) != _norm(actual.get('fabricante')) \
            and not llega.get('fabricante_url'):
        nueva.pop('fabricante_url', None)            # el enlace era del fabricante de antes
    return nueva


def _ya_lo_dice(actual, nuevo):
    """¿Lo que llega ya está dicho (igual o con menos detalle) en lo que hay? Una
    tabla que dice «IdeaCentre B320» no reemplaza «IdeaCentre B320 (Touch)», ni
    «16 GB» a «16 GB DDR4»; «16 GB DDR4» sí reemplaza «8 GB DDR4»."""
    palabras = _norm(nuevo).split()
    return bool(palabras) and set(palabras) <= set(_norm(actual).split())


def _equipo_actualizado(actual, nuevo, fila, hoy, version_fija=None):
    """Cómo queda un equipo que ya existía con lo que trae una fila del archivo
    (sin escribir nada). Lo que trae el archivo manda y lo que no trae se
    conserva; el serial y el IMEI solo se llenan si faltaban (si son otros, es
    otro equipo). Si cambia lo que dice la ficha, queda como versión nueva
    (`version_fija`: la que ya le dio esta misma importación)."""
    d = dict(actual)
    tipo = d['tipo']
    if tipo == 'otro' and tipos.es_valido(nuevo.get('tipo') or '') and nuevo['tipo'] != 'otro' \
            and not fila.get('tipo_adivinado'):
        d['tipo'] = tipo = nuevo['tipo']             # la ficha dice qué es («Impresora de carnés…»)
    permitidas = set(tipos.columnas(tipo)) | {'serial', 'imei'}
    for c in st._COLUMNAS_EQUIPO:
        valor = nuevo.get(c)
        if not valor or c not in permitidas:
            continue
        valor = str(valor)[:st._LARGOS[c]]
        if c == 'imei' and not st.imei_valido(valor):
            continue
        if d.get(c) and (c in ('serial', 'imei') or _ya_lo_dice(d[c], valor)):
            continue
        d[c] = valor
    extras = dict(d.get('extras') or {})
    for k, v in tipos.limpiar_extras(tipo, nuevo.get('extras') or {}).items():
        if v not in (None, '') and not (extras.get(k) not in (None, '') and _ya_lo_dice(extras[k], v)):
            extras[k] = v
    d['extras'] = extras
    resumen = ' '.join(str(nuevo.get('resumen_ia') or '').split())[:800]
    if resumen and not _ya_lo_dice(d.get('resumen_ia'), resumen):
        d['resumen_ia'] = resumen
    if nuevo.get('notas'):                           # solo lo que no estaba ya en las notas
        ya = {linea.strip() for linea in (d.get('notas') or '').split('\n')}
        lineas = [x.strip() for x in nuevo['notas'].split('\n') if x.strip() and x.strip() not in ya]
        if lineas:
            d['notas'] = ((d.get('notas') or '') + '\n' + '\n'.join(lineas)).strip()[:2000]
    if fila.get('ficha'):
        d['ficha'] = _unir_ficha(actual.get('ficha') or {}, fila['ficha'])
    if fila.get('ficha_codigo') and not d.get('ficha_codigo'):
        d['ficha_codigo'] = fila['ficha_codigo'][:40]
    antes, despues = ver.estado_de(actual), ver.estado_de(d)
    if d.get('ficha') and ver.diferencias(antes, despues):
        previa, llega = actual.get('ficha') or {}, fila.get('ficha') or {}
        if version_fija:
            numero = version_fija
        elif previa:
            numero = max(ver.numero(antes) + 1, _entero(llega.get('version')))
        else:
            numero = max(1, _entero(llega.get('version')))
        # La primera ficha lleva la fecha de su documento; una versión nueva, la de hoy.
        fecha = hoy.isoformat() if previa else (llega.get('fecha') or hoy.isoformat())
        d['ficha'] = {**d['ficha'], 'version': numero, 'fecha': fecha}
    elif ver.igual(antes, despues):
        d = dict(actual, notas=d.get('notas'))     # la ficha ya decía esto: a lo sumo, notas nuevas
    return d


def _actualizar_existente(cur, equipo_id, nuevo, fila, lote_id, usuario_id, hoy, hechas):
    """El archivo trae otra vez un equipo que ya existía: se actualiza con lo que
    trae y su ficha queda como versión nueva; la anterior se guarda (versiones).
    Lo que el archivo no trae se conserva, y sus fotos, documentos,
    mantenimientos, piezas, órdenes y recordatorios no se tocan. `hechas` junta lo
    de esta importación por equipo (una sola versión aunque venga en varias filas).
    Devuelve lo hecho con este equipo (para deshacerlo) o None si no cambió nada."""
    cur.execute('SELECT * FROM st_equipos WHERE id = %s FOR UPDATE', (equipo_id,))
    actual = dict(cur.fetchone())
    previa = hechas.get(equipo_id)
    d = _equipo_actualizado(actual, nuevo, fila, hoy, version_fija=previa.get('version_nueva') if previa else None)
    if d.get('ficha_codigo') and d['ficha_codigo'] != actual.get('ficha_codigo'):
        cur.execute("SELECT pg_advisory_xact_lock(hashtext('st_ficha_codigo'))")   # el mismo candado del consecutivo
        cur.execute('SELECT 1 FROM st_equipos WHERE lower(ficha_codigo) = lower(%s) AND activo AND id <> %s',
                    (d['ficha_codigo'], equipo_id))
        if cur.fetchone():
            d['ficha_codigo'] = actual.get('ficha_codigo')
    sets, params = [], []
    for c in st._COLUMNAS_EQUIPO + ('tipo', 'notas', 'resumen_ia', 'ficha_codigo'):
        if d.get(c) != actual.get(c):
            sets.append(f'{c} = %s')
            params.append(d.get(c))
    for c in ('extras', 'ficha'):
        if (d.get(c) or {}) != (actual.get(c) or {}):
            sets.append(f'{c} = %s::jsonb')
            params.append(json.dumps(d.get(c) or {}, ensure_ascii=False, default=str))
    if not sets:
        return previa
    cur.execute(f"UPDATE st_equipos SET {', '.join(sets)}, actualizado_en = NOW() WHERE id = %s", (*params, equipo_id))
    entrada = previa or {'equipo_id': equipo_id, 'antes': ver.estado_de(actual), 'notas_antes': actual.get('notas'),
                         'version_id': None, 'version_nueva': None}
    entrada['despues'] = ver.estado_de(d)
    entrada['notas'] = ({'antes': entrada['notas_antes'], 'despues': d.get('notas')}
                        if d.get('notas') != entrada['notas_antes'] else None)
    version_id = ver.registrar(cur, equipo_id, entrada['antes'], entrada['despues'], 'importacion', lote_id, usuario_id)
    entrada['version_id'] = version_id or entrada['version_id']
    nueva = (d.get('ficha') or {}).get('version')
    if nueva and nueva != (actual.get('ficha') or {}).get('version'):
        entrada['version_nueva'] = nueva
        st._evento(cur, None, equipo_id, 'equipo',
                   f'Ficha técnica: versión {nueva} desde la importación #{lote_id}'
                   + (' (la anterior queda guardada)' if actual.get('ficha') else ''), usuario_id)
    entrada['cambios'] = ver.diferencias(entrada['antes'], entrada['despues'])
    hechas[equipo_id] = entrada
    return entrada


def _poner_ficha(cur, equipo_id, fila):
    """La ficha técnica importada: si el equipo no tenía ficha, o si la que
    llega es más nueva (versión mayor o fecha posterior). El código se pone si
    ningún otro equipo activo lo tiene. Devuelve lo hecho (con la ficha que
    había, para deshacer) o None."""
    cur.execute('SELECT ficha, ficha_codigo FROM st_equipos WHERE id = %s FOR UPDATE', (equipo_id,))
    actual = cur.fetchone()
    antes = actual['ficha'] or {}

    def orden(fx):
        try:
            return int(fx.get('version') or 1), str(fx.get('fecha') or '')
        except (TypeError, ValueError):
            return 1, str(fx.get('fecha') or '')
    if antes and orden(fila['ficha']) <= orden(antes):
        return None                                  # la que tiene es igual o más reciente
    codigo = fila.get('ficha_codigo') if not actual['ficha_codigo'] else None
    if codigo:
        cur.execute("SELECT pg_advisory_xact_lock(hashtext('st_ficha_codigo'))")   # el mismo candado del consecutivo
        cur.execute('SELECT 1 FROM st_equipos WHERE lower(ficha_codigo) = lower(%s) AND activo AND id <> %s',
                    (codigo, equipo_id))
        if cur.fetchone():
            codigo = None
    cur.execute('UPDATE st_equipos SET ficha = %s::jsonb, ficha_codigo = COALESCE(ficha_codigo, %s) WHERE id = %s',
                (json.dumps(fila['ficha'], ensure_ascii=False), codigo, equipo_id))
    return {'equipo_id': equipo_id, 'codigo': codigo, 'ficha': fila['ficha'], 'antes': antes or None}


def _fotos_del_documento(cur, equipo_id, doc_id, usuario_id=None, hoja=None):
    """Las fotos que trae la ficha (el PDF, o la hoja del libro de Excel), si el
    equipo aún no tiene fotos."""
    from services import servicio_tecnico_ficha as fichas
    cur.execute('SELECT COUNT(*) AS n FROM st_fotos WHERE equipo_id = %s AND activo', (equipo_id,))
    if cur.fetchone()['n']:
        return []
    cur.execute('SELECT contenido FROM st_documentos WHERE id = %s', (doc_id,))
    fila = cur.fetchone()
    if not fila:
        return []
    ids = []
    contenido = bytes(fila['contenido'])
    imagenes = fichas.fotos_de_excel(contenido, hoja) if hoja else fichas.fotos_de_pdf(contenido, maximo=4)
    for imagen in imagenes:
        try:
            grande, mini, ancho, alto = st.preparar_foto(imagen)
        except ErrorServicio:
            continue
        cur.execute("""INSERT INTO st_fotos (equipo_id, momento, descripcion, mime, ancho, alto, bytes, contenido,
                                             miniatura, creado_por)
                       VALUES (%s, 'ficha', 'Foto de la ficha técnica importada', 'image/jpeg', %s, %s, %s, %s, %s, %s)
                       RETURNING id""", (equipo_id, ancho, alto, len(grande), grande, mini, usuario_id))
        ids.append(cur.fetchone()['id'])
    return ids


def _poner_carpeta(cur, equipo_id, carpeta, res, usuario_id, contacto_id, cache):
    """Pone el equipo en la carpeta que salió en la revisión (la crea si no
    existe). Un equipo que ya estaba en una carpeta solo cambia si se escogió
    otra a mano. Lo hecho queda en `res` para poder deshacerlo."""
    cur.execute('SELECT empresa_id FROM st_equipos WHERE id = %s', (equipo_id,))
    antes = cur.fetchone()['empresa_id']
    if carpeta.get('tipo') != 'empresa':
        if antes and carpeta.get('fuente') == 'manual':          # a «Particulares», escogido a mano
            cur.execute('UPDATE st_equipos SET empresa_id = NULL WHERE id = %s', (equipo_id,))
            res['carpetas_puestas'].append({'equipo_id': equipo_id, 'empresa_id': None, 'antes': antes})
            st._evento(cur, None, equipo_id, 'equipo', f'Carpeta: {emp.PARTICULARES}', usuario_id)
        return
    if carpeta.get('fuente') == 'existente' or (antes and carpeta.get('fuente') != 'manual'):
        return
    c = emp.clave(carpeta['nombre'])
    if c not in cache:
        vincular = None
        if contacto_id:                       # el cliente del equipo ES la empresa: queda enlazada al CRM
            cur.execute('SELECT nombre FROM crm_contactos WHERE id = %s', (contacto_id,))
            fila = cur.fetchone()
            vincular = contacto_id if fila and emp.clave(fila['nombre']) == c else None
        empresa_id, creada = emp.obtener_o_crear(cur, carpeta['nombre'], crm_contacto_id=vincular,
                                                 usuario_id=usuario_id)
        cache[c] = empresa_id
        if creada:
            res['empresas_creadas'].append(empresa_id)
    if antes == cache[c]:
        return
    cur.execute('UPDATE st_equipos SET empresa_id = %s WHERE id = %s', (cache[c], equipo_id))
    res['carpetas_puestas'].append({'equipo_id': equipo_id, 'empresa_id': cache[c], 'antes': antes})
    st._evento(cur, None, equipo_id, 'equipo', f"Carpeta: {carpeta['nombre']}", usuario_id)


def importar(lote_id, usuario_id=None, hoy=None):
    """Guarda lo aprobado del lote en UNA transacción: si algo falla, no queda
    nada a medias. Devuelve el resultado (también queda en el lote)."""
    from services import servicio_tecnico_seguimiento as seg
    hoy = hoy or date.today()
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote:
            raise ErrorServicio('La importación no existe.')
        if lote['estado'] != 'revision':
            raise ErrorServicio('Esta importación ya se procesó.')
        vista = _revisar(cur, lote['filas'], lote.get('ajustes') or {}, hoy, lote.get('archivos') or ())
        nombres_archivo = {a['id']: a['nombre'] for a in lote['archivos']}
        res = {'equipos_creados': [], 'equipos_completados': [], 'contactos_creados': [], 'mantenimientos': [],
               'seguimientos': [], 'documentos': [], 'planes': [], 'filas': {}, 'fichas': [], 'fotos': [],
               'empresas_creadas': [], 'carpetas_puestas': [], 'equipos_existentes': [], 'versiones': []}
        carpetas_cache, hechas = {}, {}
        listas = [v for v in vista['filas'] if v['clase'] == 'equipo' and v['estado'] in ('ok', 'aviso')]
        # Clientes nuevos: uno por persona, con los datos de todas sus filas.
        agenda, persona_de = _Agenda(), {}
        for v in listas:
            if v['cliente_accion'] == 'nuevo':
                persona_de[v['n']] = agenda.persona(v['cliente'])
        defecto_id = None
        equipo_de_fila, contacto_de_equipo, planes = {}, {}, {}

        def contacto_para(v):
            nonlocal defecto_id
            if v['cliente_accion'] in ('existente', 'defecto') and v['contacto']:
                return v['contacto']['id']
            if v['cliente_accion'] == 'defecto':
                if defecto_id is None:
                    defecto_id = _crear_contacto(cur, vista['defecto']['nuevo'], lote_id)
                    res['contactos_creados'].append(defecto_id)
                return defecto_id
            persona = persona_de[v['n']]
            if persona['id'] is None:
                persona['id'] = _crear_contacto(cur, persona['datos'], lote_id)
                res['contactos_creados'].append(persona['id'])
            return persona['id']

        # Primero las filas que crean o completan un equipo; después las que se juntan con otra.
        for v in sorted(listas, key=lambda x: (1 if x['mismo_que'] else 0, x['n'])):
            f = v['fila']
            try:
                datos_equipo = dict(f['equipo'])
                datos_equipo['notas'] = '\n'.join(f['notas'])[:2000] or None
                datos_equipo['resumen_ia'] = f.get('resumen')
                if v['mismo_que']:
                    equipo_id = equipo_de_fila.get(v['mismo_que'])
                    if not equipo_id:
                        raise ErrorServicio(f"la fila {v['mismo_que']} con la que se junta no se pudo importar.")
                    if equipo_id in res['equipos_creados']:       # equipo nuevo de este lote: se completa
                        _completar_existente(cur, equipo_id, datos_equipo)
                        if tipos.es_valido(f['equipo'].get('tipo') or '') and f['equipo']['tipo'] != 'otro' \
                                and not f.get('tipo_adivinado'):
                            # La hoja de su ficha dice qué es («Impresora de carnés…»): manda sobre «Otro».
                            cur.execute("UPDATE st_equipos SET tipo = %s WHERE id = %s AND tipo = 'otro'",
                                        (f['equipo']['tipo'], equipo_id))
                    else:                                         # se junta con uno que ya estaba: versión nueva
                        _actualizar_existente(cur, equipo_id, datos_equipo, f, lote_id, usuario_id, hoy, hechas)
                elif v['existente']:
                    equipo_id = v['existente']['id']
                    contacto_de_equipo[equipo_id] = v['existente']['crm_contacto_id']
                    res['equipos_existentes'].append(equipo_id)
                    _actualizar_existente(cur, equipo_id, datos_equipo, f, lote_id, usuario_id, hoy, hechas)
                else:
                    contacto_id = contacto_para(v)
                    equipo_id = st.crear_equipo(cur, contacto_id, datos_equipo, usuario_id)
                    cur.execute('UPDATE st_equipos SET importacion_id = %s WHERE id = %s', (lote_id, equipo_id))
                    res['equipos_creados'].append(equipo_id)
                    contacto_de_equipo[equipo_id] = contacto_id
                equipo_de_fila[v['n']] = equipo_id
                res['filas'][str(v['n'])] = equipo_id
                if not v['mismo_que']:
                    _poner_carpeta(cur, equipo_id, v.get('carpeta') or {}, res, usuario_id,
                                   contacto_de_equipo.get(equipo_id), carpetas_cache)
                # Hoja de vida: los mantenimientos que no tenga ya (misma fecha y tipo).
                cur.execute('SELECT fecha, tipo FROM st_mantenimientos WHERE equipo_id = %s AND activo', (equipo_id,))
                ya = {(r['fecha'].isoformat(), r['tipo']) for r in cur.fetchall()}
                for h in f['historial']:
                    if (h['fecha'], h['tipo']) in ya or _iso(h['fecha']) > hoy:
                        continue
                    cur.execute("""INSERT INTO st_mantenimientos (equipo_id, tipo, fecha, descripcion, costo,
                                                                  creado_por, importacion_id)
                                   VALUES (%s, %s, %s, %s, %s, %s, %s) RETURNING id""",
                                (equipo_id, h['tipo'], h['fecha'], h['descripcion'][:4000], h.get('costo'),
                                 usuario_id, lote_id))
                    res['mantenimientos'].append(cur.fetchone()['id'])
                    ya.add((h['fecha'], h['tipo']))
                if v['plan'][1] and (equipo_id not in planes or (not planes[equipo_id][0] and v['plan'][0])):
                    planes[equipo_id] = v['plan']      # entre filas del mismo equipo, el plan que dice cada cuánto
                if f.get('documento_id'):
                    cur.execute("""UPDATE st_documentos SET equipo_id = %s
                                   WHERE id = %s AND importacion_id = %s AND equipo_id IS NULL AND activo""",
                                (equipo_id, f['documento_id'], lote_id))
                    if cur.rowcount:
                        res['documentos'].append(f['documento_id'])
                if f.get('ficha') and equipo_id in res['equipos_creados']:   # (la de uno que ya estaba: arriba)
                    ficha = _poner_ficha(cur, equipo_id, f)
                    if ficha:
                        res['fichas'].append(ficha)
                if f.get('fotos_pdf') and f.get('documento_id'):
                    res['fotos'] += _fotos_del_documento(cur, equipo_id, f['documento_id'], usuario_id)
                elif f.get('fotos_pdf') and f.get('fotos_hoja') and f.get('archivo_id'):
                    res['fotos'] += _fotos_del_documento(cur, equipo_id, f['archivo_id'], usuario_id,
                                                         hoja=f['fotos_hoja'])
                archivo = nombres_archivo.get(f.get('archivo_id'), 'un archivo')
                st._evento(cur, None, equipo_id, 'equipo',
                           f'Importado desde «{archivo}» (importación #{lote_id})'
                           + (f": {len(f['historial'])} mantenimiento(s) de su hoja de vida" if f['historial'] else ''),
                           usuario_id)
            except ErrorServicio as exc:
                raise ErrorServicio(f"Fila {v['n']} ({v['origen']}): {exc}")
        if not equipo_de_fila:
            raise ErrorServicio('No hay filas listas para importar: corrige las que tienen error o inclúyelas.')
        # Plan de mantenimiento y su recordatorio (solo si el equipo no tiene ya uno).
        for equipo_id, (cada, proximo) in planes.items():
            cur.execute('SELECT crm_contacto_id, mant_cada_meses, mant_proximo FROM st_equipos WHERE id = %s',
                        (equipo_id,))
            actual = cur.fetchone()
            if actual['mant_cada_meses'] or actual['mant_proximo']:
                continue
            cur.execute('UPDATE st_equipos SET mant_cada_meses = %s, mant_proximo = %s WHERE id = %s',
                        (cada, proximo, equipo_id))
            res['planes'].append({'equipo_id': equipo_id, 'cada_meses': cada, 'proximo': proximo.isoformat()})
            seg_id = seg.programar(cur, 'mantenimiento', proximo, equipo_id=equipo_id,
                                   contacto_id=actual['crm_contacto_id'])
            if seg_id:
                res['seguimientos'].append(seg_id)
        # PDF escaneados que se escogieron para un equipo.
        for v in vista['filas']:
            f = v['fila']
            if v['clase'] == 'adjunto' and v['estado'] == 'ok' and equipo_de_fila.get(f.get('adjuntar_a')):
                cur.execute("""UPDATE st_documentos SET equipo_id = %s
                               WHERE id = %s AND importacion_id = %s AND equipo_id IS NULL AND activo""",
                            (equipo_de_fila[f['adjuntar_a']], f['archivo_id'], lote_id))
                if cur.rowcount:
                    res['documentos'].append(f['archivo_id'])
        # Una nota en la línea de tiempo del CRM por cliente.
        por_contacto = {}
        for equipo_id in set(equipo_de_fila.values()):
            if equipo_id not in contacto_de_equipo:
                cur.execute('SELECT crm_contacto_id FROM st_equipos WHERE id = %s', (equipo_id,))
                contacto_de_equipo[equipo_id] = cur.fetchone()['crm_contacto_id']
            por_contacto.setdefault(contacto_de_equipo[equipo_id], set()).add(equipo_id)
        for contacto_id, equipos_ids in por_contacto.items():
            st._actividad_crm(cur, contacto_id, 'nota', 'Servicio técnico: equipos importados',
                              f'{len(equipos_ids)} equipo(s) traídos desde archivos (importación #{lote_id}).',
                              usuario_id)
        res['versiones'] = [{k: x for k, x in h.items() if k != 'notas_antes'} for h in hechas.values()]
        res['resumen'] = {'equipos_creados': len(res['equipos_creados']),
                          'equipos_completados': len({c['equipo_id'] for c in res['equipos_completados']}
                                                     | set(res['equipos_existentes'])),
                          'fichas_actualizadas': sum(1 for h in hechas.values() if h.get('version_id')),
                          'contactos_creados': len(res['contactos_creados']),
                          'mantenimientos': len(res['mantenimientos']), 'recordatorios': len(res['seguimientos']),
                          'documentos': len(res['documentos']),
                          'carpetas': len({c['empresa_id'] for c in res['carpetas_puestas'] if c['empresa_id']}),
                          'carpetas_creadas': len(res['empresas_creadas'])}
        _guardar(cur, lote_id, estado='importada', resultado=res, importado_por=usuario_id)
        cur.execute('UPDATE st_importaciones SET importado_en = NOW() WHERE id = %s', (lote_id,))
    return res


# ── Deshacer (sin borrar nada) ──────────────────────────────────
def deshacer(lote_id, usuario_id=None):
    """Retira lo que creó la importación: equipos, mantenimientos, documentos y
    recordatorios quedan activo = FALSE / omitidos; los datos que completó en
    equipos que ya existían vuelven a quedar vacíos si nadie los cambió; los
    contactos que creó se archivan si no se han usado para nada más."""
    with get_db_cursor(dict_cursor=True) as cur:
        lote = obtener_lote(lote_id, cur, bloquear=True)
        if not lote:
            raise ErrorServicio('La importación no existe.')
        if lote['estado'] != 'importada':
            raise ErrorServicio('Solo se puede deshacer una importación ya hecha.')
        res = lote['resultado'] or {}
        creados = res.get('equipos_creados') or []
        conservados = []
        if creados:
            cur.execute('SELECT DISTINCT equipo_id FROM st_ordenes WHERE equipo_id = ANY(%s)', (creados,))
            conservados = [r['equipo_id'] for r in cur.fetchall()]
        retirar = [e for e in creados if e not in conservados]
        if retirar:
            cur.execute('UPDATE st_equipos SET activo = FALSE, actualizado_en = NOW() WHERE id = ANY(%s) AND activo',
                        (retirar,))
            cur.execute("""UPDATE st_seguimientos SET estado = 'omitido', hecho_en = NOW(),
                                  comentario = COALESCE(comentario, 'Importación deshecha')
                           WHERE equipo_id = ANY(%s) AND estado = 'pendiente'""", (retirar,))
            cur.execute('UPDATE st_fotos SET activo = FALSE WHERE equipo_id = ANY(%s) AND activo', (retirar,))
        cur.execute("""UPDATE st_mantenimientos SET activo = FALSE
                       WHERE importacion_id = %s AND activo AND NOT (equipo_id = ANY(%s))""", (lote_id, conservados))
        if res.get('seguimientos'):
            cur.execute("""UPDATE st_seguimientos SET estado = 'omitido', hecho_en = NOW(),
                                  comentario = COALESCE(comentario, 'Importación deshecha')
                           WHERE id = ANY(%s) AND estado = 'pendiente'""", (res['seguimientos'],))
        cur.execute("""UPDATE st_documentos SET activo = FALSE
                       WHERE importacion_id = %s AND activo AND (equipo_id IS NULL OR NOT (equipo_id = ANY(%s)))""",
                    (lote_id, conservados))
        for plan in res.get('planes') or []:
            if plan['equipo_id'] in conservados or plan['equipo_id'] not in creados:
                cur.execute("""UPDATE st_equipos SET mant_cada_meses = NULL, mant_proximo = NULL
                               WHERE id = %s AND mant_cada_meses IS NOT DISTINCT FROM %s
                                 AND mant_proximo IS NOT DISTINCT FROM %s""",
                            (plan['equipo_id'], plan['cada_meses'], plan['proximo']))
        for fx in res.get('fichas') or []:
            if fx['equipo_id'] in conservados or fx['equipo_id'] in retirar:
                continue
            cur.execute('SELECT ficha FROM st_equipos WHERE id = %s', (fx['equipo_id'],))
            actual = cur.fetchone()
            if actual and actual['ficha'] == fx['ficha']:      # nadie la cambió después: vuelve la de antes
                cur.execute("""UPDATE st_equipos SET ficha = %s::jsonb,
                                      ficha_codigo = CASE WHEN ficha_codigo = %s THEN NULL ELSE ficha_codigo END
                               WHERE id = %s""",
                            (json.dumps(fx.get('antes') or {}, ensure_ascii=False), fx.get('codigo'), fx['equipo_id']))
        if res.get('fotos'):
            cur.execute('UPDATE st_fotos SET activo = FALSE WHERE id = ANY(%s) AND activo AND NOT (equipo_id = ANY(%s))',
                        (res['fotos'], conservados))
        # Equipos que ya estaban: cada dato vuelve a la versión anterior si nadie lo cambió después.
        devueltos = ver.deshacer_importacion(cur, lote_id, [x for x in res.get('versiones') or []
                                                            if x['equipo_id'] not in retirar])
        for cambio in res.get('equipos_completados') or []:
            eid = cambio['equipo_id']
            for c, valor in (cambio.get('columnas') or {}).items():
                if c in st._COLUMNAS_EQUIPO:
                    cur.execute(f'UPDATE st_equipos SET {c} = NULL WHERE id = %s AND {c} = %s', (eid, valor))
                    devueltos += cur.rowcount
            for k, valor in (cambio.get('extras') or {}).items():
                cur.execute("UPDATE st_equipos SET extras = extras - %s WHERE id = %s AND extras->>%s = %s",
                            (k, eid, k, str(valor)))
                devueltos += cur.rowcount
            notas = cambio.get('notas')
            if notas:
                cur.execute('UPDATE st_equipos SET notas = %s WHERE id = %s AND notas = %s',
                            (notas['antes'], eid, notas['despues']))
                devueltos += cur.rowcount
        # Carpetas: los equipos que se quedan vuelven a la que tenían; las carpetas
        # que creó la importación se archivan si ya no tienen equipos.
        for cp in res.get('carpetas_puestas') or []:
            if cp['equipo_id'] in retirar:
                continue
            cur.execute('UPDATE st_equipos SET empresa_id = %s WHERE id = %s AND empresa_id IS NOT DISTINCT FROM %s',
                        (cp.get('antes'), cp['equipo_id'], cp['empresa_id']))
            devueltos += cur.rowcount
        carpetas_archivadas = 0
        for eid in res.get('empresas_creadas') or []:
            cur.execute("""UPDATE st_empresas SET activo = FALSE WHERE id = %s AND activo
                           AND NOT EXISTS (SELECT 1 FROM st_equipos WHERE empresa_id = %s AND activo)""", (eid, eid))
            carpetas_archivadas += cur.rowcount
        archivados, contactos_conservados = [], []
        for cid in res.get('contactos_creados') or []:
            cur.execute("""SELECT
                             (SELECT COUNT(*) FROM st_equipos WHERE crm_contacto_id = %s AND activo) AS equipos,
                             (SELECT COUNT(*) FROM st_ordenes WHERE crm_contacto_id = %s) AS ordenes,
                             (SELECT COUNT(*) FROM crm_actividades WHERE contacto_id = %s
                                AND asunto NOT LIKE 'Servicio técnico:%%') AS actividades""", (cid, cid, cid))
            uso = cur.fetchone()
            if uso['equipos'] or uso['ordenes'] or uso['actividades']:
                contactos_conservados.append(cid)
                continue
            cur.execute('UPDATE crm_contactos SET activo = FALSE, updated_at = NOW() WHERE id = %s AND activo', (cid,))
            if cur.rowcount:
                archivados.append(cid)
        for eid in set(retirar) | {c['equipo_id'] for c in res.get('equipos_completados') or []} | \
                {p['equipo_id'] for p in res.get('planes') or []} | set(res.get('equipos_existentes') or []) | \
                {x['equipo_id'] for x in res.get('versiones') or []}:
            st._evento(cur, None, eid, 'equipo', f'Importación #{lote_id} deshecha', usuario_id)
        res['deshecho'] = {'equipos_retirados': len(retirar), 'equipos_conservados': conservados,
                           'datos_devueltos': devueltos, 'contactos_archivados': len(archivados),
                           'contactos_conservados': len(contactos_conservados),
                           'carpetas_archivadas': carpetas_archivadas}
        _guardar(cur, lote_id, estado='deshecha', resultado=res, deshecho_por=usuario_id)
        cur.execute('UPDATE st_importaciones SET deshecho_en = NOW() WHERE id = %s', (lote_id,))
    return res['deshecho']


def se_puede_releer(lote):
    """¿Hay partes que se leyeron sin IA (y sin correcciones a mano) para leerlas con ella?"""
    editadas = {(f.get('archivo_id'), f.get('parte')) for f in lote['filas'] if f.get('editada')}
    return any(p['ia'] in ('sin_ia', 'fallo', 'omitida') and (a['id'], p['i']) not in editadas
               and (p['tipo'] == 'texto' or p.get('sin_mapear'))
               for a in lote['archivos'] for p in a['partes']) \
        or any((a.get('empresa') or {}).get('ia') in ('sin_ia', 'fallo', 'omitida') for a in lote['archivos'])


def equipos_del_resultado(lote):
    """Equipos que creó o completó la importación, para la pantalla final."""
    res = lote.get('resultado') or {}
    creados = res.get('equipos_creados') or []
    completados = [c['equipo_id'] for c in res.get('equipos_completados') or []]
    ids = list(dict.fromkeys(creados + completados + [int(x) for x in (res.get('filas') or {}).values()]))
    if not ids:
        return []
    with get_db_cursor(dict_cursor=True) as cur:
        cur.execute("""SELECT e.id, e.tipo, e.marca, e.modelo, e.serial, e.imei, e.activo, e.mant_proximo,
                              c.nombre AS cliente_nombre,
                              (SELECT COUNT(*) FROM st_mantenimientos m WHERE m.equipo_id = e.id
                                 AND m.importacion_id = %s) AS mantenimientos
                       FROM st_equipos e LEFT JOIN crm_contactos c ON c.id = e.crm_contacto_id
                       WHERE e.id = ANY(%s) ORDER BY e.id""", (lote['id'], ids))
        filas = [dict(r) for r in cur.fetchall()]
    versiones = {x['equipo_id']: x for x in res.get('versiones') or []}
    for f in filas:
        f['descripcion'] = st.descripcion_equipo(f)
        f['nuevo'] = f['id'] in creados
        hecho = versiones.get(f['id']) or {}
        f['actualizado'], f['version'] = bool(hecho.get('cambios')), hecho.get('version_nueva')
        f['completado'] = f['id'] in completados          # importaciones de antes de las versiones
    return filas


# ── Plantilla de Excel ──────────────────────────────────────────
PLANTILLA_COLUMNAS = ('Tipo de equipo', 'Marca', 'Modelo', 'Serial', 'IMEI', 'Procesador', 'RAM', 'Disco',
                      'Sistema operativo', 'Cliente', 'Teléfono', 'Correo', 'Fecha último mantenimiento',
                      'Qué se hizo', 'Próximo mantenimiento', 'Frecuencia (meses)', 'Observaciones')


def plantilla_excel():
    """Libro de Excel vacío con los títulos que el importador reconoce."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Equipos'
    hoja.append(list(PLANTILLA_COLUMNAS))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='1F3A8A')
    for k, titulo in enumerate(PLANTILLA_COLUMNAS, start=1):
        hoja.column_dimensions[_letra(k - 1)].width = max(14, len(titulo) + 4)
    nombres = ','.join(n.split(' (')[0] for _c, n, _i, _x in tipos.TIPOS)
    lista = DataValidation(type='list', formula1=f'"{nombres}"', allow_blank=True)
    hoja.add_data_validation(lista)
    lista.add('A2:A2000')
    hoja.freeze_panes = 'A2'
    ayuda = libro.create_sheet('Instrucciones')
    for linea in ('Cómo llenar la hoja «Equipos» (esta hoja no se importa):',
                  '1. Una fila por equipo. Si un equipo tuvo varios mantenimientos, repite la fila con el mismo '
                  'serial: se juntan en una sola ficha con toda su hoja de vida.',
                  '2. Fechas como 15/03/2025. Frecuencia en meses (6 = cada seis meses).',
                  '3. El cliente se busca en el CRM por correo o teléfono; si no está, se crea.',
                  '4. Puedes dejar vacío lo que no sepas. No pongas claves ni contraseñas.',
                  'Ejemplo: Portátil | HP | 240 G8 | 5CD1234XYZ | | Core i5-1135G7 | 8 GB | SSD 256 GB | Windows 11 | '
                  'Ana Gómez | 300 123 4567 | ana@correo.com | 15/03/2025 | Limpieza y cambio de pasta térmica | '
                  '15/09/2025 | 6 | Teclado con una tecla dura'):
        ayuda.append([linea])
    ayuda.column_dimensions['A'].width = 140
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
