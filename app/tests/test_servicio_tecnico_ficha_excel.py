# -*- coding: utf-8 -*-
"""Servicio Técnico: el libro de fichas técnicas en Excel del taller.

Un libro con la misma forma que los que usa el taller: una hoja «Control
Mantenimiento» (la empresa arriba y una tabla con el número de ficha de cada
equipo) y una hoja por equipo con la ficha (código, fechas, descripción con la
foto DENTRO de la celda, características, vida útil con ✅, fabricante, problema y
requerimientos), más la plantilla en blanco. Datos inventados.

Se lee completo sin IA, cada fila de la tabla se une con su hoja (por el número
de ficha), las fotos salen de su hoja (no el logo) y todo se importa y se deshace.
Contra la base de pruebas; lo creado se borra al final."""
import io
import random
import re
import string
import uuid
import zipfile
from datetime import datetime

import pytest

from tests.test_servicio_tecnico import modulo  # noqa: F401  (fixture)
from tests.test_servicio_tecnico_importar import _archivo, sin_ia  # noqa: F401  (fixtures)

MARCA = 'PYTEST-FX-' + uuid.uuid4().hex[:6].upper()
EMPRESA = 'CONTROLADORA ' + ''.join(random.choice(string.ascii_uppercase) for _ in range(6))
COD1, COD2 = f'{MARCA}-F01', f'{MARCA}-F02'
ROJO = 'FFFF4F4F'


def _imagen(color, lado=320):
    from PIL import Image
    salida = io.BytesIO()
    Image.effect_noise((lado, lado), 60).convert('RGB').save(salida, 'JPEG', quality=90)
    return salida.getvalue()


def _png_logo():
    from PIL import Image
    salida = io.BytesIO()
    Image.new('RGB', (40, 20), (18, 44, 148)).save(salida, 'PNG')
    return salida.getvalue()


def _hoja_ficha(ws, codigo, roja, sin_valores=False):
    from openpyxl.styles import Font, PatternFill
    v = (lambda x: None) if sin_valores else (lambda x: x)
    ws['C2'], ws['E2'], ws['F2'] = 'Taller QA Colombia', 'CODIGO', v(codigo)
    ws['E3'], ws['F3'] = 'FECHA', v(datetime(2025, 12, 2))
    ws['E4'], ws['F4'] = 'FECHA ULTIMO MANTENIMIENTO', v(datetime(2026, 8, 12))
    ws['E5'], ws['F5'] = 'FECHA PROXIMO MANTENIMIENTO', v(datetime(2027, 8, 12))
    ws['C6'], ws['E6'], ws['F6'] = 'FICHA TECNICA', 'VERSION', v(1)
    ws['A8'], ws['D8'] = 'DESCRIPCIÓN GENERAL DEL PRODUCTO', '#VALUE!'     # la foto va DENTRO de D8
    ws['A9'] = v('Equipo corporativo "Todo en Uno" para la gestión contable.')
    ws['A10'] = 'CARACTERÍSTICAS GLOBALES'
    filas = [('Sistema Operativo', 'Microsoft Windows 11 PRO'), ('Fabricante del Sistema', 'HP'),
             ('Modelo', 'Hp Pavilion All-in-one 24-xa1xx'), ('Nombre de Usuario', 'CONTABILIDAD'),
             ('Procesador', 'AMD Ryzen 7 3750H with Radeon Vega Mobile Gfx'), ('Memoria RAM', '8GB DDR4'),
             ('Memoria física disponible:', '4,38 GB'), ('Tarjeta Grafica', 'AMD Radeon RX Vega 10'),
             ('Unidad de Disco', 'SSD 239GB M.2')]
    for k, (etiqueta, valor) in enumerate(filas):
        ws.cell(11 + k, 1, etiqueta)
        ws.cell(11 + k, 4, v(valor))
    ws.cell(20, 4, v('HDD 932GB'))                         # segundo disco: sigue el dato de arriba
    ws['A22'] = 'VIDA UTIL'
    ws['A23'] = v('La vida útil se proyecta a largo plazo. Es fundamental:\n\n✅ Mantener el equipo limpio. '
                  '\n✅ Usar un regulador de voltaje o UPS. \n✅ Ubicar en un área ventilada.')
    ws['A24'], ws['C24'] = 'DATOS DEL FABRICANTE', v('Hp')
    ws['C25'] = v('https://www.hp.com/')
    ws['A26'], ws['C26'] = 'DESCRIPCION DEL PROBLEMA', v('Memoria RAM insuficiente.' if roja else 'N/A')
    ws['A27'], ws['C27'] = 'DETALLE TECNICO DE REQUERIMIENTOS', v('Mantenimiento Preventivo cada 4 meses')
    if roja and not sin_valores:
        for ref in ('A8', 'A10', 'A22'):
            ws[ref].fill = PatternFill('solid', fgColor=ROJO)
        ws['D16'].font = Font(color='FFFF0000')            # «Memoria RAM» en rojo


def _libro():
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    libro = Workbook()
    control = libro.active
    control.title = 'Control Mantenimiento'
    control['D2'], control['F2'], control['G2'] = 'Taller QA Colombia', 'CODIGO', 'CYBER-CONTROL-001'
    control['F3'], control['G3'] = 'EMPRESA', EMPRESA
    control['F4'], control['G4'] = 'FECHA ULTIMA ACTUALIZACION', datetime(2026, 8, 12)
    control['D5'], control['F5'], control['G5'] = 'Control de Mantenimientos', 'VERSION', 1
    for j, titulo in enumerate(['Referencia Equipo', 'Area Correspondiente', 'Numeración último mantenimiento',
                                'Nro Ficha Tecnica', 'Fecha revisión Anterior', 'Proxima revisión'], start=2):
        control.cell(6, j, titulo)
    for i, (ref, area, cod) in enumerate([('HP Pavilion All-in-One 24-xa1xx', 'Contabilidad', COD1),
                                          ('Matica MC310 (Impresora de carnés)', 'Carnetización', COD2)], start=7):
        for j, valor in enumerate([ref, area, 1, cod, datetime(2026, 8, 12), datetime(2027, 8, 12)], start=2):
            control.cell(i, j, valor)
    contab = libro.create_sheet('CONTABILIDAD')
    _hoja_ficha(contab, COD1, roja=True)
    mc = libro.create_sheet('MC310')
    _hoja_ficha(mc, COD2, roja=False)
    mc['A9'] = 'Impresora de carnés de retransferencia para la carnetización.'
    mc['D13'] = 'Matica MC310'
    foto_hoja = XLImage(io.BytesIO(_imagen('azul')))
    mc.add_image(foto_hoja, 'D9')                          # foto puesta sobre la hoja
    blanco = libro.create_sheet('Formato en blanco')
    _hoja_ficha(blanco, '', roja=False, sin_valores=True)
    for ws in libro.worksheets:                            # el logo del taller en todas las hojas
        ws.add_image(XLImage(io.BytesIO(_png_logo())), 'A1')
    salida = io.BytesIO()
    libro.save(salida)
    return _con_imagen_en_celda(salida.getvalue(), 'xl/worksheets/sheet2.xml', 'D8', _imagen('verde'))


def _con_imagen_en_celda(datos, hoja_xml, celda, imagen):
    """Pone una imagen DENTRO de una celda, como lo guarda Excel 365 (richData)."""
    entrada = zipfile.ZipFile(io.BytesIO(datos))
    salida_b = io.BytesIO()
    with zipfile.ZipFile(salida_b, 'w', zipfile.ZIP_DEFLATED) as salida:
        for item in entrada.infolist():
            contenido = entrada.read(item.filename)
            if item.filename == hoja_xml:
                xml = contenido.decode('utf-8')
                xml = re.sub(r'<c r="' + celda + r'"', '<c r="' + celda + '" vm="1"', xml, count=1)
                contenido = xml.encode('utf-8')
            salida.writestr(item, contenido)
        salida.writestr('xl/media/foto_en_celda.jpeg', imagen)
        salida.writestr('xl/metadata.xml', (
            '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<metadataTypes count="1"><metadataType name="XLRICHVALUE"/></metadataTypes>'
            '<futureMetadata name="XLRICHVALUE" count="1"><bk><extLst><ext uri="{3e2802c4}">'
            '<xlrd:rvb i="0"/></ext></extLst></bk></futureMetadata>'
            '<valueMetadata count="1"><bk><rc t="1" v="0"/></bk></valueMetadata></metadata>'))
        salida.writestr('xl/richData/rdrichvaluestructure.xml', (
            '<rvStructures count="1"><s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/>'
            '<k n="CalcOrigin" t="i"/></s></rvStructures>'))
        salida.writestr('xl/richData/rdrichvalue.xml', '<rvData count="1"><rv s="0"><v>0</v><v>5</v></rv></rvData>')
        salida.writestr('xl/richData/richValueRel.xml', (
            '<richValueRels xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<rel r:id="rId1"/></richValueRels>'))
        salida.writestr('xl/richData/_rels/richValueRel.xml.rels', (
            '<Relationships><Relationship Id="rId1" Type="image" Target="../media/foto_en_celda.jpeg"/>'
            '</Relationships>'))
    return salida_b.getvalue()


@pytest.fixture()
def limpiar(cursor):
    lotes = []
    yield lotes
    with cursor() as cur:
        cur.execute("SELECT id FROM crm_contactos WHERE nombre ILIKE %s", (EMPRESA + '%',))
        contactos = [r['id'] for r in cur.fetchall()]
        cur.execute("""SELECT id FROM st_equipos WHERE crm_contacto_id = ANY(%s) OR importacion_id = ANY(%s)
                       OR ficha_codigo ILIKE %s""", (contactos, lotes, MARCA + '%'))
        equipos = [r['id'] for r in cur.fetchall()]
        cur.execute("SELECT DISTINCT empresa_id FROM st_equipos WHERE id = ANY(%s) AND empresa_id IS NOT NULL", (equipos,))
        empresas = [r['empresa_id'] for r in cur.fetchall()]
        for tabla in ('st_documentos', 'st_mantenimientos', 'st_seguimientos', 'st_fotos', 'st_eventos'):
            cur.execute(f'DELETE FROM {tabla} WHERE equipo_id = ANY(%s)', (equipos,))
        cur.execute('DELETE FROM st_documentos WHERE importacion_id = ANY(%s)', (lotes,))
        cur.execute('DELETE FROM st_equipos WHERE id = ANY(%s)', (equipos,))
        cur.execute("""DELETE FROM st_empresas x WHERE (x.id = ANY(%s) OR x.nombre ILIKE %s)
                       AND NOT EXISTS (SELECT 1 FROM st_equipos e WHERE e.empresa_id = x.id)""", (empresas, EMPRESA + '%'))
        cur.execute('DELETE FROM st_importaciones WHERE id = ANY(%s)', (lotes,))
        cur.execute('DELETE FROM crm_actividades WHERE contacto_id = ANY(%s)', (contactos,))
        cur.execute('DELETE FROM crm_contactos WHERE id = ANY(%s)', (contactos,))


def test_el_libro_se_lee_completo_sin_ia():
    from services import servicio_tecnico_ficha as ficha
    from services import servicio_tecnico_importar as imp
    datos = _libro()
    partes = imp.partes_de_archivo('excel', datos)
    assert [p['tipo'] for p in partes] == ['tabla', 'ficha', 'ficha']     # la plantilla en blanco no cuenta
    f = partes[1]['ficha']
    assert f['codigo'] == COD1 and f['fecha'] == '2025-12-02' and f['ultimo'] == '2026-08-12'
    assert f['proximo'] == '2027-08-12' and f['version'] == 1
    assert f['descripcion'].startswith('Equipo corporativo')
    filas = dict(f['filas'])
    assert filas['Sistema Operativo'] == 'Microsoft Windows 11 PRO' and filas['Memoria física disponible'] == '4,38 GB'
    assert filas['Unidad de Disco'] == 'SSD 239GB M.2 / HDD 932GB'           # el segundo disco se une
    assert f['recomendaciones'] == ['Mantener el equipo limpio.', 'Usar un regulador de voltaje o UPS.',
                                    'Ubicar en un área ventilada.']
    assert f['vida_util'].startswith('La vida útil') and f['fabricante'] == 'Hp'
    assert f['fabricante_url'] == 'https://www.hp.com/' and f['problema'] == 'Memoria RAM insuficiente.'
    assert f['requerimientos'] == 'Mantenimiento Preventivo cada 4 meses'
    assert f['estado'] == 'inconveniente' and f['destacar'] == ['Memoria RAM']
    assert partes[2]['ficha']['estado'] == 'bueno' and partes[2]['ficha']['destacar'] == []
    # Fotos: la de dentro de la celda y la de encima de la hoja; el logo de todas las hojas, no.
    assert len(ficha.fotos_de_excel(datos, 'CONTABILIDAD')) == 1 and f['fotos'] == 1
    assert len(ficha.fotos_de_excel(datos, 'MC310')) == 1
    assert ficha.fotos_de_excel(datos, 'Control Mantenimiento') == []
    assert imp._empresa_etiquetada(partes, 'taller qa') == EMPRESA            # «EMPRESA · …» sobre la tabla


def test_el_libro_se_importa_equipo_por_equipo(flask_app, modulo, cursor, limpiar, sin_ia):
    from services import servicio_tecnico_importar as imp
    with flask_app.test_request_context('/'):
        lote_id, _ = imp.crear_lote([_archivo('fichas.xlsx', _libro())], usuario_id=1)
        limpiar.append(lote_id)
        lote = imp.obtener_lote(lote_id)
        assert lote['estado'] == 'revision' and lote['archivos'][0]['empresa']['etiqueta'] == EMPRESA
        vista = imp.revisar(lote)
        listas = [v for v in vista['filas'] if v['estado'] in ('ok', 'aviso')]
        assert len(listas) == 4 and vista['resumen']['incluidas'] == 4
        assert vista['resumen']['nuevos'] == 2 and vista['resumen']['unidas'] == 2   # tabla + hoja = 1 equipo
        raices = [v for v in listas if not v['mismo_que']]
        assert all(v['carpeta']['nombre'] == EMPRESA and v['cliente']['empresa'] == EMPRESA for v in raices)
        res = imp.importar(lote_id, usuario_id=1)
    assert res['resumen']['equipos_creados'] == 2 and len(res['fotos']) == 2 and len(res['fichas']) == 2
    with cursor() as cur:
        cur.execute("""SELECT e.*, x.nombre AS carpeta, c.nombre AS cliente,
                              (SELECT COUNT(*) FROM st_fotos f WHERE f.equipo_id = e.id AND f.activo) AS fotos,
                              (SELECT max(m.fecha) FROM st_mantenimientos m WHERE m.equipo_id = e.id AND m.activo) AS ultimo
                       FROM st_equipos e JOIN st_empresas x ON x.id = e.empresa_id
                       JOIN crm_contactos c ON c.id = e.crm_contacto_id
                       WHERE e.importacion_id = %s ORDER BY e.id""", (lote_id,))
        equipos = cur.fetchall()
    pc, impresora = equipos
    assert pc['ficha_codigo'] == COD1 and pc['ficha']['estado'] == 'inconveniente'
    assert pc['ficha']['recomendaciones'][1] == 'Usar un regulador de voltaje o UPS.'
    assert pc['sistema_operativo'] and pc['procesador'].startswith('AMD Ryzen 7') and pc['ram'].startswith('8')
    assert pc['tipo'] == 'computador' and pc['carpeta'] == EMPRESA and pc['cliente'] == EMPRESA
    assert pc['fotos'] == 1 and str(pc['ultimo']) == '2026-08-12' and str(pc['mant_proximo']) == '2027-08-12'
    assert impresora['ficha_codigo'] == COD2 and impresora['tipo'] == 'impresora' and impresora['fotos'] == 1
    assert impresora['marca'] == 'Matica'
    with flask_app.test_request_context('/'):
        d = imp.deshacer(lote_id, usuario_id=1)
    assert d['equipos_retirados'] == 2 and d['carpetas_archivadas'] == 1


def test_la_pagina_del_equipo_muestra_la_ficha(flask_app, modulo, client, cursor, limpiar, sin_ia):
    """La página del equipo se ve como la ficha del Excel: características,
    recomendaciones, último y próximo mantenimiento; y no pide el serial ni
    «datos técnicos» cuando la ficha ya los trae."""
    from services import servicio_tecnico_importar as imp
    from tests.conftest import _login
    with flask_app.test_request_context('/'):
        lote_id, _ = imp.crear_lote([_archivo('fichas.xlsx', _libro())], usuario_id=1)
        limpiar.append(lote_id)
        res = imp.importar(lote_id, usuario_id=1)
    _login(client, 2)
    html = client.get(f"/admin/servicio-tecnico/equipo/{res['equipos_creados'][0]}").get_data(as_text=True)
    assert 'class="st-card st-seccion st-fv st-fv-roja"' in html and f'Código {COD1}' in html
    assert 'Microsoft Windows 11 PRO' in html and 'st-fv-resaltada' in html            # RAM en rojo
    assert 'Usar un regulador de voltaje o UPS.' in html and 'Recomendaciones que se le dejaron' in html
    assert '12/08/2026' in html and '12/08/2027' in html and 'Último mantenimiento' in html
    assert 'Falta el serial' not in html and 'datos técnicos' not in html            # la ficha los trae
    assert 'id="registrar-mant"' in html and 'href="#registrar-mant"' in html
