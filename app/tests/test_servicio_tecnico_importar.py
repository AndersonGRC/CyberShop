# -*- coding: utf-8 -*-
"""Servicio Técnico: traer equipos desde Excel, CSV y PDF (importador).

Lectura con reglas (sin IA), lo que le llega a la IA (nada personal ni
claves), validación de lo que la IA propone, revisión, importación en una sola
transacción, deshacer sin borrar nada, documentos de la ficha y pantallas.

Contra la base de pruebas. Todo lo creado lleva la marca PYTEST-STI-<uuid> y se
borra al final; el interruptor del módulo vuelve a quedar como estaba."""
import io
import json
import uuid
from datetime import date, datetime, timedelta

import pytest
from werkzeug.datastructures import FileStorage

from tests.conftest import _login
from tests.test_servicio_tecnico import IMEI_OK, modulo  # noqa: F401  (fixture)

MARCA = 'PYTEST-STI-' + uuid.uuid4().hex[:6]
NEGOCIO = {'telefonos': ['6015550000'], 'correos': ['taller@ejemplo.com'], 'nombre': 'taller de prueba'}


# ── Archivos de prueba ──────────────────────────────────────────
def _xlsx(filas, titulo='Equipos', hojas=None):
    from openpyxl import Workbook
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo
    for f in filas:
        hoja.append(list(f))
    for nombre, filas2 in (hojas or {}).items():
        otra = libro.create_sheet(nombre)
        for f in filas2:
            otra.append(list(f))
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _pdf(lineas):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    salida = io.BytesIO()
    c = canvas.Canvas(salida, pagesize=letter)
    y = 750
    for linea in lineas:
        c.drawString(50, y, linea)
        y -= 20
    c.save()
    return salida.getvalue()


def _pdf_escaneado():
    from PIL import Image
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    img = io.BytesIO()
    Image.new('RGB', (200, 150), (230, 230, 230)).save(img, 'PNG')
    img.seek(0)
    salida = io.BytesIO()
    c = canvas.Canvas(salida, pagesize=letter)
    c.drawImage(ImageReader(img), 60, 400, width=200, height=150)
    c.save()
    return salida.getvalue()


def _archivo(nombre, datos):
    return FileStorage(stream=io.BytesIO(datos), filename=nombre)


def _filas(clase, datos, nombre='archivo'):
    from services import servicio_tecnico_importar as imp
    salida = []
    for i, parte in enumerate(imp.partes_de_archivo(clase, datos)):
        salida += imp.filas_de_parte(parte, {'id': 1, 'nombre': nombre, 'clase': clase}, i, NEGOCIO)
    return salida


# ── Fixtures ────────────────────────────────────────────────────
@pytest.fixture()
def limpiar(cursor):
    lotes = []
    yield lotes
    with cursor() as cur:
        cur.execute("SELECT id FROM crm_contactos WHERE nombre LIKE %s", (MARCA + '%',))
        contactos = [r['id'] for r in cur.fetchall()]
        cur.execute("""SELECT id FROM st_equipos WHERE crm_contacto_id = ANY(%s) OR importacion_id = ANY(%s)
                       OR serial LIKE %s""", (contactos, lotes, MARCA + '%'))
        equipos = [r['id'] for r in cur.fetchall()]
        cur.execute("SELECT id FROM st_ordenes WHERE equipo_id = ANY(%s)", (equipos,))
        ordenes = [r['id'] for r in cur.fetchall()]
        cur.execute("DELETE FROM st_documentos WHERE equipo_id = ANY(%s) OR importacion_id = ANY(%s)", (equipos, lotes))
        cur.execute("DELETE FROM st_mantenimientos WHERE equipo_id = ANY(%s)", (equipos,))
        cur.execute("DELETE FROM st_seguimientos WHERE equipo_id = ANY(%s) OR orden_id = ANY(%s)", (equipos, ordenes))
        cur.execute("DELETE FROM st_fotos WHERE equipo_id = ANY(%s)", (equipos,))
        cur.execute("DELETE FROM st_eventos WHERE equipo_id = ANY(%s) OR orden_id = ANY(%s)", (equipos, ordenes))
        cur.execute("DELETE FROM st_cambios WHERE equipo_id = ANY(%s)", (equipos,))
        cur.execute("DELETE FROM st_ordenes WHERE id = ANY(%s)", (ordenes,))
        cur.execute("SELECT DISTINCT empresa_id FROM st_equipos WHERE id = ANY(%s) AND empresa_id IS NOT NULL", (equipos,))
        empresas = [r['empresa_id'] for r in cur.fetchall()]
        cur.execute("DELETE FROM st_equipos WHERE id = ANY(%s)", (equipos,))
        cur.execute("""DELETE FROM st_empresas x WHERE (x.id = ANY(%s) OR x.nombre ILIKE %s)
                       AND NOT EXISTS (SELECT 1 FROM st_equipos e WHERE e.empresa_id = x.id)""", (empresas, MARCA + '%'))
        cur.execute("DELETE FROM st_importaciones WHERE id = ANY(%s)", (lotes,))
        cur.execute("DELETE FROM crm_actividades WHERE contacto_id = ANY(%s)", (contactos,))
        cur.execute("DELETE FROM crm_contactos WHERE id = ANY(%s)", (contactos,))


@pytest.fixture()
def dueno(client):
    _login(client, 2)
    return client


@pytest.fixture()
def sin_ia(monkeypatch):
    """La IA apagada: el importador lee con reglas y el lote pasa directo a revisión."""
    from services import servicio_tecnico_ia as st_ia
    monkeypatch.setattr(st_ia, 'estado', lambda: (False, 'La IA está apagada.'))


# ── Títulos, fechas, frecuencias ────────────────────────────────
def test_titulos_de_columna():
    from services import servicio_tecnico_importar as imp
    casos = {'Serial': 'serial', 'N° de serie': 'serial', 'S/N': 'serial', 'Marca del equipo': 'marca',
             'Proc.': 'procesador', 'Memoria RAM (GB)': 'ram', 'Disco duro': 'almacenamiento',
             'Correo electrónico': 'cliente_email', 'Cel': 'cliente_telefono', 'Clave Windows': 'secreto',
             'Contraseña': 'secreto', 'PIN': 'secreto', 'Serial monitor': 'notas',
             'Fecha último mantenimiento': 'mant_fecha', 'Próximo mantenimiento': 'mant_proximo',
             'Fecha de compra': 'notas', 'Trabajo realizado': 'mant_descripcion', 'Columna rara': None}
    for titulo, destino in casos.items():
        assert imp.destino_de(titulo) == destino, titulo


def test_fechas_frecuencias_y_tipo_de_mantenimiento():
    from services import servicio_tecnico_importar as imp
    assert imp.meses_de('Semestral') == 6 and imp.meses_de('cada 4 meses') == 4
    assert imp.meses_de('2 veces al año') == 6 and imp.meses_de('6') == 6 and imp.meses_de('anual') == 12
    assert imp.meses_de('pronto') is None
    fechas = [f[0] for f in imp.fechas_en('15/03/2025, luego 2024-11-02 y el 3 de mayo de 2023')]
    assert fechas == [date(2025, 3, 15), date(2024, 11, 2), date(2023, 5, 3)]
    assert imp.fechas_en('03/15/2025')[0][0] == date(2025, 3, 15)          # si el «mes» pasa de 12, es mes/día
    assert imp.tipo_mantenimiento('Preventivo - cambio de pasta térmica') == 'preventivo'
    assert imp.tipo_mantenimiento('Cambio de pantalla') == 'correctivo'
    assert imp.tipo_mantenimiento('Limpieza general') == 'preventivo'


def test_caracteristicas_en_texto_libre():
    from services import servicio_tecnico_importar as imp
    col, ext = imp.specs_de_texto('Core i5 8250U 8GB RAM 256 SSD Win10 Pro')
    assert col == {'procesador': 'Core i5-8250U', 'ram': '8 GB', 'almacenamiento': 'SSD 256 GB',
                   'sistema_operativo': 'Windows 10 Pro'}
    assert ext == {'tipo_disco': 'SSD SATA'}
    col, ext = imp.specs_de_texto('Windows 11 Home, Ryzen 5 5500U, 16 GB, 1TB HDD')
    assert col['sistema_operativo'] == 'Windows 11 Home' and col['ram'] == '16 GB'
    assert col['almacenamiento'] == 'HDD 1 TB' and ext['tipo_disco'] == 'HDD'


# ── Lectura con reglas (sin IA) ─────────────────────────────────
def test_inventario_en_excel_sin_claves():
    filas = _filas('excel', _xlsx([
        ['INVENTARIO 2025'], [],
        ['No.', 'Cliente', 'Celular', 'Tipo de equipo', 'Marca', 'Modelo', 'Serial', 'RAM', 'Último mantenimiento',
         'Frecuencia', 'Clave Windows', 'Observaciones'],
        [1, 'Laura', '300 123 4567', 'Portátil', 'HP', '240 G8', '5CD1XYZ', 8, datetime(2025, 3, 15), 'Semestral',
         'secreta-777', 'Bisagra floja'],
        [2, 'Carlos', '3109876543', 'Celular', 'Samsung', 'A54', 'R58NXYZ', '', '', '', 'otra-888',
         f'IMEI {IMEI_OK}'],
    ]))
    assert len(filas) == 2
    hp, cel = filas
    assert hp['equipo']['tipo'] == 'portatil' and hp['equipo']['ram'] == '8 GB' and hp['plan']['cada_meses'] == 6
    assert hp['historial'][0]['fecha'] == '2025-03-15' and hp['cliente']['telefono'] == '3001234567'
    assert 'Observaciones: Bisagra floja' in hp['notas']
    assert cel['equipo']['tipo'] == 'celular' and cel['equipo']['imei'] == IMEI_OK  # el IMEI estaba en «Observaciones»
    todo = json.dumps(filas, ensure_ascii=False)
    assert 'secreta-777' not in todo and 'otra-888' not in todo                     # la columna de claves no se importa


def test_historial_con_el_mismo_serial_es_un_solo_equipo():
    filas = _filas('excel', _xlsx([
        ['Fecha', 'Cliente', 'Equipo', 'Serial', 'Tipo de servicio', 'Trabajo realizado', 'Valor'],
        [datetime(2024, 5, 10), 'Ana', 'Portátil Lenovo IdeaPad 3', 'PF2ABC12', 'Preventivo', 'Limpieza', 80000],
        [datetime(2024, 11, 12), 'Ana', 'Portátil Lenovo IdeaPad 3', 'PF2-ABC12', 'Correctivo', 'Cambio de teclado',
         '$ 150.000'],
        ['03/02/2025', 'Jorge', 'Todo en uno HP 22-dd0', 'CN0123XYZ', 'Correctivo', 'Cambio de fuente', 120000],
    ]))
    assert len(filas) == 2
    lenovo, hp = filas
    assert lenovo['equipo']['marca'] == 'Lenovo' and lenovo['equipo']['modelo'] == 'IdeaPad 3'
    assert [(h['fecha'], h['tipo'], h['costo']) for h in lenovo['historial']] == [
        ('2024-05-10', 'preventivo', 80000.0), ('2024-11-12', 'correctivo', 150000.0)]
    assert 'filas 2, 3' in lenovo['origen']
    assert hp['equipo']['tipo'] == 'computador' and hp['historial'][0]['fecha'] == '2025-02-03'


def test_hoja_de_vida_en_excel_con_su_historial():
    filas = _filas('excel', _xlsx([
        ['HOJA DE VIDA DE EQUIPO DE CÓMPUTO'], [],
        ['Empresa:', 'Distribuidora La Esperanza SAS', '', 'NIT:', '900123456-7'],
        ['Área:', 'Contabilidad', '', 'Usuario:', 'Marta Rojas'],
        ['Tipo de equipo:', 'Computador de escritorio'],
        ['Marca', 'Lenovo', 'Modelo', 'ThinkCentre M70s'],
        ['Serial', 'MJ0ABCDE', 'Procesador', 'Intel Core i5-10400'],
        ['Memoria RAM', '16 GB DDR4', 'Disco', 'SSD 512 GB'],
        ['Contraseña', 'Conta2023*'],
        [], ['HISTORIAL DE MANTENIMIENTOS'],
        ['Fecha', 'Tipo', 'Descripción', 'Responsable'],
        [datetime(2024, 2, 20), 'Preventivo', 'Limpieza general', 'Jhon'],
        [datetime(2025, 1, 15), 'Correctivo', 'Cambio de disco por SSD', 'Mario'],
        [], ['Próximo mantenimiento:', datetime(2026, 7, 15)], ['Frecuencia:', 'Semestral'],
    ], titulo='PC-CONTA-01'))
    assert len(filas) == 1
    f = filas[0]
    assert f['equipo']['tipo'] == 'computador' and f['equipo']['marca'] == 'Lenovo'
    assert f['equipo']['modelo'] == 'ThinkCentre M70s' and f['equipo']['serial'] == 'MJ0ABCDE'
    assert f['cliente']['empresa'] == 'Distribuidora La Esperanza SAS'
    assert [(h['fecha'], h['tipo']) for h in f['historial']] == [('2024-02-20', 'preventivo'), ('2025-01-15', 'correctivo')]
    assert f['plan'] == {'cada_meses': 6, 'proximo': '2026-07-15'}
    assert 'Conta2023' not in json.dumps(f, ensure_ascii=False)


def test_pdf_con_texto_y_pdf_escaneado():
    pdf = _pdf(['TALLER - Tel: 601 555 0000', 'Cliente: Sandra Castro', 'Celular: 312 456 7890',
                'Correo: sandra@correo.com', 'Tipo de equipo: Portátil', 'Marca: Asus', 'Modelo: VivoBook 15',
                'Número de serie: N3NRKD012345', 'Procesador: Intel Core i7-1255U', 'Memoria RAM: 16 GB',
                'Historial de mantenimientos', '10/01/2025 Preventivo - Limpieza interna',
                'Próximo mantenimiento: 10/01/2026', 'Frecuencia de mantenimiento: cada 6 meses'])
    f = _filas('pdf', pdf, 'ficha.pdf')[0]
    assert f['cliente'] == {'nombre': 'Sandra Castro', 'telefono': '3124567890', 'email': 'sandra@correo.com'}
    assert f['equipo']['serial'] == 'N3NRKD012345' and f['equipo']['tipo'] == 'portatil'
    assert f['historial'] == [{'fecha': '2025-01-10', 'tipo': 'preventivo', 'descripcion': 'Limpieza interna',
                               'costo': None}]
    assert f['plan'] == {'cada_meses': 6, 'proximo': '2026-01-10'} and f['documento_id'] == 1
    escaneado = _filas('pdf', _pdf_escaneado(), 'factura.pdf')
    assert len(escaneado) == 1 and escaneado[0]['clase'] == 'adjunto'


def test_cliente_del_documento_sin_el_membrete_del_taller():
    from services import servicio_tecnico_importar as imp
    texto = '\n'.join(['TALLER CYBER - Tel: 601 555 0000 - correo: taller@cyber.com - NIT 900111222',
                       'Cliente: Sandra Castro', 'Celular: 312 456 7890', 'Teléfono fijo: 601 444 3322',
                       'Correo: sandra@correo.com', 'C.C. 1.027.150.819'])
    sin_negocio = {'telefonos': [], 'correos': [], 'nombre': ''}
    assert imp.cliente_de_texto(texto, sin_negocio) == {
        'nombre': 'Sandra Castro', 'telefono': '6014443322', 'whatsapp': '3124567890',
        'email': 'sandra@correo.com', 'documento': '1.027.150.819'}
    # Si dice quién es el cliente pero no trae su celular, no se toma el del membrete.
    membrete = '\n'.join(['TALLER Tel: 601 555 0000', 'Cliente: Pedro Ruiz', 'Marca: HP'])
    assert imp.cliente_de_texto(membrete, sin_negocio) == {'nombre': 'Pedro Ruiz'}


def test_aviso_si_el_celular_es_de_otra_persona_en_el_crm(flask_app, cursor, limpiar):
    from services import servicio_tecnico_importar as imp
    with cursor() as cur:
        cur.execute("""INSERT INTO crm_contactos (tipo, nombre, telefono, whatsapp, origen)
                       VALUES ('cliente', %s, '311 222 3344', '311 222 3344', 'pytest') RETURNING id""",
                    (MARCA + ' Ana Vélez',))
        cid = cur.fetchone()['id']
    fila = imp._cerrar(dict(imp._fila_vacia(), n=1, archivo='x.xlsx', origen='fila 2',
                            cliente={'nombre': 'Laura Gómez', 'telefono': '3112223344'},
                            equipo={'tipo': 'portatil', 'marca': 'HP', 'modelo': '240 G8', 'extras': {}}))
    with flask_app.test_request_context('/'):
        with cursor() as cur:
            vista = imp._revisar(cur, [fila], {}, date.today())
    v = vista['filas'][0]
    assert v['cliente_accion'] == 'existente' and v['contacto']['id'] == cid
    assert any('ese celular o correo es de' in texto for nivel, texto in v['mensajes'] if nivel == 'aviso')


# ── IA: qué recibe y qué se acepta ──────────────────────────────
class _IAFalsa:
    MSG_MOTOR_PREPARANDO = 'El motor de IA se está preparando'
    enviado = []
    respuesta = None

    @classmethod
    def _chat(cls, system, user, **kw):
        cls.enviado.append(system + '\n' + user)
        return cls.respuesta, None

    @staticmethod
    def estado_ia():
        return True, None


def test_ia_no_recibe_datos_personales_ni_claves_y_no_inventa(monkeypatch):
    from services import servicio_tecnico_ia as st_ia
    from services import servicio_tecnico_importar as imp
    _IAFalsa.enviado = []
    _IAFalsa.respuesta = json.dumps({'equipos': [{
        'tipo': 'portatil',
        'campos': {'marca': 'Asus', 'modelo': 'VivoBook 15', 'serial': 'INVENTADO999',
                   'procesador': 'Intel Core i9-99999', 'extra.tipo_disco': 'SSD NVMe'},
        'mantenimientos': [{'fecha': '2025-01-10', 'tipo': 'preventivo', 'que_se_hizo': 'Limpieza interna'},
                           {'fecha': '2024-02-02', 'tipo': 'correctivo', 'que_se_hizo': 'Cambio de pantalla'}],
        'proximo_mantenimiento': '2026-01-10', 'cada_meses': 6, 'resumen': 'Portátil Asus en buen estado'}]})
    monkeypatch.setattr(st_ia, '_ia', lambda: _IAFalsa)
    texto = ('Cliente: Sandra Milena Castro\nCelular: 312 456 7890\nCorreo: sandra@correo.com\n'
             'C.C. 1.027.150.819\nContraseña: Gato2024*\nTipo de equipo: Portátil\nMarca: Asus\n'
             'Modelo: VivoBook 15 con disco NVMe\nNúmero de serie: N3NRKD012345\nHistorial\n'
             '10/01/2025 Preventivo - Limpieza interna\nPróximo mantenimiento: 10/01/2026\nFrecuencia: semestral')
    parte = {'tipo': 'texto', 'titulo': 'PDF', 'texto': texto, 'paginas': [texto]}
    reglas = imp._filas_de_texto(parte, NEGOCIO)
    ia = st_ia.interpretar_documento(imp.sin_secretos(texto), nombres=imp._nombres_para_tapar(parte, reglas))
    enviado = '\n'.join(_IAFalsa.enviado)
    for dato in ('Sandra', 'Milena', 'Castro', '312 456 7890', 'sandra@correo.com', '1.027.150.819', 'Gato2024'):
        assert dato not in enviado, dato
    eq = ia['equipos'][0]
    assert 'serial' not in eq['columnas'] and 'procesador' not in eq['columnas']   # no están en el documento
    assert eq['columnas'] == {'marca': 'Asus', 'modelo': 'VivoBook 15'}
    assert [h['fecha'] for h in eq['historial']] == ['2025-01-10']                # 2024-02-02 no está
    assert eq['proximo'] == '2026-01-10' and eq['cada_meses'] == 6
    assert ia['descartados'] >= 3 and 'no se usaron' in ia['aviso']
    filas = imp.filas_de_parte(parte, {'id': 7, 'nombre': 'ficha.pdf', 'clase': 'pdf'}, 0, NEGOCIO, ia=ia)
    assert len(filas) == 1 and filas[0]['fuente'] == 'ia' and filas[0]['resumen'] == 'Portátil Asus en buen estado'
    assert filas[0]['equipo']['serial'] == 'N3NRKD012345'                         # lo de las reglas manda


def test_ia_columnas_raras_con_ejemplos_sin_nombres(monkeypatch):
    from services import servicio_tecnico_ia as st_ia
    from services import servicio_tecnico_importar as imp
    _IAFalsa.enviado = []
    _IAFalsa.respuesta = json.dumps({'columnas': {'2': 'procesador', '3': 'inventado', '4': 'notas'}})
    monkeypatch.setattr(st_ia, '_ia', lambda: _IAFalsa)
    parte = imp.partes_de_archivo('excel', _xlsx([
        ['Cliente', 'Marca', 'Micro del PC', 'Cosa', 'Resp. del equipo'],
        ['Juan Pérez', 'HP', 'Core i3-1115G4', 'abc', 'Carlos Gómez'],
    ]))[0]
    columnas, sin_mapear = imp._columnas_de_tabla(parte)
    assert sin_mapear == [2, 3, 4]
    nombres = ('Juan Pérez',)
    pedido = [{'col': j, 'encabezado': parte['titulos'][j], 'ejemplos': imp._ejemplos_seguros(parte, j, nombres)}
              for j in sin_mapear]
    mapeo = st_ia.mapear_columnas(pedido, imp.DESTINOS)
    assert mapeo['mapeo'] == {2: 'procesador', 4: 'notas'}                       # «inventado» no es un campo
    enviado = '\n'.join(_IAFalsa.enviado)
    assert 'Juan' not in enviado and 'Carlos Gómez' not in enviado and 'Core i3-1115G4' in enviado
    filas = imp._filas_de_tabla(parte, mapeo['mapeo'])
    assert filas[0]['equipo']['procesador'] == 'Core i3-1115G4'


# ── Lote completo: revisar, importar y deshacer ─────────────────
def _contar(cur):
    cuentas = {}
    for tabla in ('st_equipos', 'st_mantenimientos', 'st_documentos', 'crm_contactos', 'st_seguimientos'):
        cur.execute(f'SELECT COUNT(*) AS n FROM {tabla}')
        cuentas[tabla] = cur.fetchone()['n']
    return cuentas


def test_importar_y_deshacer_sin_borrar(flask_app, modulo, cursor, limpiar, sin_ia):
    from services import servicio_tecnico_importar as imp
    from services import servicio_tecnico_service as st
    hoy = date.today()
    serial_viejo = f'{MARCA}-EX1'
    with flask_app.test_request_context('/'):
        existente = st.registrar_equipo({'cliente': {'nombre': MARCA + ' Previo', 'telefono': '3000000001'},
                                         'equipo': {'tipo': 'portatil', 'marca': 'Dell', 'serial': serial_viejo}}, None)
        inventario = _xlsx([
            ['Cliente', 'Celular', 'Correo', 'Tipo de equipo', 'Marca', 'Modelo', 'Serial', 'RAM',
             'Último mantenimiento', 'Frecuencia', 'Clave'],
            [MARCA + ' Laura', '300 765 4321', 'laura-sti@ejemplo.com', 'Portátil', 'HP', '240 G8', MARCA + '-HP1', 8,
             hoy - timedelta(days=40), 6, 'clave-secreta-77'],
            [MARCA + ' Laura', '300 765 4321', '', 'Celular', 'Samsung', 'A54', '', '', '', '', ''],
            ['', '', '', 'Portátil', 'Dell', 'Latitude 5420', serial_viejo, 16, hoy - timedelta(days=10), '', ''],
            ['', '', '', 'Impresora', 'Epson', 'L3250', MARCA + '-EP1', '', '', '', ''],
        ])
        ficha = _pdf(['Cliente: ' + MARCA + ' Sandra', 'Celular: 312 000 1122', 'Tipo de equipo: Portátil',
                      'Marca: Asus', 'Modelo: VivoBook 15', 'Número de serie: ' + MARCA + '-AS1',
                      'Historial de mantenimientos',
                      (hoy - timedelta(days=100)).strftime('%d/%m/%Y') + ' Preventivo - Limpieza interna',
                      'Próximo mantenimiento: ' + (hoy + timedelta(days=60)).strftime('%d/%m/%Y')])
        lote_id, errores = imp.crear_lote([_archivo('inventario.xlsx', inventario), _archivo('ficha.pdf', ficha),
                                           _archivo('escaneo.pdf', _pdf_escaneado()), _archivo('viejo.xls', b'xx')],
                                          usuario_id=1)
        limpiar.append(lote_id)
        assert len(errores) == 1 and '.xls' in errores[0]
        lote = imp.obtener_lote(lote_id)
        assert lote['estado'] == 'revision'                                    # sin IA: directo a revisar
        vista = imp.revisar(lote)
        por_titulo = {v['titulo']: v for v in vista['filas']}
        assert por_titulo['Portátil · HP 240 G8']['estado'] == 'ok'
        assert por_titulo['Portátil · Dell Latitude 5420']['existente']['id'] == existente
        assert por_titulo['Impresora · Epson L3250']['estado'] == 'error'       # no dice de quién es
        assert por_titulo['PDF sin texto']['clase'] == 'adjunto'
        r = vista['resumen']
        assert (r['incluidas'], r['error'], r['existentes'], r['clientes_nuevos']) == (4, 1, 1, 2)

        # Cliente para los equipos sin dueño, el escaneo va a la impresora, y una corrección.
        imp.guardar_ajustes(lote_id, {'cliente_nombre': MARCA + ' Empresa', 'cliente_telefono': '3015550000'})
        n_epson = por_titulo['Impresora · Epson L3250']['n']
        n_escaneo = por_titulo['PDF sin texto']['n']
        imp.editar_fila(lote_id, n_escaneo, {'accion': 'adjuntar', 'con': str(n_epson)})
        imp.editar_fila(lote_id, n_epson, {'accion': 'corregir', 'modelo': 'L3250 EcoTank'})
        n_samsung = por_titulo['Celular · Samsung A54']['n']
        imp.editar_fila(lote_id, n_samsung, {'accion': 'quitar'})
        imp.editar_fila(lote_id, n_samsung, {'accion': 'incluir'})
        with pytest.raises(st.ErrorServicio):
            imp.editar_fila(lote_id, n_epson, {'accion': 'corregir', 'imei': '123456789012345'})
        vista = imp.revisar(imp.obtener_lote(lote_id))
        r = vista['resumen']
        assert (r['incluidas'], r['error'], r['clientes_nuevos'], r['documentos']) == (5, 0, 3, 2)

        with cursor() as cur:
            antes = _contar(cur)
        res = imp.importar(lote_id, usuario_id=1)
        assert res['resumen']['equipos_creados'] == 4 and res['resumen']['equipos_completados'] == 1
        assert res['resumen']['contactos_creados'] == 3 and res['resumen']['mantenimientos'] == 3
        with pytest.raises(st.ErrorServicio):
            imp.importar(lote_id, usuario_id=1)                                  # no se importa dos veces

    with cursor() as cur:
        cur.execute('SELECT * FROM st_equipos WHERE importacion_id = %s ORDER BY id', (lote_id,))
        creados = {f"{e['marca']} {e['modelo']}": dict(e) for e in cur.fetchall()}
        assert set(creados) == {'HP 240 G8', 'Samsung A54', 'Epson L3250 EcoTank', 'Asus VivoBook 15'}
        assert creados['HP 240 G8']['crm_contacto_id'] == creados['Samsung A54']['crm_contacto_id']  # un solo cliente
        assert creados['HP 240 G8']['mant_cada_meses'] == 6
        assert creados['Asus VivoBook 15']['mant_proximo'] == hoy + timedelta(days=60)
        cur.execute('SELECT * FROM st_equipos WHERE id = %s', (existente,))
        dell = cur.fetchone()
        assert dell['modelo'] == 'Latitude 5420' and dell['ram'] == '16 GB' and dell['marca'] == 'Dell'
        cur.execute("SELECT COUNT(*) AS n FROM st_mantenimientos WHERE importacion_id = %s AND activo", (lote_id,))
        assert cur.fetchone()['n'] == 3
        cur.execute("""SELECT COUNT(*) AS n FROM st_seguimientos WHERE equipo_id = ANY(%s) AND estado = 'pendiente'
                       AND tipo = 'mantenimiento'""", ([e['id'] for e in creados.values()],))
        assert cur.fetchone()['n'] >= 2
        cur.execute('SELECT equipo_id, mime FROM st_documentos WHERE importacion_id = %s ORDER BY id', (lote_id,))
        docs = cur.fetchall()
        destino = {d['equipo_id'] for d in docs if d['mime'] == 'application/pdf'}
        assert destino == {creados['Asus VivoBook 15']['id'], creados['Epson L3250 EcoTank']['id']}
        cur.execute("SELECT nombre, telefono FROM crm_contactos WHERE nombre LIKE %s ORDER BY nombre", (MARCA + '%',))
        assert [c['nombre'] for c in cur.fetchall()] == [MARCA + ' Empresa', MARCA + ' Laura', MARCA + ' Previo',
                                                          MARCA + ' Sandra']
        cur.execute("SELECT string_agg(COALESCE(notas, ''), ' ') AS t FROM st_equipos WHERE importacion_id = %s",
                    (lote_id,))
        assert 'clave-secreta-77' not in cur.fetchone()['t']
        despues = _contar(cur)
        assert despues['st_equipos'] == antes['st_equipos'] + 4

    with flask_app.test_request_context('/'):
        d = imp.deshacer(lote_id, usuario_id=1)
        assert d['equipos_retirados'] == 4 and d['contactos_archivados'] == 3 and d['datos_devueltos'] >= 2
    with cursor() as cur:
        assert _contar(cur) == despues                                          # nada se borró
        cur.execute('SELECT COUNT(*) AS n FROM st_equipos WHERE importacion_id = %s AND activo', (lote_id,))
        assert cur.fetchone()['n'] == 0
        cur.execute('SELECT activo, modelo, ram FROM st_equipos WHERE id = %s', (existente,))
        dell = cur.fetchone()
        assert dell['activo'] and dell['modelo'] is None and dell['ram'] is None  # vuelve a como estaba
        cur.execute('SELECT COUNT(*) AS n FROM st_mantenimientos WHERE importacion_id = %s AND activo', (lote_id,))
        assert cur.fetchone()['n'] == 0
        cur.execute('SELECT COUNT(*) AS n FROM st_documentos WHERE importacion_id = %s AND activo', (lote_id,))
        assert cur.fetchone()['n'] == 0
        cur.execute("SELECT COUNT(*) AS n FROM crm_contactos WHERE nombre LIKE %s AND activo", (MARCA + '%',))
        assert cur.fetchone()['n'] == 1                                          # solo «Previo», que no era del lote
        cur.execute("""SELECT COUNT(*) AS n FROM st_seguimientos WHERE equipo_id = ANY(%s) AND estado = 'pendiente'""",
                    ([e['id'] for e in creados.values()],))
        assert cur.fetchone()['n'] == 0
    with pytest.raises(st.ErrorServicio):
        imp.deshacer(lote_id)


def test_deshacer_conserva_los_equipos_con_ordenes(flask_app, modulo, cursor, limpiar, sin_ia):
    from services import servicio_tecnico_importar as imp
    from services import servicio_tecnico_service as st
    with flask_app.test_request_context('/'):
        lote_id, _ = imp.crear_lote([_archivo('uno.xlsx', _xlsx([
            ['Cliente', 'Celular', 'Tipo', 'Marca', 'Modelo', 'Serial', 'Fecha mantenimiento', 'Trabajo realizado'],
            [MARCA + ' Olga', '3112223344', 'Portátil', 'Acer', 'Aspire 5', MARCA + '-AC1', '01/02/2025', 'Limpieza']]))],
            usuario_id=1)
        limpiar.append(lote_id)
        imp.importar(lote_id, usuario_id=1)
        with cursor() as cur:
            cur.execute('SELECT id, crm_contacto_id FROM st_equipos WHERE importacion_id = %s', (lote_id,))
            eq = cur.fetchone()
        st.crear_orden({'crm_contacto_id': eq['crm_contacto_id'], 'equipo_id': eq['id'],
                        'falla_reportada': 'No enciende'}, None)
        d = imp.deshacer(lote_id, usuario_id=1)
        assert d['equipos_retirados'] == 0 and d['equipos_conservados'] == [eq['id']]
    with cursor() as cur:
        cur.execute('SELECT activo FROM st_equipos WHERE id = %s', (eq['id'],))
        assert cur.fetchone()['activo']
        cur.execute('SELECT COUNT(*) AS n FROM st_mantenimientos WHERE equipo_id = %s AND activo', (eq['id'],))
        assert cur.fetchone()['n'] == 1                                         # su hoja de vida se queda
        cur.execute('SELECT activo FROM crm_contactos WHERE id = %s', (eq['crm_contacto_id'],))
        assert cur.fetchone()['activo']


def test_lote_con_ia_lee_parte_por_parte(flask_app, modulo, cursor, limpiar, monkeypatch):
    from services import servicio_tecnico_ia as st_ia
    from services import servicio_tecnico_importar as imp
    monkeypatch.setattr(st_ia, 'estado', lambda: (True, None))
    llamadas = []

    def interpretar(texto, tipo_sugerido=None, nombres=(), hoy=None):
        llamadas.append((texto, nombres))
        if len(llamadas) == 1:                                                  # la primera vez el motor está frío
            return {'equipos': [], 'ia': False, 'aviso': 'preparando', 'reintentar': True, 'descartados': 0}
        return {'equipos': [{'tipo': 'tablet', 'columnas': {'marca': 'Lenovo'}, 'extras': {}, 'historial': [],
                             'proximo': None, 'cada_meses': None, 'resumen': 'Tablet Lenovo'}],
                'ia': True, 'aviso': None, 'reintentar': False, 'descartados': 0}
    monkeypatch.setattr(st_ia, 'interpretar_documento', interpretar)
    empresas = []

    def identificar(encabezado, nombres=(), negocio='', archivo=''):
        empresas.append((encabezado, list(nombres)))
        return {'ia': True, 'aviso': None, 'reintentar': False, 'empresa': None, 'clasificacion': {}}
    monkeypatch.setattr(st_ia, 'identificar_empresa', identificar)
    with flask_app.test_request_context('/'):
        lote_id, _ = imp.crear_lote([_archivo('t.txt', ('Cliente: ' + MARCA + ' Rita\nMarca: Lenovo\nModelo: Tab M10\n'
                                                       'Clave: 1234\nCorreo: rita.qa@ejemplo.com\n').encode('utf-8'))],
                                    usuario_id=1)
        limpiar.append(lote_id)
        assert imp.obtener_lote(lote_id)['estado'] == 'leyendo'
        paso = imp.leer_siguiente(lote_id)                                     # la parte y, después, la empresa
        assert paso['estado'] == 'pendiente' and paso['esperar'] == 20 and paso['pendientes'] == 2
        paso = imp.leer_siguiente(lote_id)
        assert paso['estado'] == 'hecha' and paso['pendientes'] == 1
        paso = imp.leer_siguiente(lote_id)                                     # ¿de qué empresa es? (IA local)
        assert paso['parte'] == 'Empresa' and paso['estado'] == 'hecha' and paso['pendientes'] == 0
        assert imp.leer_siguiente(lote_id)['listo'] is True
        encabezado, dudosos = empresas[-1]
        assert '1234' not in encabezado and 'rita.qa@' not in encabezado      # sin claves ni correos
        assert dudosos == [MARCA + ' Rita']                                     # nombre dudoso: la IA lo clasifica
        texto, nombres = llamadas[-1]
        assert '1234' not in texto                                           # la línea de la clave no sale
        assert MARCA + ' Rita' in nombres                                    # el nombre se tapa antes de la IA
        lote = imp.obtener_lote(lote_id)
        assert lote['estado'] == 'revision'
        fila = lote['filas'][0]
        assert fila['equipo']['tipo'] == 'tablet' and fila['fuente'] == 'ia' and fila['equipo']['modelo'] == 'Tab M10'


# ── Pantallas ───────────────────────────────────────────────────
def test_pantallas_del_importador(flask_app, modulo, dueno, cursor, limpiar, sin_ia):
    html = dueno.get('/admin/servicio-tecnico/equipos/importar').get_data(as_text=True)
    assert 'Traer equipos desde Excel o PDF' in html and 'name="archivos"' in html
    r = dueno.get('/admin/servicio-tecnico/importacion/plantilla.xlsx')
    assert r.status_code == 200 and r.mimetype.endswith('spreadsheetml.sheet')
    from openpyxl import load_workbook

    from services import servicio_tecnico_importar as imp
    libro = load_workbook(io.BytesIO(r.data))
    assert [c.value for c in libro['Equipos'][1]] == list(imp.PLANTILLA_COLUMNAS)
    assert len(imp.partes_de_archivo('excel', r.data)) == 0                      # vacía: no trae equipos de ejemplo

    datos = _xlsx([['Cliente', 'Celular', 'Tipo de equipo', 'Marca', 'Modelo', 'Serial'],
                   [MARCA + ' Nora', '3159990000', 'Tablet', 'Samsung', 'Tab A9', MARCA + '-TB1']])
    r = dueno.post('/admin/servicio-tecnico/equipos/importar', data={'archivos': (io.BytesIO(datos), 'tabs.xlsx')},
                   content_type='multipart/form-data')
    assert r.status_code == 302 and '/importacion/' in r.headers['Location']
    lote_id = int(r.headers['Location'].rstrip('/').split('/')[-1])
    limpiar.append(lote_id)
    base = f'/admin/servicio-tecnico/importacion/{lote_id}'
    html = dueno.get(base).get_data(as_text=True)
    assert 'Revisa lo que encontramos' in html and 'Tablet · Samsung Tab A9' in html and 'Importar 1 equipo' in html
    r = dueno.post(f'{base}/fila/1', data={'accion': 'quitar'})
    assert r.status_code == 302 and r.headers['Location'].endswith('#fila-1')
    assert 'Volver a incluir' in dueno.get(base + '?ver=quitadas').get_data(as_text=True)
    dueno.post(f'{base}/fila/1', data={'accion': 'incluir'})
    dueno.post(f'{base}/fila/1', data={'accion': 'corregir', 'modelo': 'Galaxy Tab A9'})
    assert 'Samsung Galaxy Tab A9' in dueno.get(base).get_data(as_text=True)
    r = dueno.post(f'{base}/importar')
    assert r.status_code == 302
    html = dueno.get(base).get_data(as_text=True)
    assert 'Listo: equipos importados' in html and 'Deshacer la importación' in html
    with cursor() as cur:
        cur.execute('SELECT id FROM st_equipos WHERE importacion_id = %s', (lote_id,))
        equipo_id = cur.fetchone()['id']
    ficha = dueno.get(f'/admin/servicio-tecnico/equipo/{equipo_id}').get_data(as_text=True)
    assert f'importación #{lote_id}' in ficha and 'id="documentos"' in ficha
    assert 'Traer desde Excel o PDF' in dueno.get('/admin/servicio-tecnico/equipos').get_data(as_text=True)
    r = dueno.post(f'{base}/deshacer')
    assert r.status_code == 302 and 'Importación deshecha' in dueno.get(base).get_data(as_text=True)


def test_documentos_de_la_ficha(flask_app, modulo, dueno, cursor, limpiar):
    from services import servicio_tecnico_service as st
    with flask_app.test_request_context('/'):
        equipo_id = st.registrar_equipo({'cliente': {'nombre': MARCA + ' Docs'},
                                         'equipo': {'tipo': 'impresora', 'marca': 'Epson'}}, None)
    url = f'/admin/servicio-tecnico/equipo/{equipo_id}/documentos'
    r = dueno.post(url, data={'documentos': [(io.BytesIO(_pdf(['Factura 001'])), 'factura.pdf'),
                                             (io.BytesIO(b'<html>no es pdf</html>'), 'falso.pdf'),
                                             (io.BytesIO(b'MZ...'), 'virus.exe')],
                              'descripcion': 'Factura de compra'}, content_type='multipart/form-data')
    assert r.status_code == 302 and r.headers['Location'].endswith('#documentos')
    docs = st.documentos_de_equipo(equipo_id)
    assert [d['nombre'] for d in docs] == ['factura.pdf'] and docs[0]['descripcion'] == 'Factura de compra'
    r = dueno.get(f'/admin/servicio-tecnico/documento/{docs[0]["id"]}')
    assert r.status_code == 200 and r.mimetype == 'application/pdf' and r.data.startswith(b'%PDF')
    assert r.headers['Content-Disposition'].startswith('inline') and 'private' in r.headers['Cache-Control']
    assert 'factura.pdf' in dueno.get(f'/admin/servicio-tecnico/equipo/{equipo_id}').get_data(as_text=True)
    r = dueno.post(f'/admin/servicio-tecnico/documento/{docs[0]["id"]}/quitar')
    assert r.status_code == 302
    assert dueno.get(f'/admin/servicio-tecnico/documento/{docs[0]["id"]}').status_code == 404
    with cursor() as cur:
        cur.execute('SELECT COUNT(*) AS n FROM st_documentos WHERE equipo_id = %s', (equipo_id,))
        assert cur.fetchone()['n'] == 1                                           # retirado, no borrado


def test_modulo_apagado_no_deja_importar(modulo, dueno):
    modulo(False)
    assert dueno.get('/admin/servicio-tecnico/equipos/importar').status_code == 302
