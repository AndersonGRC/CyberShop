# -*- coding: utf-8 -*-
"""Servicio Técnico: versiones de la ficha técnica.

Volver a traer el libro de fichas de un cliente no crea equipos repetidos: el
equipo que ya estaba se actualiza y su ficha queda como versión nueva; la
anterior se guarda y se puede ver en PDF. Lo que el archivo no trae se conserva
y las fotos, mantenimientos y piezas no se tocan. Subir otra vez el mismo
archivo no crea otra versión, deshacer devuelve todo y un serial distinto nunca
se junta. Datos inventados, contra la base de pruebas; lo creado se borra al final."""
import io
from datetime import date, datetime

from tests.conftest import _login
from tests.test_servicio_tecnico import modulo  # noqa: F401  (fixture)
from tests.test_servicio_tecnico_ficha_excel import COD1, EMPRESA, _hoja_ficha, _libro, limpiar  # noqa: F401
from tests.test_servicio_tecnico_importar import _archivo, _xlsx, sin_ia  # noqa: F401  (fixtures)

FILAS_NUEVAS = [('Sistema Operativo', 'Microsoft Windows 11 PRO'), ('Fabricante del Sistema', 'HP'),
                ('Modelo', 'Hp Pavilion All-in-one 24-xa1xx'), ('Nombre de Usuario', 'CONTABILIDAD'),
                ('Procesador', 'AMD Ryzen 7 3750H with Radeon Vega Mobile Gfx'), ('Memoria RAM', '16GB DDR4'),
                ('Memoria física disponible:', '9,10 GB'), ('Unidad de Disco', 'SSD 480GB M.2')]


def _libro_actualizado(filas_pc=None):
    """El mismo libro del taller meses después: la ficha de Contabilidad con más
    RAM, otro disco y ya sin la fila de la tarjeta gráfica (azul: sin fallas)."""
    from openpyxl import Workbook
    libro = Workbook()
    control = libro.active
    control.title = 'Control Mantenimiento'
    control['D2'], control['F2'], control['G2'] = 'Taller QA Colombia', 'CODIGO', 'CYBER-CONTROL-001'
    control['F3'], control['G3'] = 'EMPRESA', EMPRESA
    control['F4'], control['G4'] = 'FECHA ULTIMA ACTUALIZACION', datetime(2026, 9, 30)
    control['D5'], control['F5'], control['G5'] = 'Control de Mantenimientos', 'VERSION', 2
    for j, titulo in enumerate(['Referencia Equipo', 'Area Correspondiente', 'Numeración último mantenimiento',
                                'Nro Ficha Tecnica', 'Fecha revisión Anterior', 'Proxima revisión'], start=2):
        control.cell(6, j, titulo)
    for j, valor in enumerate(['HP Pavilion All-in-One 24-xa1xx', 'Contabilidad', 2, COD1, datetime(2026, 9, 30),
                               datetime(2027, 9, 30)], start=2):
        control.cell(7, j, valor)
    _hoja_ficha(libro.create_sheet('CONTABILIDAD'), COD1, roja=False, filas=filas_pc or FILAS_NUEVAS)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _importar(flask_app, limpiar, datos, nombre='fichas.xlsx', antes_de_importar=None):
    from services import servicio_tecnico_importar as imp
    with flask_app.test_request_context('/'):
        lote_id, _ = imp.crear_lote([_archivo(nombre, datos)], usuario_id=1)
        limpiar.append(lote_id)
        vista = imp.revisar(imp.obtener_lote(lote_id))
        if antes_de_importar:
            vista = antes_de_importar(imp, lote_id, vista)
        return lote_id, vista, imp.importar(lote_id, usuario_id=1)


def test_reimportar_deja_version_nueva_y_guarda_la_anterior(flask_app, modulo, cursor, limpiar, sin_ia):
    from services import servicio_tecnico_importar as imp
    _importar(flask_app, limpiar, _libro())
    with cursor() as cur:
        cur.execute("SELECT id FROM st_equipos WHERE ficha_codigo = %s AND activo", (COD1,))
        pc = cur.fetchone()['id']
        cur.execute("SELECT crm_contacto_id FROM st_equipos WHERE id = %s", (pc,))
        contacto = cur.fetchone()['crm_contacto_id']
        # Lo que el taller le hizo después: una pieza cambiada y un mantenimiento a mano.
        cur.execute("INSERT INTO st_cambios (equipo_id, componente, detalle) VALUES (%s, 'Disco', 'Se cambió a SSD')",
                    (pc,))
        cur.execute("""INSERT INTO st_mantenimientos (equipo_id, tipo, fecha, descripcion)
                       VALUES (%s, 'correctivo', %s, 'Cambio de pasta térmica')""", (pc, date(2026, 9, 1)))
        cur.execute('SELECT ficha, ram FROM st_equipos WHERE id = %s', (pc,))
        antes = cur.fetchone()
    assert antes['ficha']['version'] == 1 and antes['ficha']['estado'] == 'inconveniente'

    lote2, vista, res = _importar(flask_app, limpiar, _libro_actualizado())
    fila = next(v for v in vista['filas'] if v['existente'] and not v['mismo_que'])
    assert fila['existente']['id'] == pc and fila['version_actual'] == 1 and fila['version_nueva'] == 2
    cambiadas = {c['etiqueta'] for c in fila['cambios']}
    assert {'Memoria RAM', 'Unidad de Disco', 'Estado'} <= cambiadas and 'Tarjeta Grafica' not in cambiadas
    assert any('versión 2' in texto and 'versión 1 queda guardada' in texto for _n, texto in fila['mensajes'])
    assert res['equipos_creados'] == [] and [x['equipo_id'] for x in res['versiones']] == [pc]
    assert res['resumen']['fichas_actualizadas'] == 1

    with cursor() as cur:
        cur.execute('SELECT * FROM st_equipos WHERE id = %s', (pc,))
        e = cur.fetchone()
        cur.execute('SELECT COUNT(*) AS n FROM st_equipos WHERE crm_contacto_id = %s AND activo', (contacto,))
        assert cur.fetchone()['n'] == 2                                  # el PC y la impresora: nada repetido
        cur.execute('SELECT COUNT(*) AS n FROM st_cambios WHERE equipo_id = %s', (pc,))
        assert cur.fetchone()['n'] == 1                                  # la pieza sigue
        cur.execute('SELECT COUNT(*) AS n FROM st_mantenimientos WHERE equipo_id = %s AND activo', (pc,))
        assert cur.fetchone()['n'] == 3                                  # el de la ficha, el manual y el de la tabla
        cur.execute('SELECT COUNT(*) AS n FROM st_fotos WHERE equipo_id = %s AND activo', (pc,))
        assert cur.fetchone()['n'] == 1                                  # la foto sigue
        cur.execute("""SELECT version, origen, activo, estado FROM st_fichas_versiones WHERE equipo_id = %s
                       ORDER BY id""", (pc,))
        versiones = cur.fetchall()
    f = e['ficha']
    assert f['version'] == 2 and f['fecha'] == date.today().isoformat() and f['estado'] == 'bueno'
    assert f['filas']['Memoria RAM'] == '16GB DDR4' and f['destacar'] == []
    assert f['filas']['Tarjeta Grafica'] == 'AMD Radeon RX Vega 10'       # lo que no trae se conserva
    assert '16' in e['ram'] and e['almacenamiento'].startswith('SSD 480')
    assert [(v['version'], v['origen'], v['activo']) for v in versiones] == [(1, 'anterior', True),
                                                                          (2, 'importacion', True)]
    assert versiones[0]['estado']['ficha']['filas']['Memoria RAM'] == '8GB DDR4'   # la versión 1, guardada

    # El mismo archivo otra vez: ya está al día, no hay versión nueva.
    _l3, vista3, res3 = _importar(flask_app, limpiar, _libro_actualizado(), nombre='otra-vez.xlsx')
    fila3 = next(v for v in vista3['filas'] if v['existente'] and not v['mismo_que'])
    assert fila3['cambios'] == [] and any('no cambia' in texto for _n, texto in fila3['mensajes'])
    assert res3['versiones'] == [] and res3['equipos_creados'] == []
    with cursor() as cur:
        cur.execute('SELECT COUNT(*) AS n FROM st_fichas_versiones WHERE equipo_id = %s AND activo', (pc,))
        assert cur.fetchone()['n'] == 2

    # Deshacer la importación de la versión 2: todo vuelve a la versión 1; nada se borra.
    with flask_app.test_request_context('/'):
        d = imp.deshacer(lote2, usuario_id=1)
    assert d['datos_devueltos'] >= 3
    with cursor() as cur:
        cur.execute('SELECT ficha, ram FROM st_equipos WHERE id = %s', (pc,))
        despues = cur.fetchone()
        cur.execute('SELECT origen, activo FROM st_fichas_versiones WHERE equipo_id = %s ORDER BY id', (pc,))
        assert [(v['origen'], v['activo']) for v in cur.fetchall()] == [('anterior', True), ('importacion', False)]
        cur.execute('SELECT COUNT(*) AS n FROM st_cambios WHERE equipo_id = %s', (pc,))
        assert cur.fetchone()['n'] == 1
    assert despues['ficha'] == antes['ficha'] and despues['ram'] == antes['ram']


def test_la_tabla_sin_numero_se_junta_con_su_hoja(flask_app, modulo, cursor, limpiar, sin_ia):
    """Un libro cuya tabla de control perdió los números de ficha (p. ej. guardado
    con otro programa): las hojas encuentran sus equipos por el número y las filas
    de la tabla se juntan con ellas. No se repite ningún equipo."""
    _importar(flask_app, limpiar, _libro())
    sin_numero = [('HP Pavilion All-in-One 24-xa1xx', 'Contabilidad', ''),
                  ('Matica MC310 (Impresora de carnés)', 'Carnetización', '')]
    _lote, vista, res = _importar(flask_app, limpiar, _libro(filas_control=sin_numero), nombre='sin-numeros.xlsx')
    tabla = [v for v in vista['filas'] if 'Control' in v['origen']]
    assert len(tabla) == 2 and all(v['mismo_que'] for v in tabla) and vista['resumen']['nuevos'] == 0
    assert res['equipos_creados'] == []


def test_lo_que_dice_menos_no_reemplaza_lo_que_hay():
    """Una tabla con menos detalle no borra lo que ya está más completo; un dato
    distinto sí manda; el serial nunca se cambia."""
    from services import servicio_tecnico_importar as imp
    actual = {'tipo': 'computador', 'marca': 'Lenovo', 'modelo': 'IdeaCentre B320 (Touch)', 'ram': '8 GB DDR4',
              'sistema_operativo': 'Windows 11 PRO', 'serial': 'QA-1', 'extras': {'tipo_disco': 'SSD NVMe'},
              'ficha': {}, 'notas': None, 'resumen_ia': None, 'ficha_codigo': None}
    d = imp._equipo_actualizado(actual, {'modelo': 'IdeaCentre B320', 'ram': '16 GB DDR4', 'serial': 'QA-2',
                                         'sistema_operativo': 'Windows 11', 'extras': {'tipo_disco': 'HDD'}},
                                {}, date.today())
    assert d['modelo'] == 'IdeaCentre B320 (Touch)' and d['sistema_operativo'] == 'Windows 11 PRO'
    assert d['ram'] == '16 GB DDR4' and d['serial'] == 'QA-1' and d['extras']['tipo_disco'] == 'HDD'


def test_un_serial_distinto_nunca_se_junta(flask_app, modulo, cursor, limpiar, sin_ia):
    """Mismo cliente y modelo: sin serial en el archivo se actualiza el que estaba;
    con OTRO serial es otro equipo y se registra aparte."""
    from services import servicio_tecnico_service as st
    with flask_app.test_request_context('/'):
        dell = st.registrar_equipo({'cliente': {'nombre': EMPRESA, 'telefono': '3000000777'},
                                    'equipo': {'tipo': 'portatil', 'marca': 'Dell', 'modelo': 'Latitude 5420',
                                               'serial': 'QASER-111', 'ram': '8 GB'}}, None)
    inventario = _xlsx([['Cliente', 'Tipo de equipo', 'Marca', 'Modelo', 'Serial', 'RAM'],
                        [EMPRESA, 'Portátil', 'Dell', 'Latitude 5420', 'QASER-999', 16],
                        [EMPRESA, 'Portátil', 'Dell', 'Latitude 5420', '', 32]])
    _lote, vista, res = _importar(flask_app, limpiar, inventario, nombre='inventario.xlsx')
    con_serial, sin_serial = [v for v in vista['filas'] if v['clase'] == 'equipo']
    assert not con_serial['existente'] and sin_serial['existente']['id'] == dell
    assert res['equipos_existentes'] == [dell] and len(res['equipos_creados']) == 1
    with cursor() as cur:
        cur.execute('SELECT serial, ram FROM st_equipos WHERE id = %s', (dell,))
        e = cur.fetchone()
    assert e['serial'] == 'QASER-111' and e['ram'].startswith('32')       # el serial no se cambia; la RAM sí


def test_escoger_a_mano_cual_era(flask_app, modulo, cursor, limpiar, sin_ia):
    """Si el modelo no coincide, el equipo se crearía nuevo; en la revisión se
    puede escoger cuál de los del cliente es, y ese se actualiza."""
    from services import servicio_tecnico_service as st
    with flask_app.test_request_context('/'):
        epson = st.registrar_equipo({'cliente': {'nombre': EMPRESA, 'telefono': '3000000778'},
                                     'equipo': {'tipo': 'impresora', 'marca': 'Epson', 'modelo': 'L3250'}}, None)
    inventario = _xlsx([['Cliente', 'Tipo de equipo', 'Marca', 'Modelo'],
                        [EMPRESA, 'Impresora', 'Epson', 'EcoTank L4260']])

    def escoger(imp, lote_id, vista):
        fila = next(v for v in vista['filas'] if v['clase'] == 'equipo')
        assert not fila['existente'] and [o['id'] for o in fila['opciones']] == [epson]
        imp.editar_fila(lote_id, fila['n'], {'accion': 'usar_equipo', 'equipo_id': str(epson)})
        vista = imp.revisar(imp.obtener_lote(lote_id))
        fila = next(v for v in vista['filas'] if v['clase'] == 'equipo')
        assert fila['existente']['id'] == epson and any('Escogiste' in t for _n, t in fila['mensajes'])
        return vista
    _lote, _vista, res = _importar(flask_app, limpiar, inventario, nombre='impresoras.xlsx', antes_de_importar=escoger)
    assert res['equipos_creados'] == [] and res['equipos_existentes'] == [epson]
    with cursor() as cur:
        cur.execute('SELECT modelo FROM st_equipos WHERE id = %s', (epson,))
        assert cur.fetchone()['modelo'] == 'EcoTank L4260'
        cur.execute("SELECT estado->'equipo'->>'modelo' AS modelo FROM st_fichas_versiones WHERE equipo_id = %s "
                    "AND origen = 'anterior'", (epson,))
        assert cur.fetchone()['modelo'] == 'L3250'                         # lo de antes, guardado


def test_guardar_a_mano_tambien_guarda_la_version_y_se_ve_en_pdf(flask_app, modulo, client, cursor, limpiar, sin_ia):
    from services import servicio_tecnico_ficha as fichas
    from services import servicio_tecnico_versiones as versiones
    _l, _v, res = _importar(flask_app, limpiar, _libro())
    with cursor() as cur:
        cur.execute("SELECT id FROM st_equipos WHERE ficha_codigo = %s AND activo", (COD1,))
        pc = cur.fetchone()['id']
    with flask_app.test_request_context('/'):
        fichas.guardar(pc, {'estado': 'bueno', 'version': '3', 'etiqueta_0': 'Memoria RAM', 'fila_0': '32 GB DDR4'},
                       usuario_id=1)
        lista = versiones.de_equipo(pc)
    assert [(v['version'], v['origen']) for v in lista] == [(3, 'manual'), (1, 'anterior')] and lista[0]['actual']
    assert any(c['etiqueta'] == 'Memoria RAM' and c['ahora'] == '32 GB DDR4' for c in lista[0]['cambios'])
    _login(client, 2)
    html = client.get(f'/admin/servicio-tecnico/equipo/{pc}').get_data(as_text=True)
    assert 'Historial de la ficha (2 versiones guardadas)' in html and 'La de hoy' in html
    r = client.get(f"/admin/servicio-tecnico/equipo/{pc}/ficha/version/{lista[1]['id']}.pdf")
    assert r.status_code == 200 and r.data.startswith(b'%PDF') and 'version 1' in r.headers['Content-Disposition']
    assert client.get(f'/admin/servicio-tecnico/equipo/{pc}/ficha/version/999999999.pdf').status_code == 404
